from pathlib import Path
from datetime import date, datetime, timedelta
import argparse
import calendar

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUTPUT_DIR = Path(__file__).resolve().parents[1] / "输出"

INDEXES = ["沪深300", "中证500", "标普500", "纳斯达克100"]
CODES = ["000300.SH", "000905.SH", "SPX.GI", "NDX.GI"]
A_SHARE_CODES = ["000300.SH", "000905.SH"]
NON_OPERATING_INDUSTRIES = {"金融", "能源", "房地产"}

MANUAL_INPUTS = {
    # 每月初更新时，Wind 能自动抓 A 股和指数行情估值；美股一致预期公开源不稳定，先在这里手工维护。
    "baseline_start_year": 2025,
    "current_year": 2026,
    "us_profit_expectation_previous": {"SPX.GI": 0.2100, "NDX.GI": None},
    "us_profit_expectation_latest": {"SPX.GI": 0.2260, "NDX.GI": 0.3030},
    "us_q1_profit_growth": {"SPX.GI": 0.2770, "NDX.GI": 0.5100},
    "non_operating_revenue_share": {"000300.SH": 0.40, "000905.SH": 0.10},
    "non_operating_profit_share": {"000300.SH": 0.60, "000905.SH": 0.16},
    "assumed_non_operating_profit_growth": {"000300.SH": 0.05, "000905.SH": 0.00},
    "ex_non_operating_pe": {"000300.SH": 22.7, "000905.SH": 34.2},
    "recommendation": {"000300.SH": "中配", "000905.SH": "高配", "SPX.GI": "低配", "NDX.GI": "高配"},
}

def month_bounds(year, month):
    last_day = calendar.monthrange(year, month)[1]
    return date(year, month, 1), date(year, month, last_day)


def add_months(d, months):
    month = d.month - 1 + months
    year = d.year + month // 12
    month = month % 12 + 1
    day = min(d.day, calendar.monthrange(year, month)[1])
    return date(year, month, day)


def fmt_date(d):
    return d.strftime("%Y-%m-%d")


def yyyymmdd(d):
    return d.strftime("%Y%m%d")


def month_label(d):
    return f"{d.month}月底"


def output_path(target_end):
    return OUTPUT_DIR / f"四大指数配置可视化比较_截止{target_end.year}年{target_end.month}月底.xlsx"


def start_wind():
    from WindPy import w

    w.start()
    return w


def wind_tdays(w, start, end):
    data = w.tdays(fmt_date(start), fmt_date(end), "")
    if data.ErrorCode != 0:
        raise RuntimeError(f"Wind tdays failed: {data.ErrorCode} {getattr(data, 'Data', None)}")
    return [t for t in data.Data[0]]


def last_trading_day_of_month(w, year, month):
    start, end = month_bounds(year, month)
    days = wind_tdays(w, start, end)
    if not days:
        raise RuntimeError(f"Wind did not return trading days for {year}-{month:02d}")
    return days[-1]


def resolve_target_dates(w, explicit_target=None):
    if explicit_target:
        target = datetime.strptime(explicit_target, "%Y-%m-%d").date()
        target_end = target
    else:
        today = date.today()
        previous_month = add_months(date(today.year, today.month, 1), -1)
        target_end = last_trading_day_of_month(w, previous_month.year, previous_month.month)
    previous_month = add_months(date(target_end.year, target_end.month, 1), -1)
    previous_end = last_trading_day_of_month(w, previous_month.year, previous_month.month)
    return target_end, previous_end


def fetch_wsd_series(w, codes, field, start, end, options="PriceAdj=F"):
    data = w.wsd(",".join(codes), field, fmt_date(start), fmt_date(end), options)
    if data.ErrorCode != 0:
        raise RuntimeError(f"Wind wsd {field} failed: {data.ErrorCode} {getattr(data, 'Data', None)}")
    result = {}
    if len(data.Data) == 1 and len(codes) > 1 and len(data.Times) == 1:
        for code, value in zip(codes, data.Data[0]):
            result[code] = [(data.Times[0], value)]
    else:
        for code, values in zip(codes, data.Data):
            result[code] = list(zip(data.Times, values))
    return result


def last_value(series):
    for _, value in reversed(series):
        if value is not None:
            return value
    return None


def value_on(w, codes, field, query_date):
    data = fetch_wsd_series(w, codes, field, query_date, query_date)
    return {code: last_value(series) for code, series in data.items()}


def percentile(values, q):
    clean = sorted(v for v in values if v is not None)
    if not clean:
        return None
    if len(clean) == 1:
        return clean[0]
    pos = (len(clean) - 1) * q
    lower = int(pos)
    upper = min(lower + 1, len(clean) - 1)
    frac = pos - lower
    return clean[lower] * (1 - frac) + clean[upper] * frac


def percentile_rank(values, current):
    clean = [v for v in values if v is not None]
    if not clean or current is None:
        return None
    return sum(1 for v in clean if v <= current) / len(clean)


def pct_change(current, base):
    if current is None or base in (None, 0):
        return None
    return current / base - 1


def implied_earnings(close, pe):
    if close is None or pe in (None, 0):
        return None
    return close / pe


def fetch_wss_percent(w, codes, field, trade_date=None, rpt_date=None):
    options = []
    if trade_date:
        options.append(f"tradeDate={yyyymmdd(trade_date)}")
    if rpt_date:
        options.append(f"rptDate={yyyymmdd(rpt_date)}")
    data = w.wss(codes, field, ";".join(options))
    if data.ErrorCode != 0:
        raise RuntimeError(f"Wind wss {field} failed: {data.ErrorCode} {getattr(data, 'Data', None)}")
    values = data.Data[0]
    return {code: (value / 100 if value is not None else None) for code, value in zip(codes, values)}


def fetch_non_operating_weight(w, index_code, query_date):
    constituent = w.wset(
        "indexconstituent",
        f"date={fmt_date(query_date)};windcode={index_code};field=wind_code,i_weight",
    )
    if constituent.ErrorCode != 0:
        raise RuntimeError(f"Wind indexconstituent failed for {index_code}: {constituent.ErrorCode}")
    codes = constituent.Data[0]
    weights = constituent.Data[1]
    total = 0
    for start in range(0, len(codes), 200):
        batch_codes = codes[start:start + 200]
        batch_weights = weights[start:start + 200]
        industries = w.wss(batch_codes, "wicsname2024", f"tradeDate={yyyymmdd(query_date)};industryType=1;")
        if industries.ErrorCode != 0:
            raise RuntimeError(f"Wind wicsname2024 failed for {index_code}: {industries.ErrorCode}")
        for industry, weight in zip(industries.Data[0], batch_weights):
            if industry in NON_OPERATING_INDUSTRIES and weight is not None:
                total += weight
    return total / 100


def fetch_market_snapshot(w, target_end, previous_end):
    base_year = MANUAL_INPUTS["baseline_start_year"]
    current_year = MANUAL_INPUTS["current_year"]
    baseline_end = last_trading_day_of_month(w, base_year - 1, 12)
    current_year_start_base = last_trading_day_of_month(w, current_year - 1, 12)
    same_month_last_year = last_trading_day_of_month(w, target_end.year - 1, target_end.month)
    ten_year_start = target_end - timedelta(days=3653)

    close_target = value_on(w, CODES, "close", target_end)
    close_previous = value_on(w, CODES, "close", previous_end)
    close_baseline = value_on(w, CODES, "close", baseline_end)
    close_current_year_base = value_on(w, CODES, "close", current_year_start_base)
    close_same_month_last_year = value_on(w, CODES, "close", same_month_last_year)

    pe_target = value_on(w, CODES, "pe_ttm", target_end)
    pe_previous = value_on(w, CODES, "pe_ttm", previous_end)
    pe_baseline = value_on(w, CODES, "pe_ttm", baseline_end)
    pe_current_year_base = value_on(w, CODES, "pe_ttm", current_year_start_base)
    pe_same_month_last_year = value_on(w, CODES, "pe_ttm", same_month_last_year)
    pe_history = fetch_wsd_series(w, CODES, "pe_ttm", ten_year_start, target_end)

    expected_previous = fetch_wss_percent(w, A_SHARE_CODES, "west_netprofit_yoy", trade_date=previous_end)
    expected_latest = fetch_wss_percent(w, A_SHARE_CODES, "west_netprofit_yoy", trade_date=target_end)
    q1_report_date = date(current_year, 3, 31)
    q1_profit_growth = fetch_wss_percent(w, A_SHARE_CODES, "yoynetprofit", trade_date=target_end, rpt_date=q1_report_date)

    non_op_previous = {code: fetch_non_operating_weight(w, code, previous_end) for code in A_SHARE_CODES}
    non_op_latest = {code: fetch_non_operating_weight(w, code, target_end) for code in A_SHARE_CODES}

    rows_by_code = {}
    for code in CODES:
        pe_values = [value for _, value in pe_history[code] if value is not None]
        target_earnings = implied_earnings(close_target[code], pe_target[code])
        baseline_earnings = implied_earnings(close_baseline[code], pe_baseline[code])
        current_year_base_earnings = implied_earnings(close_current_year_base[code], pe_current_year_base[code])
        same_month_earnings = implied_earnings(close_same_month_last_year[code], pe_same_month_last_year[code])
        rows_by_code[code] = {
            "close": close_target[code],
            "previous_close": close_previous[code],
            "pe": pe_target[code],
            "previous_pe": pe_previous[code],
            "previous_baseline_return": pct_change(close_previous[code], close_baseline[code]),
            "baseline_return": pct_change(close_target[code], close_baseline[code]),
            "current_year_return": pct_change(close_target[code], close_current_year_base[code]),
            "baseline_earnings_growth": pct_change(current_year_base_earnings, baseline_earnings),
            "ttm_earnings_yoy": pct_change(target_earnings, same_month_earnings),
            "pe_p25": percentile(pe_values, 0.25),
            "pe_median": percentile(pe_values, 0.50),
            "pe_p75": percentile(pe_values, 0.75),
            "pe_rank": percentile_rank(pe_values, pe_target[code]),
        }

    for code in A_SHARE_CODES:
        rows_by_code[code]["expected_previous"] = expected_previous.get(code)
        rows_by_code[code]["expected_latest"] = expected_latest.get(code)
        rows_by_code[code]["q1_profit_growth"] = q1_profit_growth.get(code)
        rows_by_code[code]["non_op_previous"] = non_op_previous.get(code)
        rows_by_code[code]["non_op_latest"] = non_op_latest.get(code)

    for code in ["SPX.GI", "NDX.GI"]:
        rows_by_code[code]["expected_previous"] = MANUAL_INPUTS["us_profit_expectation_previous"].get(code)
        rows_by_code[code]["expected_latest"] = MANUAL_INPUTS["us_profit_expectation_latest"].get(code)
        rows_by_code[code]["q1_profit_growth"] = MANUAL_INPUTS["us_q1_profit_growth"].get(code)
        rows_by_code[code]["non_op_previous"] = None
        rows_by_code[code]["non_op_latest"] = None

    return {
        "target_end": target_end,
        "previous_end": previous_end,
        "baseline_end": baseline_end,
        "current_year_start_base": current_year_start_base,
        "same_month_last_year": same_month_last_year,
        "rows_by_code": rows_by_code,
    }


def values_for_codes(rows_by_code, key):
    return [rows_by_code[code].get(key) for code in CODES]


def build_rows(snapshot):
    rows_by_code = snapshot["rows_by_code"]
    target_end = snapshot["target_end"]
    previous_end = snapshot["previous_end"]
    base_year = MANUAL_INPUTS["baseline_start_year"]
    current_year = MANUAL_INPUTS["current_year"]
    previous_label = month_label(previous_end)
    target_label = month_label(target_end)
    return [
        ("指数代码", *CODES, "Wind"),
        ("行情截至日期", *[fmt_date(target_end)] * 4, "Wind；月末最后交易日"),
        (f"{base_year}年至今涨幅", *values_for_codes(rows_by_code, "baseline_return"), f"Wind close；相对 {fmt_date(snapshot['baseline_end'])}"),
        (f"{current_year}年初至今涨幅", *values_for_codes(rows_by_code, "current_year_return"), f"Wind close；相对 {fmt_date(snapshot['current_year_start_base'])}"),
        (f"{base_year}年滚动盈利增长", *values_for_codes(rows_by_code, "baseline_earnings_growth"), "由指数点位 / PE_TTM 反推"),
        (f"截止{target_label}滚动盈利同比", *values_for_codes(rows_by_code, "ttm_earnings_yoy"), f"由指数点位 / PE_TTM 反推；对比 {fmt_date(snapshot['same_month_last_year'])}"),
        (f"{current_year}年Q1利润增速", *values_for_codes(rows_by_code, "q1_profit_growth"), "A股：Wind YOYNETPROFIT；标普500/纳斯达克100：手工输入公开资料"),
        (f"{current_year}年利润增速预期：截止{previous_label}", *values_for_codes(rows_by_code, "expected_previous"), "A股：Wind west_netprofit_yoy；美股：手工输入公开资料"),
        (f"{current_year}年利润增速预期：截止{target_label}", *values_for_codes(rows_by_code, "expected_latest"), "A股：Wind west_netprofit_yoy；美股：手工输入公开资料"),
        (f"{current_year}年利润增速预期变动", None, None, None, None, f"公式：截止{target_label}减截止{previous_label}"),
        (f"非实业权重：截止{previous_label}", *values_for_codes(rows_by_code, "non_op_previous"), "A股：Wind 成分权重；金融+能源+房地产；美股留空"),
        (f"非实业权重：截止{target_label}", *values_for_codes(rows_by_code, "non_op_latest"), "A股：Wind 成分权重；金融+能源+房地产；美股留空"),
        ("非实业收入占比", *[MANUAL_INPUTS["non_operating_revenue_share"].get(code) for code in CODES], "手工输入；沿用 Q1 数据"),
        ("非实业利润占比", *[MANUAL_INPUTS["non_operating_profit_share"].get(code) for code in CODES], "手工输入；沿用 Q1 数据"),
        ("假设非实业利润增速", *[MANUAL_INPUTS["assumed_non_operating_profit_growth"].get(code) for code in CODES], "手工输入假设"),
        ("剔除非实业后利润增速", None, None, None, None, "A股公式动态计算；美股留空"),
        ("加权后预期利润增速", None, None, None, None, "A股公式动态计算；美股直接引用整体利润预期"),
        (f"截止{target_label} PE_TTM", *values_for_codes(rows_by_code, "pe"), f"Wind，{fmt_date(target_end)} 收盘估值"),
        ("剔除非实业后 PE", *[MANUAL_INPUTS["ex_non_operating_pe"].get(code) for code in CODES], "手工输入；沿用 Q1 数据"),
        ("指数整体 PE/G", None, None, None, None, "公式：PE_TTM ÷ (加权后预期利润增速 × 100)"),
        ("剔除非实业 PE/G", None, None, None, None, "公式：剔除非实业后 PE ÷ (剔除非实业后利润增速 × 100)"),
        ("过去十年 PE 25%分位数", *values_for_codes(rows_by_code, "pe_p25"), "Wind PE_TTM 过去十年 25% 分位"),
        ("过去十年估值中位数", *values_for_codes(rows_by_code, "pe_median"), "Wind PE_TTM 过去十年中位数"),
        ("过去十年 PE 75%分位数", *values_for_codes(rows_by_code, "pe_p75"), "Wind PE_TTM 过去十年 75% 分位"),
        (f"截止{target_label} PE 十年分位", *values_for_codes(rows_by_code, "pe_rank"), "Wind PE_TTM 过去十年"),
        ("指数整体泡沫评估", None, None, None, None, "综合指数整体 PE/G 与 PE 十年分位；缺利润预期时仅按 PE 分位提示"),
        ("建议配置", *[MANUAL_INPUTS["recommendation"].get(code) for code in CODES], "手工输入；结合利润增速、PE/G 与 PE 分位的主观配置结论"),
    ]

PERCENT_KEYWORDS = (
    "涨幅",
    "滚动盈利",
    "利润增速",
    "权重",
    "收入占比",
    "利润占比",
    "十年分位",
)


def is_percent_label(label):
    return any(keyword in str(label) for keyword in PERCENT_KEYWORDS)


def style_header(ws, row, start_col=1, end_col=6):
    for cell in ws.iter_cols(min_col=start_col, max_col=end_col, min_row=row, max_row=row):
        c = cell[0]
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.font = Font(name="PingFang SC", color="FFFFFF", bold=True)
        c.alignment = Alignment(horizontal="center")


def apply_grid(ws, min_row, max_row, min_col, max_col):
    thin = Side(style="thin", color="D9E1F2")
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def apply_readable_table(ws, min_row, max_row, min_col, max_col):
    ws.sheet_view.showGridLines = False
    zebra_fill = PatternFill("solid", fgColor="FBFCFE")
    first_col_fill = PatternFill("solid", fgColor="F4F7FA")
    source_fill = PatternFill("solid", fgColor="FAFAFA")
    for row_idx in range(min_row + 1, max_row + 1):
        if row_idx % 2 == 0:
            for col_idx in range(min_col, max_col + 1):
                ws.cell(row_idx, col_idx).fill = zebra_fill
        ws.cell(row_idx, min_col).fill = first_col_fill
        ws.cell(row_idx, min_col).font = Font(bold=True, color="1F2937")
        ws.cell(row_idx, max_col).fill = source_fill
        ws.cell(row_idx, max_col).font = Font(size=9, color="666666")
    for row_idx in range(min_row + 1, max_row + 1):
        ws.row_dimensions[row_idx].height = 23


def theme_fill(theme, tint=0.0):
    return PatternFill("solid", fgColor=Color(theme=theme, tint=tint))


def style_summary_sections(ws, labels):
    category_fills = {
        "basic": PatternFill("solid", fgColor="F4F7FA"),
        "return": theme_fill(3, 0.7999816888943144),
        "profit": theme_fill(6, 0.7999816888943144),
        "non_operating": theme_fill(7, 0.7999816888943144),
        "valuation": theme_fill(9, 0.7999816888943144),
        "conclusion": theme_fill(0),
    }
    label_names = set(labels.keys())
    category_rows = {
        "basic": {label for label in label_names if label in {"指数代码", "行情截至日期"}},
        "return": {label for label in label_names if "涨幅" in label},
        "profit": {label for label in label_names if "滚动盈利" in label or "利润增速" in label},
        "non_operating": {
            label for label in label_names
            if "非实业" in label or label in {"假设非实业利润增速", "加权后预期利润增速"}
        },
        "valuation": {label for label in label_names if "PE" in label or "估值" in label},
        "conclusion": {label for label in label_names if label in {"指数整体泡沫评估", "建议配置"}},
    }
    section_top = Side(style="medium", color="9FB3C8")
    thin = Side(style="thin", color="D9E1F2")
    dark = Side(style="medium", color="7B8794")
    section_starts = set()
    for label in label_names:
        if "年至今涨幅" in label and "年初" not in label:
            section_starts.add(label)
        if "Q1利润增速" in label:
            section_starts.add(label)
        if label.startswith("非实业权重："):
            section_starts.add(label)
        if "PE_TTM" in label:
            section_starts.add(label)
        if label == "指数整体泡沫评估":
            section_starts.add(label)
    key_rows = {
        "加权后预期利润增速",
        "指数整体 PE/G",
        "指数整体泡沫评估",
        "建议配置",
    }
    key_rows.update({label for label in label_names if "利润增速预期：截止" in label and "变动" not in label})
    key_rows.update({label for label in label_names if "PE 十年分位" in label})
    for category, row_labels in category_rows.items():
        fill = category_fills[category]
        for label in row_labels:
            row_idx = labels.get(label)
            if not row_idx:
                continue
            for col_idx in range(1, 7):
                cell = ws.cell(row_idx, col_idx)
                cell.fill = fill
                cell.font = Font(name="PingFang SC", size=10, color="1F2937")
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for label in section_starts:
        row_idx = labels.get(label)
        if not row_idx:
            continue
        for col_idx in range(1, 7):
            cell = ws.cell(row_idx, col_idx)
            cell.border = Border(left=thin, right=thin, top=section_top, bottom=thin)
            cell.font = Font(name="PingFang SC", size=10, bold=True, color="1F2937")
    for label in key_rows:
        row_idx = labels.get(label)
        if not row_idx:
            continue
        for col_idx in range(1, 7):
            cell = ws.cell(row_idx, col_idx)
            cell.font = Font(name="PingFang SC", size=10, bold=True, color="111827")
            if label in {"指数整体泡沫评估", "建议配置"}:
                cell.border = Border(left=thin, right=thin, top=dark, bottom=dark)
    for col_idx in range(2, 6):
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row_idx, col_idx).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_idx in range(2, ws.max_row + 1):
        ws.cell(row_idx, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(row_idx, 6).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(row_idx, 6).font = Font(name="PingFang SC", size=9, color="4B5563")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 15
    for col_idx in range(3, 6):
        ws.column_dimensions[get_column_letter(col_idx)].width = 13
    ws.column_dimensions["F"].width = 62
    ws.row_dimensions[1].height = 28
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 24
    if labels.get("指数整体泡沫评估"):
        ws.row_dimensions[labels["指数整体泡沫评估"]].height = 30
    if labels.get("建议配置"):
        ws.row_dimensions[labels["建议配置"]].height = 28
    ws.sheet_properties.tabColor = "1F4E78"


def add_bar_chart(ws, source_ws, title, rows, anchor, y_title):
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = title
    chart.y_axis.title = y_title
    chart.x_axis.title = "指数"
    data = Reference(source_ws, min_col=1, max_col=5, min_row=rows[0], max_row=rows[-1])
    chart.add_data(data, titles_from_data=True, from_rows=True)
    chart.set_categories(Reference(source_ws, min_col=2, max_col=5, min_row=1, max_row=1))
    chart.height = 7.2
    chart.width = 13.5
    ws.add_chart(chart, anchor)


def find_label(labels, predicate):
    matches = [label for label in labels if predicate(label)]
    if len(matches) != 1:
        raise KeyError(f"Expected exactly one label, got {matches}")
    return matches[0]


def build(snapshot):
    rows = build_rows(snapshot)
    wb = Workbook()
    summary = wb.active
    summary.title = "配置比较"

    summary.append(["指标", *INDEXES, "数据来源与说明"])
    for row in rows:
        summary.append(list(row))

    labels = {summary.cell(row, 1).value: row for row in range(2, summary.max_row + 1)}
    expected_previous_label = find_label(labels, lambda label: "利润增速预期：截止" in label and month_label(snapshot["previous_end"]) in label)
    expected_latest_label = find_label(labels, lambda label: "利润增速预期：截止" in label and month_label(snapshot["target_end"]) in label)
    expected_change_label = find_label(labels, lambda label: "利润增速预期变动" in label)
    expected_previous = labels[expected_previous_label]
    expected_latest = labels[expected_latest_label]
    expected_change = labels[expected_change_label]
    for col in range(2, 6):
        letter = get_column_letter(col)
        summary.cell(
            expected_change,
            col,
            f'=IF(OR({letter}{expected_latest}="",{letter}{expected_previous}=""),"",{letter}{expected_latest}-{letter}{expected_previous})',
        )

    for col in range(2, 4):
        letter = get_column_letter(col)
        non_operating_weight_label = find_label(labels, lambda label: label.startswith("非实业权重：截止") and month_label(snapshot["target_end"]) in label)
        non_operating_weight = labels[non_operating_weight_label]
        non_operating_profit_share = labels["非实业利润占比"]
        non_operating_profit_growth = labels["假设非实业利润增速"]
        operating_growth = labels["剔除非实业后利润增速"]
        fundamental_return = labels["加权后预期利润增速"]

        summary.cell(
            operating_growth,
            col,
            f'=IFERROR(({letter}{expected_latest}-{letter}{non_operating_profit_share}*{letter}{non_operating_profit_growth})/(1-{letter}{non_operating_profit_share}),"")',
        )
        summary.cell(
            fundamental_return,
            col,
            f'=IFERROR({letter}{operating_growth}*(1-{letter}{non_operating_weight})+{letter}{non_operating_profit_growth}*{letter}{non_operating_weight},"")',
        )
    summary.cell(labels[expected_change_label], 6, f"公式：截止{month_label(snapshot['target_end'])}利润增速预期 - 截止{month_label(snapshot['previous_end'])}数值；缺历史截面则留空")
    summary.cell(labels["剔除非实业后利润增速"], 6, "公式：(整体利润增速预期 - 非实业利润占比 × 非实业利润增速) ÷ (1 - 非实业利润占比)")
    summary.cell(labels["加权后预期利润增速"], 6, "A股公式：剔除非实业后增速 × (1 - 非实业权重) + 非实业利润增速 × 非实业权重")

    pe_ttm_label = find_label(labels, lambda label: "PE_TTM" in label)
    pe_percentile_label = find_label(labels, lambda label: "PE 十年分位" in label)
    pe_ttm = labels[pe_ttm_label]
    ex_operating_pe = labels["剔除非实业后 PE"]
    overall_peg = labels["指数整体 PE/G"]
    ex_operating_peg = labels["剔除非实业 PE/G"]
    bubble_assessment = labels["指数整体泡沫评估"]
    operating_growth = labels["剔除非实业后利润增速"]
    fundamental_return = labels["加权后预期利润增速"]
    pe_percentile = labels[pe_percentile_label]
    for col in range(2, 6):
        letter = get_column_letter(col)
        summary.cell(overall_peg, col, f'=IFERROR({letter}{pe_ttm}/({letter}{fundamental_return}*100),"")')
        summary.cell(ex_operating_peg, col, f'=IFERROR({letter}{ex_operating_pe}/({letter}{operating_growth}*100),"")')
        summary.cell(
            bubble_assessment,
            col,
            (
                f'=IF({letter}{overall_peg}="",'
                f'IF({letter}{pe_percentile}>=0.85,"缺PEG，PE高分位",IF({letter}{pe_percentile}>=0.7,"缺PEG，PE偏高","缺PEG，PE不高")),'
                f'IF(AND({letter}{pe_percentile}>=0.85,{letter}{overall_peg}>=1.5),"泡沫较高",'
                f'IF(AND({letter}{pe_percentile}>=0.85,{letter}{overall_peg}>=1),"偏泡沫",'
                f'IF(AND({letter}{pe_percentile}>=0.85,{letter}{overall_peg}<1),"估值高但盈利可消化",'
                f'IF(AND({letter}{pe_percentile}>=0.7,{letter}{overall_peg}>=1),"偏热",'
                f'IF(AND({letter}{pe_percentile}>=0.7,{letter}{overall_peg}<1),"偏热但盈利可消化","泡沫不明显"))))))'
            ),
        )
    for col in range(4, 6):
        letter = get_column_letter(col)
        expected_latest = labels[expected_latest_label]
        fundamental_return = labels["加权后预期利润增速"]
        summary.cell(fundamental_return, col, f'=IFERROR({letter}{expected_latest},"")')

    style_header(summary, 1)
    apply_grid(summary, 1, summary.max_row, 1, 6)
    apply_readable_table(summary, 1, summary.max_row, 1, 6)
    summary.freeze_panes = "B2"
    summary.auto_filter.ref = f"A1:F{len(rows) + 1}"
    widths = [30, 16, 16, 16, 16, 48]
    for col, width in enumerate(widths, 1):
        summary.column_dimensions[get_column_letter(col)].width = width

    for row in range(2, summary.max_row + 1):
        label = summary.cell(row, 1).value
        if is_percent_label(label):
            for col in range(2, 6):
                summary.cell(row, col).number_format = "0.0%"
    for row in range(2, summary.max_row + 1):
        label = str(summary.cell(row, 1).value)
        if ("PE" in label and "十年分位" not in label) or label == "过去十年估值中位数":
            for col in range(2, 6):
                summary.cell(row, col).number_format = "0.00"
    style_summary_sections(summary, labels)

    visual = wb.create_sheet("可视化")
    visual.sheet_view.showGridLines = False
    visual.sheet_properties.tabColor = "5B7FA3"
    visual["A1"] = "四大指数配置可视化比较"
    visual["A1"].font = Font(size=16, bold=True, color="1F4E78")
    visual["A2"] = f"表格为 Wind 截止{month_label(snapshot['target_end'])}市场数据；利润增速预期中，A股来自 Wind，美股来自脚本顶部手工输入区。非实业 = 金融 + 能源 + 房地产。"
    visual["A2"].alignment = Alignment(wrap_text=True)
    visual.column_dimensions["A"].width = 100

    baseline_return_label = find_label(labels, lambda label: "年至今涨幅" in label and "年初" not in label)
    current_return_label = find_label(labels, lambda label: "年初至今涨幅" in label)
    baseline_earnings_label = find_label(labels, lambda label: "滚动盈利增长" in label)
    ttm_earnings_label = find_label(labels, lambda label: "滚动盈利同比" in label)
    add_bar_chart(visual, summary, "指数涨幅比较", [labels[baseline_return_label], labels[current_return_label]], "A4", "涨幅")
    add_bar_chart(visual, summary, f"截止{month_label(snapshot['target_end'])} PE 与过去十年中位数", [labels[pe_ttm_label], labels["过去十年估值中位数"]], "J4", "PE_TTM")
    add_bar_chart(visual, summary, "滚动盈利增长比较", [labels[baseline_earnings_label], labels[ttm_earnings_label]], "A19", "盈利增长")

    previous_label = month_label(snapshot["previous_end"])
    target_label = month_label(snapshot["target_end"])
    comparison = wb.create_sheet(f"截止{previous_label}与{target_label}比较")
    comparison.append([
        "指数",
        f"截止{previous_label}：{MANUAL_INPUTS['baseline_start_year']}年至今涨幅", f"截止{target_label}：{MANUAL_INPUTS['baseline_start_year']}年至今涨幅", "变动",
        f"截止{previous_label}：PE", f"截止{target_label}：PE_TTM", "PE 变动", "PE 相对变动",
        f"截止{previous_label}：非实业权重", f"截止{target_label}：非实业权重", "变动",
        f"截止{previous_label}：利润增速预期", f"截止{target_label}：利润增速预期", "变动",
    ])
    previous_non_op_label = find_label(labels, lambda label: label.startswith("非实业权重：截止") and previous_label in label)
    latest_non_op_label = find_label(labels, lambda label: label.startswith("非实业权重：截止") and target_label in label)
    rows_by_code = snapshot["rows_by_code"]
    for row_num, (name, code) in enumerate(zip(INDEXES, CODES), start=2):
        summary_col = row_num
        previous_baseline_return = rows_by_code[code].get("previous_baseline_return")
        previous_pe = rows_by_code[code].get("previous_pe")
        previous_non_operating = rows_by_code[code].get("non_op_previous")
        comparison.append([
            name,
            previous_baseline_return, f"='配置比较'!{get_column_letter(summary_col)}{labels[baseline_return_label]}", f"=C{row_num}-B{row_num}",
            previous_pe, f"='配置比较'!{get_column_letter(summary_col)}{labels[pe_ttm_label]}", f"=F{row_num}-E{row_num}", f"=IFERROR(F{row_num}/E{row_num}-1,\"\")",
            previous_non_operating, f"='配置比较'!{get_column_letter(summary_col)}{labels[latest_non_op_label]}", f'=IF(OR(I{row_num}="",J{row_num}=""),"",J{row_num}-I{row_num})',
            f"='配置比较'!{get_column_letter(summary_col)}{labels[expected_previous_label]}", f"='配置比较'!{get_column_letter(summary_col)}{labels[expected_latest_label]}", f'=IF(OR(L{row_num}="",M{row_num}=""),"",M{row_num}-L{row_num})',
        ])
    comparison["A7"] = "说明"
    comparison["B7"] = f"变动列为截止{target_label}减截止{previous_label}。A股与市场数据优先来自 Wind；美股利润预期来自脚本顶部手工输入区。"
    comparison.merge_cells("B7:N7")
    comparison["B7"].alignment = Alignment(wrap_text=True)
    style_header(comparison, 1, 1, 14)
    apply_grid(comparison, 1, comparison.max_row, 1, 14)
    comparison.sheet_view.showGridLines = False
    comparison.sheet_properties.tabColor = "5B7FA3"
    comparison.freeze_panes = "B2"
    comparison.auto_filter.ref = "A1:N5"
    comparison.column_dimensions["A"].width = 16
    for col in range(2, 15):
        comparison.column_dimensions[get_column_letter(col)].width = 17
    comparison.row_dimensions[7].height = 34
    note_fill = PatternFill("solid", fgColor="FFFDF3")
    for row in range(2, 6):
        comparison.cell(row, 1).fill = PatternFill("solid", fgColor="F4F7FA")
        comparison.cell(row, 1).font = Font(bold=True, color="1F2937")
        if row % 2 == 0:
            for col in range(2, 15):
                comparison.cell(row, col).fill = PatternFill("solid", fgColor="FBFCFE")
    for col in range(1, 15):
        comparison.cell(7, col).fill = note_fill
    for row in range(2, 6):
        for col in [2, 3, 4, 8, 9, 10, 11, 12, 13, 14]:
            comparison.cell(row, col).number_format = "0.0%"
        for col in [5, 6, 7]:
            comparison.cell(row, col).number_format = "0.00"
    forecast_chart = BarChart()
    forecast_chart.type = "col"
    forecast_chart.style = 10
    forecast_chart.title = f"截止{previous_label}与截止{target_label}利润增速预期"
    forecast_chart.y_axis.title = "利润增速预期"
    forecast_chart.x_axis.title = "指数"
    forecast_chart.add_data(Reference(comparison, min_col=12, max_col=13, min_row=1, max_row=5), titles_from_data=True)
    forecast_chart.set_categories(Reference(comparison, min_col=1, min_row=2, max_row=5))
    forecast_chart.height = 7.2
    forecast_chart.width = 15
    comparison.add_chart(forecast_chart, "A9")

    assumptions = wb.create_sheet("口径说明")
    assumptions.append(["项目", "说明"])
    notes = [
        ("非实业定义", "金融、能源、房地产三个板块。"),
        ("市场数据", f"指数点位、PE_TTM、过去十年估值分位来自 WindPy；四个指数均使用 {fmt_date(snapshot['target_end'])} 收盘，即月末最后交易日。"),
        (f"{MANUAL_INPUTS['current_year']}年Q1利润增速", "沪深300和中证500使用 Wind 指数级字段 YOYNETPROFIT，按一季报 rptDate 获取报告期净利润同比增速。标普500和纳斯达克100使用脚本顶部手工输入区。"),
        ("利润增速预期", f"沪深300和中证500使用 Wind 指数级字段 west_netprofit_yoy，分别按 tradeDate={yyyymmdd(snapshot['previous_end'])} 和 tradeDate={yyyymmdd(snapshot['target_end'])} 获取 FY1 一致预期净利润增速及其变动。美股使用脚本顶部手工输入区。"),
        ("手工输入项", "美股利润增速预期、非实业收入占比、非实业利润占比、剔除后 PE、假设非实业利润增速、建议配置在脚本顶部 MANUAL_INPUTS 中维护。"),
        ("盈利增长", f"{MANUAL_INPUTS['baseline_start_year']} 年滚动盈利增长和截止{target_label}滚动盈利同比由指数点位 / PE_TTM 反推，用作指数整体盈利变化的近似观察。"),
        ("剔除非实业后利润增速", "A 股使用公式动态计算：(指数整体利润增速 - 非实业利润占比 × 假设非实业利润增速) ÷ (1 - 非实业利润占比)。"),
        (f"截止{target_label}非实业权重", "沪深300和中证500使用 Wind 月末成分权重，按万得一级行业汇总金融、能源、房地产。美股留空。"),
        ("空白单元格", "表示 Wind 本次未能按统一口径取得，且手工输入区未提供，不进行推测填充。"),
    ]
    for item in notes:
        assumptions.append(item)
    style_header(assumptions, 1, 1, 2)
    apply_grid(assumptions, 1, assumptions.max_row, 1, 2)
    apply_readable_table(assumptions, 1, assumptions.max_row, 1, 2)
    assumptions.sheet_properties.tabColor = "9AA9B7"
    assumptions.column_dimensions["A"].width = 20
    assumptions.column_dimensions["B"].width = 115
    for row in range(2, assumptions.max_row + 1):
        assumptions.row_dimensions[row].height = 32

    raw = wb.create_sheet("原始数据")
    raw.append([
        "指数", "代码", "行情日期", "收盘点位", "PE_TTM",
        f"{MANUAL_INPUTS['baseline_start_year']}年至今",
        f"{MANUAL_INPUTS['current_year']}年初至今",
        f"{MANUAL_INPUTS['baseline_start_year']}滚动盈利增长",
        f"截止{target_label}滚动盈利同比",
        f"{MANUAL_INPUTS['current_year']}年Q1利润增速",
        "PE十年分位",
        f"截止{previous_label}利润增速预期",
        f"截止{target_label}利润增速预期",
    ])
    raw_rows = []
    for name, code in zip(INDEXES, CODES):
        row = rows_by_code[code]
        raw_rows.append((
            name,
            code,
            fmt_date(snapshot["target_end"]),
            row.get("close"),
            row.get("pe"),
            row.get("baseline_return"),
            row.get("current_year_return"),
            row.get("baseline_earnings_growth"),
            row.get("ttm_earnings_yoy"),
            row.get("q1_profit_growth"),
            row.get("pe_rank"),
            row.get("expected_previous"),
            row.get("expected_latest"),
        ))
    for row in raw_rows:
        raw.append(row)
    style_header(raw, 1, 1, 13)
    apply_grid(raw, 1, raw.max_row, 1, 13)
    raw.sheet_view.showGridLines = False
    raw.sheet_properties.tabColor = "9AA9B7"
    for row in range(2, raw.max_row + 1):
        raw.cell(row, 1).fill = PatternFill("solid", fgColor="F4F7FA")
        raw.cell(row, 1).font = Font(bold=True, color="1F2937")
        if row % 2 == 0:
            for col in range(2, 14):
                raw.cell(row, col).fill = PatternFill("solid", fgColor="FBFCFE")
    for col in range(1, 14):
        raw.column_dimensions[get_column_letter(col)].width = 18
    for row in range(2, raw.max_row + 1):
        for col in range(6, 14):
            raw.cell(row, col).number_format = "0.0%"

    output = output_path(snapshot["target_end"])
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(output)

    # Reload once to catch malformed formulas, relationships, or chart serialization.
    checked = load_workbook(output, data_only=False)
    assert checked.sheetnames == ["配置比较", "可视化", f"截止{previous_label}与{target_label}比较", "口径说明", "原始数据"]
    assert checked["配置比较"]["B2"].value == "000300.SH"
    print(output)
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="每月更新四大指数配置比较表。默认更新到上个月最后一个交易日。")
    parser.add_argument("--target-date", help="指定行情截至日期，格式 YYYY-MM-DD；不填则自动使用上个月最后一个交易日。")
    args = parser.parse_args()
    wind = start_wind()
    try:
        target_end, previous_end = resolve_target_dates(wind, args.target_date)
        snapshot = fetch_market_snapshot(wind, target_end, previous_end)
    finally:
        wind.close()
    build(snapshot)
