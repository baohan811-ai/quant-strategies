import os
import json
import sqlite3
import time
import urllib.request
import pandas as pd
import numpy as np
from datetime import datetime, timedelta
from openpyxl.styles import Font
from openpyxl.drawing.image import Image as OpenpyxlImage

from 维护工具.local_market_db import MARKET_DB_PATH, init_market_db, load_price_matrix

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(SCRIPT_DIR)
CACHE_DIR = os.path.join(BASE_DIR, "缓存")
os.makedirs(CACHE_DIR, exist_ok=True)

# =========================
# 0. 参数
# =========================
MAX_HOLDINGS = 20
INITIAL_WEIGHT = 1 / MAX_HOLDINGS
TRANSACTION_COST_RATE = 0.0025
STOP_DRAWDOWN = 0.2
PEAK_RETRACE_SELL_DRAWDOWN = 0.1
LOW_EFFICIENCY_MIN_HOLDING_DAYS = 60
LOW_EFFICIENCY_MAX_PROFIT_THRESHOLD = 0.05
LOW_EFFICIENCY_CURRENT_PROFIT_THRESHOLD = 0.00

# 评分系统
# total_score =
#   MA5_MA60_GAP_WEIGHT * (MA5 / MA60 - 1)
#   + MA60_5D_TREND_WEIGHT * (MA60 / MA60.shift(5) - 1)
#   + VOLUME_RATIO_SCORE_WEIGHT * clip(成交量 / 20日均量 - 1, 0, VOLUME_RATIO_MAX_BONUS_BASE)
MA5_MA60_GAP_WEIGHT = 0.8
MA60_5D_TREND_WEIGHT = 0.8
VOLUME_RATIO_SCORE_WEIGHT = 0.5
VOLUME_RATIO_LOOKBACK = 20
VOLUME_RATIO_MAX_BONUS_BASE = 1.0

# 金叉触发日最大涨幅过滤。
# 已用本地中证800缓存行情做网格回测：3%~15%中，5%在收益、回撤、夏普之间最优；
# 简单放宽到5.5%~7%会增加追高交易，组合表现反而下降。
SIGNAL_MAX_DAILY_RETURN = 0.05

LOOKBACK_DAYS = 1600
TRADE_START_DATE = "2023-04-03"  # 必须写成字符串，例如 "2025-01-01"；None 表示沿用当前逻辑
CACHE_PREFIX = "中证800"
TUSHARE_TOKEN = os.environ.get("TUSHARE_TOKEN", "")
TUSHARE_INDEX_CODE = "000906.SH"
TUSHARE_API_URL = "http://api.tushare.pro"
TUSHARE_SLEEP_SECONDS = 0.25
CONSTITUENT_CACHE_PATH = os.path.join(CACHE_DIR, "中证800_tushare_constituents.pkl")
TUSHARE_DAILY_RANGE_CACHE = {}
TUSHARE_REALTIME_CACHE = {}
TUSHARE_REALTIME_BATCH_SIZE = 400
TUSHARE_INDEX_DAILY_CACHE = {}
USE_HISTORICAL_CONSTITUENTS = True


def require_tushare_token():
    if not TUSHARE_TOKEN:
        raise RuntimeError("未设置环境变量 TUSHARE_TOKEN")


def tushare_query(api_name, params=None, fields=None, max_retry=3):
    require_tushare_token()
    payload = {
        "api_name": api_name,
        "token": TUSHARE_TOKEN,
        "params": params or {},
        "fields": fields or "",
    }
    data = json.dumps(payload).encode("utf-8")
    last_error = None
    for attempt in range(1, max_retry + 1):
        try:
            req = urllib.request.Request(
                TUSHARE_API_URL,
                data=data,
                headers={"Content-Type": "application/json"},
                method="POST",
            )
            with urllib.request.urlopen(req, timeout=30) as response:
                result = json.loads(response.read().decode("utf-8"))
            if result.get("code") != 0:
                raise RuntimeError(result.get("msg", result))
            payload_data = result.get("data") or {}
            return pd.DataFrame(
                payload_data.get("items", []),
                columns=payload_data.get("fields", []),
            )
        except Exception as exc:
            last_error = exc
            if attempt < max_retry:
                time.sleep(TUSHARE_SLEEP_SECONDS * attempt)
                continue
            raise RuntimeError(f"Tushare 接口 {api_name} 调用失败：{last_error}") from exc


def tushare_code_to_wind(ts_code):
    return str(ts_code).strip()


def wind_code_to_tushare(wind_code):
    return str(wind_code).strip()


def get_tushare_trade_dates(start_date, end_date):
    cal = tushare_query(
        "trade_cal",
        {
            "exchange": "SSE",
            "start_date": pd.Timestamp(start_date).strftime("%Y%m%d"),
            "end_date": pd.Timestamp(end_date).strftime("%Y%m%d"),
            "is_open": "1",
        },
        "cal_date",
    )
    if cal.empty:
        raise RuntimeError(f"Tushare 交易日历为空：{start_date} ~ {end_date}")
    return sorted(pd.to_datetime(cal["cal_date"]).dt.strftime("%Y-%m-%d").tolist())


def get_latest_tushare_trading_dates(end_date):
    query_end = pd.Timestamp(end_date)
    query_start = query_end - pd.Timedelta(days=30)
    return get_tushare_trade_dates(query_start, query_end)


def load_stock_basic_name_map():
    frames = []
    for list_status in ["L", "D", "P"]:
        df = tushare_query(
            "stock_basic",
            {"exchange": "", "list_status": list_status},
            "ts_code,name",
        )
        if not df.empty:
            frames.append(df)
        time.sleep(TUSHARE_SLEEP_SECONDS)
    if not frames:
        return {}
    stock_basic = pd.concat(frames, ignore_index=True).drop_duplicates("ts_code")
    return dict(zip(stock_basic["ts_code"], stock_basic["name"]))


def get_index_constituents_from_tushare(index_code, end_date):
    query_end = pd.Timestamp(end_date)
    query_start = query_end - pd.Timedelta(days=120)
    weights = tushare_query(
        "index_weight",
        {
            "index_code": index_code,
            "start_date": query_start.strftime("%Y%m%d"),
            "end_date": query_end.strftime("%Y%m%d"),
        },
        "index_code,con_code,trade_date,weight",
    )
    if weights.empty:
        if os.path.exists(CONSTITUENT_CACHE_PATH):
            print(f"Tushare 中证800成分为空，沿用本地缓存：{CONSTITUENT_CACHE_PATH}")
            cached = pd.read_pickle(CONSTITUENT_CACHE_PATH)
            return cached["wind_code"].tolist(), cached["sec_name"].tolist()
        raise RuntimeError("Tushare 未返回中证800成分股，且没有本地成分缓存可用。")

    latest_weight_date = weights["trade_date"].max()
    latest = weights.loc[weights["trade_date"] == latest_weight_date].copy()
    latest = latest.sort_values("con_code").drop_duplicates("con_code")
    name_map = load_stock_basic_name_map()
    latest["wind_code"] = latest["con_code"].map(tushare_code_to_wind)
    latest["sec_name"] = latest["con_code"].map(name_map).fillna(latest["con_code"])
    cache_df = latest[["wind_code", "sec_name"]].reset_index(drop=True)
    cache_df.to_pickle(CONSTITUENT_CACHE_PATH)
    print(f"Tushare 中证800成分日期：{pd.Timestamp(latest_weight_date).strftime('%Y-%m-%d')}")
    return cache_df["wind_code"].tolist(), cache_df["sec_name"].tolist()


def fetch_tushare_index_close(index_code, start_date, end_date):
    cache_key = (
        index_code,
        pd.Timestamp(start_date).strftime("%Y-%m-%d"),
        pd.Timestamp(end_date).strftime("%Y-%m-%d"),
    )
    if cache_key in TUSHARE_INDEX_DAILY_CACHE:
        return TUSHARE_INDEX_DAILY_CACHE[cache_key].copy()

    index_daily = tushare_query(
        "index_daily",
        {
            "ts_code": index_code,
            "start_date": pd.Timestamp(start_date).strftime("%Y%m%d"),
            "end_date": pd.Timestamp(end_date).strftime("%Y%m%d"),
        },
        "ts_code,trade_date,close",
    )
    if index_daily.empty:
        result = pd.Series(dtype=float, name=f"{index_code}收盘价")
        TUSHARE_INDEX_DAILY_CACHE[cache_key] = result
        return result.copy()

    index_daily["trade_date"] = pd.to_datetime(index_daily["trade_date"])
    index_daily["close"] = pd.to_numeric(index_daily["close"], errors="coerce")
    close = (
        index_daily
        .dropna(subset=["close"])
        .drop_duplicates("trade_date", keep="last")
        .set_index("trade_date")["close"]
        .sort_index()
    )
    close.name = f"{index_code}收盘价"
    TUSHARE_INDEX_DAILY_CACHE[cache_key] = close
    return close.copy()

# =========================
# 2. 中证800成分股
# =========================
stock_codes, stock_names = get_index_constituents_from_tushare(TUSHARE_INDEX_CODE, datetime.today())
code_to_name = dict(zip(stock_codes, stock_names))

print("中证800股票数量：", len(stock_codes))


def load_historical_constituent_snapshots(universe_name, query_start_date, query_end_date):
    with sqlite3.connect(MARKET_DB_PATH) as conn:
        prior_snapshot = pd.read_sql_query(
            """
            SELECT snapshot_date, wind_code, sec_name
            FROM universe_constituents_snapshot
            WHERE universe_name = ?
              AND snapshot_date = (
                  SELECT MAX(snapshot_date)
                  FROM universe_constituents_snapshot
                  WHERE universe_name = ?
                    AND snapshot_date < ?
              )
            """,
            conn,
            params=[universe_name, universe_name, query_start_date],
        )
        range_snapshots = pd.read_sql_query(
            """
            SELECT snapshot_date, wind_code, sec_name
            FROM universe_constituents_snapshot
            WHERE universe_name = ?
              AND snapshot_date >= ?
              AND snapshot_date <= ?
            ORDER BY snapshot_date, wind_code
            """,
            conn,
            params=[universe_name, query_start_date, query_end_date],
        )

    snapshots = pd.concat([prior_snapshot, range_snapshots], ignore_index=True)
    if snapshots.empty:
        return snapshots

    snapshots["snapshot_date"] = pd.to_datetime(snapshots["snapshot_date"])
    snapshots = snapshots.drop_duplicates(["snapshot_date", "wind_code"], keep="last")
    snapshots = snapshots.sort_values(["snapshot_date", "wind_code"]).reset_index(drop=True)
    return snapshots


def build_daily_universe_member_matrix(snapshots, trade_dates, codes):
    member = pd.DataFrame(False, index=trade_dates, columns=codes)
    if snapshots.empty or len(trade_dates) == 0:
        return member

    snapshots = snapshots.sort_values("snapshot_date")
    snapshot_dates = snapshots["snapshot_date"].drop_duplicates().tolist()
    for i, snapshot_date in enumerate(snapshot_dates):
        next_snapshot_date = snapshot_dates[i + 1] if i + 1 < len(snapshot_dates) else None
        if next_snapshot_date is None:
            active_dates = trade_dates[trade_dates >= snapshot_date]
        else:
            active_dates = trade_dates[(trade_dates >= snapshot_date) & (trade_dates < next_snapshot_date)]
        if len(active_dates) == 0:
            continue

        snapshot_codes = snapshots.loc[snapshots["snapshot_date"] == snapshot_date, "wind_code"].tolist()
        active_codes = [code for code in snapshot_codes if code in member.columns]
        if active_codes:
            member.loc[active_dates, active_codes] = True

    return member


def get_limit_up_ratio(code, stock_name):
    stock_name = str(stock_name).upper()
    if "ST" in stock_name:
        return 0.05
    if code.endswith(".BJ"):
        return 0.30
    if code.startswith("300") or code.startswith("688"):
        return 0.20
    return 0.10


def get_limit_down_ratio(code, stock_name):
    return get_limit_up_ratio(code, stock_name)


def sanitize_score_component(df):
    return df.replace([np.inf, -np.inf], np.nan)


def calculate_golden_cross_score_details(ma5_df, ma60_df, volume_df):
    ma5_ma60_gap_score = MA5_MA60_GAP_WEIGHT * (ma5_df / ma60_df - 1)
    ma60_5d_trend_score = MA60_5D_TREND_WEIGHT * (ma60_df / ma60_df.shift(5) - 1)
    volume_ma = volume_df.rolling(VOLUME_RATIO_LOOKBACK).mean()
    volume_ratio = volume_df / volume_ma.replace(0, np.nan)
    volume_ratio_score = VOLUME_RATIO_SCORE_WEIGHT * (
        (volume_ratio - 1).clip(lower=0, upper=VOLUME_RATIO_MAX_BONUS_BASE)
    )
    total_score = ma5_ma60_gap_score + ma60_5d_trend_score + volume_ratio_score
    return {
        "ma5_ma60_gap_score": sanitize_score_component(ma5_ma60_gap_score),
        "ma60_5d_trend_score": sanitize_score_component(ma60_5d_trend_score),
        "volume_ratio": sanitize_score_component(volume_ratio),
        "volume_ratio_score": sanitize_score_component(volume_ratio_score),
        "total_score": sanitize_score_component(total_score),
    }


def evaluate_intraday_sell_signal(holding_info, open_price, low_price, high_price):
    entry_price = holding_info["entry_price"]
    prev_peak_price = holding_info["peak_price"]
    max_profit = prev_peak_price / entry_price - 1 if entry_price > 0 else -np.inf
    sell_price = np.nan
    is_breakeven_exit = False
    sell_reason_text = ""
    monitor_level = np.nan

    stop_loss_price = entry_price * (1 - STOP_DRAWDOWN)
    if low_price <= stop_loss_price:
        sell_price = open_price if open_price < stop_loss_price else stop_loss_price
        sell_reason_text = "固定止损触发"
        monitor_level = stop_loss_price
    elif prev_peak_price > 0:
        retrace_price = prev_peak_price * (1 - PEAK_RETRACE_SELL_DRAWDOWN)
        monitor_level = retrace_price
        if low_price <= retrace_price:
            sell_price = open_price if open_price < retrace_price else retrace_price
            sell_reason_text = "前高回撤10%卖出"

    return {
        "sell_price": sell_price,
        "is_breakeven_exit": is_breakeven_exit,
        "sell_reason_text": sell_reason_text,
        "monitor_level": monitor_level,
        "max_profit": max_profit,
        "prev_peak_price": prev_peak_price,
    }

# =========================
# 3. 时间区间
# =========================
def parse_date(date_str):
    if isinstance(date_str, pd.Timestamp):
        return date_str.date()
    if isinstance(date_str, datetime):
        return date_str.date()
    if hasattr(date_str, "year") and hasattr(date_str, "month") and hasattr(date_str, "day") and not isinstance(date_str, str):
        return date_str
    if not isinstance(date_str, str):
        raise TypeError('TRADE_START_DATE 必须是 "YYYY-MM-DD" 格式的字符串，例如 "2025-01-01"。')
    return datetime.strptime(date_str, "%Y-%m-%d").date()


end_dt = datetime.today()
end_date = end_dt.strftime("%Y-%m-%d")

if TRADE_START_DATE is not None:
    trade_start_dt = parse_date(TRADE_START_DATE)
    trade_start_ts = pd.Timestamp(trade_start_dt)
    data_start_dt = trade_start_dt - timedelta(days=LOOKBACK_DAYS)
else:
    trade_start_dt = None
    trade_start_ts = None
    data_start_dt = end_dt.date() - timedelta(days=LOOKBACK_DAYS)

start_date = data_start_dt.strftime("%Y-%m-%d")

print(f"数据开始日期: {start_date}")
print(f"数据结束日期: {end_date}")
if TRADE_START_DATE is not None:
    print(f"开始建仓日期: {TRADE_START_DATE}")
    if trade_start_dt > end_dt.date():
        raise ValueError("TRADE_START_DATE 不能晚于 end_date，请检查参数设置。")

# =========================
# 4. 行情获取
# =========================
def get_latest_complete_price_date(db_path=MARKET_DB_PATH, adjusted="F"):
    if not os.path.exists(db_path):
        return None
    with sqlite3.connect(db_path) as conn:
        row = conn.execute("""
            SELECT MAX(trade_date)
            FROM daily_prices
            WHERE adjusted = ?
              AND open IS NOT NULL
              AND high IS NOT NULL
              AND low IS NOT NULL
              AND close IS NOT NULL
              AND volume IS NOT NULL
              AND amt IS NOT NULL
        """, (adjusted,)).fetchone()
    return row[0] if row and row[0] else None


def upsert_tushare_daily_prices(conn, df, adjusted="F"):
    if df.empty:
        return 0
    updated_at = datetime.now().isoformat(timespec="seconds")
    rows = [
        (
            row.trade_date,
            row.wind_code,
            row.open,
            row.high,
            row.low,
            row.close,
            row.volume,
            row.amt,
            adjusted,
            updated_at,
        )
        for row in df.itertuples(index=False)
    ]
    conn.executemany("""
        INSERT INTO daily_prices (
            trade_date, wind_code, open, high, low, close, volume, amt, adjusted, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(trade_date, wind_code, adjusted) DO UPDATE SET
            open = excluded.open,
            high = excluded.high,
            low = excluded.low,
            close = excluded.close,
            volume = excluded.volume,
            amt = excluded.amt,
            updated_at = excluded.updated_at
    """, rows)
    conn.commit()
    return len(rows)


def fetch_tushare_daily_range(start_date, end_date):
    cache_key = (
        pd.Timestamp(start_date).strftime("%Y-%m-%d"),
        pd.Timestamp(end_date).strftime("%Y-%m-%d"),
    )
    if cache_key in TUSHARE_DAILY_RANGE_CACHE:
        return TUSHARE_DAILY_RANGE_CACHE[cache_key].copy()

    trade_dates = get_tushare_trade_dates(start_date, end_date)
    frames = []
    for trade_date in trade_dates:
        trade_date_compact = pd.Timestamp(trade_date).strftime("%Y%m%d")
        print(f"拉取 Tushare daily: {trade_date}")
        daily = tushare_query(
            "daily",
            {"trade_date": trade_date_compact},
            "ts_code,trade_date,open,high,low,close,vol,amount",
        )
        if daily.empty:
            print(f"  Tushare daily 为空，跳过：{trade_date}")
            continue
        adj = tushare_query(
            "adj_factor",
            {"trade_date": trade_date_compact},
            "ts_code,trade_date,adj_factor",
        )
        if not adj.empty:
            daily = daily.merge(adj, on=["ts_code", "trade_date"], how="left")
        else:
            daily["adj_factor"] = np.nan
        frames.append(daily)
        time.sleep(TUSHARE_SLEEP_SECONDS)

    if not frames:
        empty = pd.DataFrame()
        TUSHARE_DAILY_RANGE_CACHE[cache_key] = empty
        return empty

    raw = pd.concat(frames, ignore_index=True)
    raw["trade_date"] = pd.to_datetime(raw["trade_date"]).dt.strftime("%Y-%m-%d")
    raw["wind_code"] = raw["ts_code"].map(tushare_code_to_wind)
    for field in ["open", "high", "low", "close", "vol", "amount", "adj_factor"]:
        raw[field] = pd.to_numeric(raw[field], errors="coerce")
    raw = raw.rename(columns={"vol": "volume", "amount": "amt"})
    result = raw[["trade_date", "wind_code", "open", "high", "low", "close", "volume", "amt", "adj_factor"]]
    TUSHARE_DAILY_RANGE_CACHE[cache_key] = result
    return result.copy()


def align_recent_prices_to_local_anchor(raw_df, anchor_date, codes):
    if raw_df.empty:
        return raw_df

    code_set = set(codes)
    raw_df = raw_df.loc[raw_df["wind_code"].isin(code_set)].copy()
    if raw_df.empty:
        return raw_df

    with sqlite3.connect(MARKET_DB_PATH) as conn:
        anchor = pd.read_sql_query("""
            SELECT wind_code, close AS anchor_close
            FROM daily_prices
            WHERE adjusted = 'F'
              AND trade_date = ?
              AND close IS NOT NULL
        """, conn, params=[anchor_date])

    raw_df = raw_df.merge(anchor, on="wind_code", how="left")
    raw_anchor = raw_df.loc[raw_df["trade_date"] == anchor_date, ["wind_code", "close", "adj_factor"]]
    raw_anchor = raw_anchor.rename(columns={"close": "raw_anchor_close", "adj_factor": "anchor_adj_factor"})
    raw_df = raw_df.merge(raw_anchor, on="wind_code", how="left")

    denominator = raw_df["raw_anchor_close"] * raw_df["anchor_adj_factor"]
    multiplier = raw_df["anchor_close"] / denominator.replace(0, np.nan)
    multiplier = multiplier.where(multiplier.notna(), 1.0)
    adjusted_multiplier = multiplier * raw_df["adj_factor"].where(raw_df["adj_factor"].notna(), 1.0)
    for field in ["open", "high", "low", "close"]:
        raw_df[field] = raw_df[field] * adjusted_multiplier

    result = raw_df.loc[raw_df["trade_date"] > anchor_date].copy()
    return result[["trade_date", "wind_code", "open", "high", "low", "close", "volume", "amt"]]


def fetch_tushare_realtime_snapshot(codes, trade_date, src="sina"):
    cache_key = (pd.Timestamp(trade_date).strftime("%Y-%m-%d"), src, tuple(codes))
    if cache_key in TUSHARE_REALTIME_CACHE:
        return TUSHARE_REALTIME_CACHE[cache_key].copy()

    try:
        import tushare as ts
    except ImportError as exc:
        print("未安装 tushare 包，跳过盘中实时行情补齐。可执行：python3 -m pip install tushare")
        TUSHARE_REALTIME_CACHE[cache_key] = pd.DataFrame()
        return pd.DataFrame()

    require_tushare_token()
    ts.set_token(TUSHARE_TOKEN)
    frames = []
    target_date_compact = pd.Timestamp(trade_date).strftime("%Y%m%d")
    for i in range(0, len(codes), TUSHARE_REALTIME_BATCH_SIZE):
        batch = codes[i:i + TUSHARE_REALTIME_BATCH_SIZE]
        print(f"拉取 Tushare realtime_quote: {i}-{i + len(batch)} [{trade_date}]")
        try:
            batch_df = ts.realtime_quote(ts_code=",".join(batch), src=src)
        except Exception as exc:
            print(f"  实时行情批次失败：{type(exc).__name__}: {exc}")
            continue
        if batch_df is None or batch_df.empty:
            continue
        frames.append(batch_df)
        time.sleep(TUSHARE_SLEEP_SECONDS)

    if not frames:
        TUSHARE_REALTIME_CACHE[cache_key] = pd.DataFrame()
        return pd.DataFrame()

    raw = pd.concat(frames, ignore_index=True)
    raw.columns = [str(col).upper() for col in raw.columns]
    required = {"TS_CODE", "DATE", "OPEN", "PRE_CLOSE", "PRICE", "HIGH", "LOW", "VOLUME", "AMOUNT"}
    if not required.issubset(set(raw.columns)):
        print(f"Tushare realtime_quote 字段不完整，跳过：{sorted(raw.columns)}")
        TUSHARE_REALTIME_CACHE[cache_key] = pd.DataFrame()
        return pd.DataFrame()

    raw = raw.loc[raw["DATE"].astype(str) == target_date_compact].copy()
    if raw.empty:
        print(f"Tushare realtime_quote 未返回 {trade_date} 的行情。")
        TUSHARE_REALTIME_CACHE[cache_key] = pd.DataFrame()
        return pd.DataFrame()

    raw["wind_code"] = raw["TS_CODE"].map(tushare_code_to_wind)
    for field in ["OPEN", "PRE_CLOSE", "PRICE", "HIGH", "LOW", "VOLUME", "AMOUNT"]:
        raw[field] = pd.to_numeric(raw[field], errors="coerce")
    raw = raw.loc[raw["PRICE"].notna() & raw["PRICE"].gt(0)].copy()
    if raw.empty:
        TUSHARE_REALTIME_CACHE[cache_key] = pd.DataFrame()
        return pd.DataFrame()

    latest_complete = get_latest_complete_price_date()
    if latest_complete is None:
        TUSHARE_REALTIME_CACHE[cache_key] = pd.DataFrame()
        return pd.DataFrame()
    with sqlite3.connect(MARKET_DB_PATH) as conn:
        anchor = pd.read_sql_query("""
            SELECT wind_code, close AS anchor_close
            FROM daily_prices
            WHERE adjusted = 'F'
              AND trade_date = ?
              AND close IS NOT NULL
        """, conn, params=[latest_complete])
    raw = raw.merge(anchor, on="wind_code", how="left")

    scale = raw["anchor_close"] / raw["PRE_CLOSE"].replace(0, np.nan)
    scale = scale.where(scale.notna(), 1.0)
    realtime = pd.DataFrame({
        "trade_date": pd.Timestamp(trade_date).strftime("%Y-%m-%d"),
        "wind_code": raw["wind_code"],
        "open": raw["OPEN"] * scale,
        "high": raw["HIGH"] * scale,
        "low": raw["LOW"] * scale,
        "close": raw["PRICE"] * scale,
        "volume": raw["VOLUME"],
        "amt": raw["AMOUNT"],
    })
    realtime = realtime.drop_duplicates("wind_code", keep="last")
    TUSHARE_REALTIME_CACHE[cache_key] = realtime
    return realtime.copy()


def ensure_market_data_updated_from_tushare(end_date, codes):
    init_market_db(MARKET_DB_PATH)
    recent_trading_dates = get_latest_tushare_trading_dates(end_date)
    latest_trading_date = recent_trading_dates[-1]
    latest_complete_target_date = (
        recent_trading_dates[-2]
        if len(recent_trading_dates) >= 2
        else latest_trading_date
    )
    latest_price_date = get_latest_complete_price_date()
    if latest_price_date is not None and latest_price_date >= latest_complete_target_date:
        print(f"本地行情数据库已是最新：{latest_price_date}")
        return latest_trading_date
    if latest_price_date is None:
        raise RuntimeError("本地行情数据库没有可锚定的历史行情，不能只用 Tushare 增量补最近几日。")

    update_start = latest_price_date
    print(
        "用 Tushare 补齐本地行情数据库："
        f"锚点={latest_price_date}，目标={latest_complete_target_date}"
    )
    raw = fetch_tushare_daily_range(update_start, latest_complete_target_date)
    aligned = align_recent_prices_to_local_anchor(raw, latest_price_date, codes)
    if aligned.empty:
        print("Tushare 未返回可写入的新增行情。")
        return latest_trading_date
    with sqlite3.connect(MARKET_DB_PATH) as conn:
        inserted = upsert_tushare_daily_prices(conn, aligned)
    print(f"Tushare 新增/更新行情行：{inserted}")
    return latest_trading_date


end_date = ensure_market_data_updated_from_tushare(end_date, stock_codes)
end_dt = pd.Timestamp(end_date).to_pydatetime()

historical_constituent_snapshots = pd.DataFrame()
if USE_HISTORICAL_CONSTITUENTS:
    constituent_start_date = TRADE_START_DATE if TRADE_START_DATE is not None else start_date
    historical_constituent_snapshots = load_historical_constituent_snapshots(
        CACHE_PREFIX,
        constituent_start_date,
        end_date,
    )
    if historical_constituent_snapshots.empty:
        print("未找到中证800历史成分快照，回退为当前中证800成分。")
    else:
        snapshot_codes = historical_constituent_snapshots["wind_code"].drop_duplicates().tolist()
        snapshot_names = (
            historical_constituent_snapshots
            .dropna(subset=["sec_name"])
            .drop_duplicates("wind_code", keep="last")
            .set_index("wind_code")["sec_name"]
            .to_dict()
        )
        stock_codes = snapshot_codes
        code_to_name.update(snapshot_names)
        first_snapshot_date = historical_constituent_snapshots["snapshot_date"].min().strftime("%Y-%m-%d")
        last_snapshot_date = historical_constituent_snapshots["snapshot_date"].max().strftime("%Y-%m-%d")
        print(
            f"使用中证800历史月频成分：{first_snapshot_date} ~ {last_snapshot_date}，"
            f"历史并集股票数量：{len(stock_codes)}"
        )


def get_cache_path(field):
    return os.path.join(CACHE_DIR, f"{CACHE_PREFIX}_{field}_PriceAdjF.pkl")


def sanitize_price_df(df):
    if df.empty:
        return df

    df = df.apply(pd.to_numeric, errors="coerce")
    # 价格字段中 0 和负值都视为无效，避免 pct_change / 比值计算报错。
    df = df.where(df > 0, np.nan)
    df.index = pd.to_datetime(df.index)
    df = df.sort_index()
    df = df.loc[:, ~df.columns.duplicated()]
    return df


def load_cached_df(field):
    cache_path = get_cache_path(field)
    if not os.path.exists(cache_path):
        return pd.DataFrame()

    df = pd.read_pickle(cache_path)
    if df.empty:
        return df

    return sanitize_price_df(df)


def save_cached_df(field, df):
    cache_path = get_cache_path(field)
    df = sanitize_price_df(df)
    df.to_pickle(cache_path)


def has_recent_all_nan_rows(df, recent_days=3):
    if df.empty:
        return False
    recent_df = df.tail(recent_days)
    return (recent_df.notna().sum(axis=1) == 0).any()


def drop_all_nan_rows(df):
    if df.empty:
        return df
    return sanitize_price_df(df.loc[df.notna().sum(axis=1) > 0])


def last_row_all_nan(df):
    if df.empty:
        return False
    return df.iloc[-1].notna().sum() == 0


def supplement_with_tushare_daily(df, codes, field, trade_date):
    if field not in {"open", "low", "high", "close", "volume", "amt"}:
        return df
    if not df.empty and df.index[-1].strftime("%Y-%m-%d") >= trade_date:
        return df

    latest_complete = get_latest_complete_price_date()
    if latest_complete is None:
        return df
    raw = fetch_tushare_daily_range(latest_complete, trade_date)
    aligned = align_recent_prices_to_local_anchor(raw, latest_complete, codes)
    if aligned.empty or aligned["trade_date"].max() < trade_date:
        realtime = fetch_tushare_realtime_snapshot(codes, trade_date)
        if not realtime.empty:
            aligned = pd.concat([aligned, realtime], ignore_index=True)
            print(f"{field} 使用 Tushare 盘中实时行情临时补齐：{trade_date}")
    if aligned.empty:
        return df

    matrix = aligned.pivot(index="trade_date", columns="wind_code", values=field)
    matrix.index = pd.to_datetime(matrix.index)
    matrix = matrix.reindex(columns=codes)
    combined_df = pd.concat([df, matrix], axis=0)
    combined_df = sanitize_price_df(combined_df)
    combined_df = combined_df[~combined_df.index.duplicated(keep="last")]
    combined_df = combined_df.loc[:, [c for c in codes if c in combined_df.columns]]
    return drop_all_nan_rows(combined_df)


def get_price_df_with_cache(codes, field):
    print(f"{field} 从本地行情数据库读取：{MARKET_DB_PATH}")
    df = load_price_matrix(
        CACHE_PREFIX,
        field,
        codes=codes,
        start_date=start_date,
        end_date=end_date,
        target_columns=codes,
        prefer_sqlite=True,
        fallback_pickle=False,
    )
    df = sanitize_price_df(df)
    df = drop_all_nan_rows(df)
    df = supplement_with_tushare_daily(df, codes, field, end_date)
    df = drop_all_nan_rows(df)
    if df.empty:
        raise ValueError(f"{field} 从本地行情数据库读取为空，请先运行 update_local_market_db.py 更新日行情。")
    return df

# =========================
# 5. 行情
# =========================
open_df = get_price_df_with_cache(stock_codes, "open")
low_df = get_price_df_with_cache(stock_codes, "low")
high_df = get_price_df_with_cache(stock_codes, "high")
close_df = get_price_df_with_cache(stock_codes, "close")
volume_df = get_price_df_with_cache(stock_codes, "volume")
amt_df = get_price_df_with_cache(stock_codes, "amt")

print("行情维度：", close_df.shape)
latest_data_ts = close_df.index[-1]
if latest_data_ts.strftime("%Y-%m-%d") != end_date:
    print(f"实际最新行情日期为 {latest_data_ts.strftime('%Y-%m-%d')}，策略输出日期同步调整。")
end_dt = latest_data_ts.to_pydatetime()
end_date = latest_data_ts.strftime("%Y-%m-%d")

if USE_HISTORICAL_CONSTITUENTS and not historical_constituent_snapshots.empty:
    universe_member = build_daily_universe_member_matrix(
        historical_constituent_snapshots,
        close_df.index,
        close_df.columns,
    )
    latest_member_count = int(universe_member.loc[latest_data_ts].sum())
    print(f"最新交易日中证800历史成分掩码数量：{latest_member_count}")
else:
    universe_member = pd.DataFrame(True, index=close_df.index, columns=close_df.columns)

# =========================
# 6. 均线
# =========================
ma5 = close_df.rolling(5).mean()
ma60 = close_df.rolling(60).mean()
ma120 = close_df.rolling(120).mean()

# =========================
# 9. 信号
# =========================
spread = ma5 - ma60
signal = np.sign(spread)
daily_ret = close_df.pct_change()
ma60_up = ma60 > ma60.shift(1)
ma120_up = ma120 > ma120.shift(1)
limit_gain = daily_ret <= SIGNAL_MAX_DAILY_RETURN
avg_amount_20d = amt_df.rolling(VOLUME_RATIO_LOOKBACK).mean()
score_details = calculate_golden_cross_score_details(ma5, ma60, volume_df)
score_ma5_ma60_gap = score_details["ma5_ma60_gap_score"]
score_ma60_5d_trend = score_details["ma60_5d_trend_score"]
score_volume_ratio = score_details["volume_ratio"]
score_volume_ratio_score = score_details["volume_ratio_score"]
candidate_score = score_details["total_score"]

cross = signal.diff()

golden_cross_raw = cross == 2
golden_cross_candidate = golden_cross_raw & ma60_up & ma120_up
golden_signal = golden_cross_candidate & limit_gain & universe_member
hard_filter_excluded_signal = golden_cross_raw & universe_member & ~golden_signal
buy_signal = golden_signal.shift(1).fillna(False).astype(bool)

# =========================
# 10. 持仓状态机
# 新开仓固定初始仓位，买入后仓位随涨跌自然漂移
# =========================
position = pd.DataFrame(0.0, index=close_df.index, columns=close_df.columns)
stop_signal = pd.DataFrame(False, index=close_df.index, columns=close_df.columns)
sell_reason = pd.DataFrame("", index=close_df.index, columns=close_df.columns)
sell_trigger_price = pd.DataFrame(np.nan, index=close_df.index, columns=close_df.columns)
sell_float_pnl = pd.DataFrame(np.nan, index=close_df.index, columns=close_df.columns)

current_holdings = {}
pending_open_sell_signals = {}
cash = 1.0
portfolio_values = []
turnover_records = []
closed_trade_returns = []
closed_trade_outcomes = []
closed_trade_reasons = []
closed_trade_holding_days = []
latest_trade_date = close_df.index[-1]
latest_intraday_stop_monitor_records = []
latest_low_efficiency_sell_plan_records = []
prev_date = None

for date in close_df.index:
    traded_amount = 0.0
    sold_today = set()

    # 先把老持仓按昨收到今收更新市值
    if prev_date is not None:
        for code in list(current_holdings.keys()):
            prev_price = close_df.at[prev_date, code]
            price = close_df.at[date, code]
            if pd.isna(prev_price) or pd.isna(price) or prev_price <= 0:
                continue
            current_holdings[code]["value"] *= price / prev_price

    # 执行上一交易日收盘后确认的低效持仓卖出信号：今日开盘卖出。
    for code, signal_info in list(pending_open_sell_signals.items()):
        if code not in current_holdings:
            del pending_open_sell_signals[code]
            continue

        open_price = open_df.at[date, code]
        close_price = close_df.at[date, code]
        if pd.isna(open_price) or pd.isna(close_price) or open_price <= 0 or close_price <= 0:
            continue

        holding_info = current_holdings[code]
        sell_price = open_price
        sell_value = holding_info["value"] * (sell_price / close_price)
        buy_cost = holding_info["cost_basis"] * (1 + TRANSACTION_COST_RATE)
        sell_proceeds = sell_value * (1 - TRANSACTION_COST_RATE)
        trade_return = sell_proceeds / buy_cost - 1 if buy_cost > 0 else 0
        holding_days = close_df.index.get_loc(date) - close_df.index.get_loc(holding_info["entry_date"])
        closed_trade_returns.append(trade_return)
        if trade_return > 0:
            trade_outcome = "win"
        else:
            trade_outcome = "loss"
        closed_trade_outcomes.append(trade_outcome)
        closed_trade_reasons.append(signal_info["reason"])
        closed_trade_holding_days.append(holding_days)
        cash += sell_proceeds
        traded_amount += sell_value
        stop_signal.at[date, code] = True
        sell_reason.at[date, code] = signal_info["reason"]
        sell_trigger_price.at[date, code] = sell_price
        sell_float_pnl.at[date, code] = sell_price / holding_info["entry_price"] - 1 if holding_info["entry_price"] > 0 else np.nan
        sold_today.add(code)
        del current_holdings[code]
        del pending_open_sell_signals[code]

    # 先用前一日信号在今日开盘买入；已持仓股票重复触发时允许继续加仓。
    portfolio_before_buy = cash + sum(info["value"] for info in current_holdings.values())
    available_slots = MAX_HOLDINGS - len(current_holdings)
    can_open_new_position = trade_start_ts is None or date >= trade_start_ts
    if can_open_new_position and cash > 0:
        buy_candidates = buy_signal.columns[buy_signal.loc[date]].tolist()
        buy_candidates = [
            code for code in buy_candidates
            if code not in sold_today and (code in current_holdings or available_slots > 0)
        ]

        if buy_candidates:
            score_prev = candidate_score.loc[prev_date, buy_candidates].dropna().sort_values(ascending=False)
            opened_new_positions = 0
            max_candidates = len([code for code in buy_candidates if code in current_holdings]) + available_slots
            for code in score_prev.head(max_candidates).index:
                is_existing_holding = code in current_holdings
                if not is_existing_holding and opened_new_positions >= available_slots:
                    continue

                open_price = open_df.at[date, code]
                high_price = high_df.at[date, code]
                close_price = close_df.at[date, code]
                prev_close = close_df.at[prev_date, code] if prev_date is not None else np.nan
                if (
                    pd.isna(open_price)
                    or pd.isna(high_price)
                    or pd.isna(close_price)
                    or open_price <= 0
                    or high_price <= 0
                    or close_price <= 0
                ):
                    continue

                if pd.notna(prev_close) and prev_close > 0:
                    limit_ratio = get_limit_up_ratio(code, code_to_name.get(code, code))
                    limit_up_price = prev_close * (1 + limit_ratio)
                    if open_price >= limit_up_price * 0.999:
                        continue

                target_value = portfolio_before_buy * INITIAL_WEIGHT
                max_affordable = cash / (1 + TRANSACTION_COST_RATE)
                buy_value = min(target_value, max_affordable)

                if buy_value <= 0:
                    break

                cash -= buy_value * (1 + TRANSACTION_COST_RATE)
                traded_amount += buy_value
                added_value_at_close = buy_value * (close_price / open_price)

                if is_existing_holding:
                    holding_info = current_holdings[code]
                    old_entry_price = holding_info["entry_price"]
                    old_cost_basis = holding_info["cost_basis"]
                    old_share_proxy = old_cost_basis / old_entry_price if old_entry_price > 0 else 0
                    added_share_proxy = buy_value / open_price
                    new_share_proxy = old_share_proxy + added_share_proxy
                    holding_info["entry_price"] = (
                        (old_cost_basis + buy_value) / new_share_proxy
                        if new_share_proxy > 0 else open_price
                    )
                    holding_info["entry_day_close_above_cost"] = (
                        holding_info["entry_day_close_above_cost"] or close_price >= open_price
                    )
                    holding_info["peak_price"] = max(holding_info["peak_price"], high_price)
                    holding_info["value"] += added_value_at_close
                    holding_info["cost_basis"] += buy_value
                else:
                    opened_new_positions += 1
                    current_holdings[code] = {
                        "entry_price": open_price,
                        "entry_date": date,
                        "entry_day_close_above_cost": close_price >= open_price,
                        "peak_price": high_price,
                        "value": added_value_at_close,
                        "cost_basis": buy_value
                    }

    # 最后检查老持仓是否在今日盘中触发回撤卖出
    for code in list(current_holdings.keys()):
        open_price = open_df.at[date, code]
        low_price = low_df.at[date, code]
        high_price = high_df.at[date, code]
        close_price = close_df.at[date, code]
        prev_close = close_df.at[prev_date, code] if prev_date is not None else np.nan
        if pd.isna(open_price) or pd.isna(low_price) or pd.isna(high_price) or pd.isna(close_price):
            continue

        holding_info = current_holdings[code]
        entry_price = holding_info["entry_price"]
        entry_date = holding_info["entry_date"]
        if date == latest_trade_date and date <= entry_date:
            latest_intraday_stop_monitor_records.append({
                "代码": code,
                "名称": code_to_name.get(code, code),
                "开仓日期": entry_date,
                "开仓价": entry_price,
                "前高价": holding_info["peak_price"],
                "历史最大浮盈": np.nan,
                "监控价位": np.nan,
                "今日开盘价": open_price,
                "今日盘中最低": low_price,
                "今日盘中最高": high_price,
                "最新价": close_price,
                "是否触发卖出": "否",
                "触发原因": "当日新开仓，不参与盘中卖出监控",
                "监控卖价": np.nan,
            })
        if date <= entry_date:
            continue

        sell_eval = evaluate_intraday_sell_signal(holding_info, open_price, low_price, high_price)
        prev_peak_price = sell_eval["prev_peak_price"]
        max_profit = sell_eval["max_profit"]
        sell_price = sell_eval["sell_price"]
        is_breakeven_exit = sell_eval["is_breakeven_exit"]
        sell_reason_text = sell_eval["sell_reason_text"]
        monitor_level = sell_eval["monitor_level"]

        limit_down_blocked = False

        if pd.notna(sell_price):
            if pd.notna(prev_close) and prev_close > 0:
                limit_ratio = get_limit_down_ratio(code, code_to_name.get(code, code))
                limit_down_price = prev_close * (1 - limit_ratio)
                if close_price <= limit_down_price * 1.001:
                    limit_down_blocked = True

        if date == latest_trade_date:
            latest_intraday_stop_monitor_records.append({
                "代码": code,
                "名称": code_to_name.get(code, code),
                "开仓日期": entry_date,
                "开仓价": entry_price,
                "前高价": prev_peak_price,
                "历史最大浮盈": max_profit,
                "监控价位": monitor_level,
                "今日开盘价": open_price,
                "今日盘中最低": low_price,
                "今日盘中最高": high_price,
                "最新价": close_price,
                "是否触发卖出": "是" if pd.notna(sell_price) and not limit_down_blocked else "否",
                "触发原因": "跌停附近封死，无法卖出" if limit_down_blocked else (sell_reason_text if pd.notna(sell_price) else ""),
                "监控卖价": sell_price if not limit_down_blocked else np.nan,
            })

        if pd.notna(sell_price) and limit_down_blocked:
            continue

        if pd.notna(sell_price):

            sell_value = holding_info["value"] * (sell_price / close_price)
            buy_cost = holding_info["cost_basis"] * (1 + TRANSACTION_COST_RATE)
            sell_proceeds = sell_value * (1 - TRANSACTION_COST_RATE)
            trade_return = sell_proceeds / buy_cost - 1 if buy_cost > 0 else 0
            holding_days = close_df.index.get_loc(date) - close_df.index.get_loc(entry_date)
            closed_trade_returns.append(trade_return)
            if is_breakeven_exit:
                trade_outcome = "draw"
            elif trade_return > 0:
                trade_outcome = "win"
            else:
                trade_outcome = "loss"
            closed_trade_outcomes.append(trade_outcome)
            closed_trade_reasons.append(sell_reason_text)
            closed_trade_holding_days.append(holding_days)
            cash += sell_proceeds
            traded_amount += sell_value
            stop_signal.at[date, code] = True
            sell_reason.at[date, code] = sell_reason_text
            sell_trigger_price.at[date, code] = sell_price
            sell_float_pnl.at[date, code] = sell_price / entry_price - 1 if entry_price > 0 else np.nan
            sold_today.add(code)
            del current_holdings[code]
            pending_open_sell_signals.pop(code, None)
            continue

        # 当日高点只能在收盘后确认，因此仅用于更新下一交易日可用的峰值。
        updated_peak_price = max(prev_peak_price, high_price)
        current_holdings[code]["peak_price"] = updated_peak_price

        holding_days = close_df.index.get_loc(date) - close_df.index.get_loc(entry_date)
        max_profit_after_update = updated_peak_price / entry_price - 1 if entry_price > 0 else -np.inf
        current_profit = close_price / entry_price - 1 if entry_price > 0 else np.nan
        is_low_efficiency_holding = (
            holding_days > LOW_EFFICIENCY_MIN_HOLDING_DAYS
            and max_profit_after_update <= LOW_EFFICIENCY_MAX_PROFIT_THRESHOLD
            and current_profit <= LOW_EFFICIENCY_CURRENT_PROFIT_THRESHOLD
        )

        if is_low_efficiency_holding:
            low_efficiency_reason = (
                f"低效持仓卖出：持有超过{LOW_EFFICIENCY_MIN_HOLDING_DAYS}个交易日，"
                f"历史最高浮盈未超过{LOW_EFFICIENCY_MAX_PROFIT_THRESHOLD:.0%}且当前浮盈不超过"
                f"{LOW_EFFICIENCY_CURRENT_PROFIT_THRESHOLD:.0%}"
            )
            pending_open_sell_signals[code] = {
                "signal_date": date,
                "reason": low_efficiency_reason,
                "holding_days": holding_days,
                "max_profit": max_profit_after_update,
                "current_profit": current_profit,
            }
            if date == latest_trade_date:
                latest_low_efficiency_sell_plan_records.append({
                    "代码": code,
                    "名称": code_to_name.get(code, code),
                    "开仓日期": entry_date,
                    "开仓价": entry_price,
                    "最新价": close_price,
                    "持有交易日": holding_days,
                    "历史最高浮盈": max_profit_after_update,
                    "当前浮盈": current_profit,
                    "计划动作": "下一交易日开盘卖出",
                    "触发原因": low_efficiency_reason,
                })
        else:
            pending_open_sell_signals.pop(code, None)

    portfolio_value = cash + sum(info["value"] for info in current_holdings.values())
    portfolio_values.append(portfolio_value)
    turnover_records.append(traded_amount / portfolio_value if portfolio_value > 0 else 0.0)

    for code, info in current_holdings.items():
        position.at[date, code] = info["value"] / portfolio_value if portfolio_value > 0 else 0.0

    prev_date = date

# =========================
# 11. 收益与净值
# =========================
nav = pd.Series(portfolio_values, index=close_df.index, name="净值")
strategy_ret = nav.pct_change().fillna(0)
turnover = pd.Series(turnover_records, index=close_df.index, name="换手率")

if trade_start_ts is not None:
    valid_analysis_dates = nav.index[nav.index >= trade_start_ts]
    if len(valid_analysis_dates) == 0:
        raise ValueError("TRADE_START_DATE 晚于当前获取到的全部行情日期，请检查参数设置。")
    analysis_start_date = valid_analysis_dates[0]
else:
    analysis_start_date = nav.index[0]

nav_analysis = nav.loc[analysis_start_date:]
strategy_ret_analysis = strategy_ret.loc[analysis_start_date:]
turnover_analysis = turnover.loc[analysis_start_date:]
position_analysis = position.loc[analysis_start_date:]

benchmark_name = "中证800指数"
benchmark_close = fetch_tushare_index_close(TUSHARE_INDEX_CODE, analysis_start_date, end_date)
benchmark_close = benchmark_close.reindex(nav_analysis.index).ffill()
if benchmark_close.notna().sum() >= 2:
    benchmark_nav = benchmark_close / benchmark_close.dropna().iloc[0]
else:
    benchmark_name = "中证800等权基准"
    benchmark_close_analysis = close_df.loc[analysis_start_date:]
    benchmark_nav = (
        benchmark_close_analysis
        .pct_change()
        .mean(axis=1, skipna=True)
        .fillna(0)
        .add(1)
        .cumprod()
    )
benchmark_nav.name = f"{benchmark_name}净值"
return_curve_df = pd.DataFrame({
    "日期": nav_analysis.index.strftime("%Y-%m-%d"),
    "策略累计收益": nav_analysis / nav_analysis.iloc[0] - 1,
    f"{benchmark_name}累计收益": benchmark_nav / benchmark_nav.iloc[0] - 1,
    "策略净值": nav_analysis,
    f"{benchmark_name}净值": benchmark_nav,
})

# =========================
# ⭐ 15. 年化收益（修正版）
# =========================
has_position = position_analysis.sum(axis=1).gt(0)
if has_position.any():
    first_trade_date = has_position.idxmax()
    nav_active = nav_analysis.loc[first_trade_date:]
    ret_active = strategy_ret_analysis.loc[first_trade_date:]
    position_active = position_analysis.loc[first_trade_date:]
    turnover_active = turnover_analysis.loc[first_trade_date:]

    annual_ret = nav_active.iloc[-1] ** (252 / len(nav_active)) - 1
    annual_vol = ret_active.std() * np.sqrt(252)
    sharpe = annual_ret / annual_vol if annual_vol != 0 else 0
    rolling_max = nav_active.cummax()
    drawdown = nav_active / rolling_max - 1
    max_dd = drawdown.min()
    holding_count = (position_active > 0).sum(axis=1)
    avg_holding = holding_count.mean()
    avg_turnover = turnover_active.mean()
    annual_turnover = avg_turnover * 252
    win_rate = (ret_active > 0).mean()
    trade_count = len(closed_trade_returns)
    if trade_count > 0:
        trade_outcome_series = pd.Series(closed_trade_outcomes)
        trade_win_count = (trade_outcome_series == "win").sum()
        trade_draw_count = (trade_outcome_series == "draw").sum()
        trade_loss_count = (trade_outcome_series == "loss").sum()
        decisive_trade_count = trade_outcome_series.isin(["win", "loss"]).sum()
        trade_win_rate = (
            trade_win_count / decisive_trade_count
            if decisive_trade_count > 0 else 0
        )
    else:
        trade_win_count = 0
        trade_draw_count = 0
        trade_loss_count = 0
        trade_win_rate = 0
else:
    annual_ret = 0
    annual_vol = 0
    sharpe = 0
    max_dd = 0
    avg_holding = 0
    annual_turnover = 0
    win_rate = 0
    trade_count = 0
    trade_win_count = 0
    trade_draw_count = 0
    trade_loss_count = 0
    trade_win_rate = 0

stats = pd.DataFrame({
    "指标": ["年化收益","年化波动","夏普比率","最大回撤","平均持仓数","年化换手率","日胜率","平仓笔数","单笔盈利数","单笔平局数","单笔亏损数","单笔胜率(平局不计入)"],
    "数值": [annual_ret, annual_vol, sharpe, max_dd, avg_holding, annual_turnover, win_rate, trade_count, trade_win_count, trade_draw_count, trade_loss_count, trade_win_rate]
})

if trade_count > 0:
    close_reason_stats = pd.DataFrame({
        "平仓原因": closed_trade_reasons,
        "单笔收益": closed_trade_returns,
        "持有天数": closed_trade_holding_days,
    })
    close_reason_stats = (
        close_reason_stats
        .groupby("平仓原因", dropna=False)
        .agg(
            平仓笔数=("平仓原因", "size"),
            平均单笔收益=("单笔收益", "mean"),
            平均持有天数=("持有天数", "mean"),
        )
        .reset_index()
    )
    close_reason_stats["平仓占比"] = close_reason_stats["平仓笔数"] / trade_count
    close_reason_stats = close_reason_stats[
        ["平仓原因", "平仓笔数", "平仓占比", "平均单笔收益", "平均持有天数"]
    ].sort_values(by="平仓笔数", ascending=False)
else:
    close_reason_stats = pd.DataFrame(
        columns=["平仓原因", "平仓笔数", "平仓占比", "平均单笔收益", "平均持有天数"]
    )

print("\n【策略指标】")
print(stats)

# =========================
# 17. 持仓（全历史）
# =========================
def extract_position(df, code_map):
    records = []
    for date in df.index:
        row = df.loc[date]
        holdings = row[row > 0]
        for code, weight in holdings.items():
            records.append({
                "日期": date,
                "代码": code,
                "名称": code_map.get(code, code),
                "权重": weight
            })
    return pd.DataFrame(records)

position_df = extract_position(position_analysis, code_to_name)
position_df = position_df.sort_values(by="日期", ascending=False)

# =========================
# 18. 当前持仓
# =========================
today = position.index[-1]

holding_records = []
for code, info in current_holdings.items():
    latest_price = close_df.at[today, code]
    entry_price = info["entry_price"]
    float_pnl = latest_price / entry_price - 1 if entry_price > 0 and pd.notna(latest_price) else np.nan
    max_float_pnl = info["peak_price"] / entry_price - 1 if entry_price > 0 else np.nan
    holding_days = close_df.index.get_loc(today) - close_df.index.get_loc(info["entry_date"])
    holding_records.append({
        "代码": code,
        "名称": code_to_name.get(code, code),
        "开仓日期": info["entry_date"],
        "开仓价": entry_price,
        "最新价": latest_price,
        "最新市值": info["value"],
        "组合权重": position.at[today, code] if code in position.columns else 0.0,
        "持有交易日": holding_days,
        "历史最高浮盈": max_float_pnl,
        "浮赢浮亏": float_pnl
    })

current_holding_df = pd.DataFrame(holding_records)
if not current_holding_df.empty:
    current_holding_df = current_holding_df.sort_values(by="组合权重", ascending=False)
    current_used_position = current_holding_df["组合权重"].sum()
else:
    current_used_position = 0.0

# =========================
# 19. 当日信号
# =========================
golden_trigger_today = golden_signal.loc[today]
golden_exec_today = buy_signal.loc[today]
hard_filter_excluded_today = hard_filter_excluded_signal.loc[today]
stop_today = stop_signal.loc[today]

golden_trigger_list = golden_trigger_today[golden_trigger_today].index.tolist()
golden_exec_list = golden_exec_today[golden_exec_today].index.tolist()
hard_filter_excluded_list = hard_filter_excluded_today[hard_filter_excluded_today].index.tolist()
stop_list = stop_today[stop_today].index.tolist()

golden_trigger_df = pd.DataFrame({
    "代码": golden_trigger_list,
    "名称": [code_to_name.get(c, c) for c in golden_trigger_list],
    "评分": [candidate_score.at[today, c] for c in golden_trigger_list],
    "MA5相对MA60强度": [score_ma5_ma60_gap.at[today, c] for c in golden_trigger_list],
    "MA60近5日趋势分": [score_ma60_5d_trend.at[today, c] for c in golden_trigger_list],
    "量比": [score_volume_ratio.at[today, c] for c in golden_trigger_list],
    "量比加分": [score_volume_ratio_score.at[today, c] for c in golden_trigger_list],
    "20日平均成交额": [avg_amount_20d.at[today, c] for c in golden_trigger_list],
    "信号": "最新金叉触发"
})
if not golden_trigger_df.empty:
    golden_trigger_df = golden_trigger_df.sort_values(by="评分", ascending=False, na_position="last")

golden_exec_df = pd.DataFrame({
    "代码": golden_exec_list,
    "名称": [code_to_name.get(c, c) for c in golden_exec_list],
    "评分": [candidate_score.shift(1).at[today, c] for c in golden_exec_list],
    "MA5相对MA60强度": [score_ma5_ma60_gap.shift(1).at[today, c] for c in golden_exec_list],
    "MA60近5日趋势分": [score_ma60_5d_trend.shift(1).at[today, c] for c in golden_exec_list],
    "量比": [score_volume_ratio.shift(1).at[today, c] for c in golden_exec_list],
    "量比加分": [score_volume_ratio_score.shift(1).at[today, c] for c in golden_exec_list],
    "20日平均成交额": [avg_amount_20d.shift(1).at[today, c] for c in golden_exec_list],
    "信号": "昨日金叉今日执行"
})
if not golden_exec_df.empty:
    golden_exec_df = golden_exec_df.sort_values(by="评分", ascending=False, na_position="last")

def build_hard_filter_excluded_record(signal_date, code, prompt_text, record_type):
    excluded_reasons = []
    if not bool(ma60_up.at[signal_date, code]):
        excluded_reasons.append("MA60未向上")
    if not bool(ma120_up.at[signal_date, code]):
        excluded_reasons.append("MA120未向上")
    if not bool(limit_gain.at[signal_date, code]):
        excluded_reasons.append("当日涨幅超过阈值")

    return {
        "记录类型": record_type,
        "信号日期": signal_date,
        "最新交易日": today,
        "代码": code,
        "名称": code_to_name.get(code, code),
        "评分": candidate_score.at[signal_date, code],
        "信号日涨幅": daily_ret.at[signal_date, code],
        "最新交易日涨幅": daily_ret.at[today, code] if code in daily_ret.columns else np.nan,
        "涨幅阈值": SIGNAL_MAX_DAILY_RETURN,
        "超出阈值": daily_ret.at[signal_date, code] - SIGNAL_MAX_DAILY_RETURN,
        "MA60是否向上": "是" if bool(ma60_up.at[signal_date, code]) else "否",
        "MA120是否向上": "是" if bool(ma120_up.at[signal_date, code]) else "否",
        "涨幅是否合格": "是" if bool(limit_gain.at[signal_date, code]) else "否",
        "剔除原因": "、".join(excluded_reasons),
        "MA5相对MA60强度": score_ma5_ma60_gap.at[signal_date, code],
        "MA60近5日趋势分": score_ma60_5d_trend.at[signal_date, code],
        "量比": score_volume_ratio.at[signal_date, code],
        "量比加分": score_volume_ratio_score.at[signal_date, code],
        "20日平均成交额": avg_amount_20d.at[signal_date, code],
        "提示": prompt_text
    }

hard_filter_excluded_records = []
for code in hard_filter_excluded_list:
    hard_filter_excluded_records.append(
        build_hard_filter_excluded_record(
            today,
            code,
            "今日发生MA5上穿MA60，但未通过硬性过滤条件",
            "今日剔除"
        )
    )

prev_trade_date = close_df.index[-2] if len(close_df.index) >= 2 else None
if prev_trade_date is not None:
    prev_filtered = hard_filter_excluded_signal.loc[prev_trade_date]
    prev_filtered_list = prev_filtered[prev_filtered].index.tolist()
    for code in prev_filtered_list:
        hard_filter_excluded_records.append(
            build_hard_filter_excluded_record(
                prev_trade_date,
                code,
                "昨日发生MA5上穿MA60但被剔除，因此今日不执行买入",
                "昨日剔除"
            )
        )

hard_filter_excluded_df = pd.DataFrame(hard_filter_excluded_records)
if not hard_filter_excluded_df.empty:
    hard_filter_excluded_df = hard_filter_excluded_df.sort_values(
        by=["信号日期", "评分"],
        ascending=[False, False],
        na_position="last"
    )

stop_df = pd.DataFrame({
    "代码": stop_list,
    "名称": [code_to_name.get(c, c) for c in stop_list],
    "触发原因": [sell_reason.at[today, c] for c in stop_list],
    "触发卖价": [sell_trigger_price.at[today, c] for c in stop_list],
    "浮赢浮亏": [sell_float_pnl.at[today, c] for c in stop_list],
    "信号": "最新交易日卖出"
})
latest_intraday_stop_monitor_df = pd.DataFrame(latest_intraday_stop_monitor_records)
if not latest_intraday_stop_monitor_df.empty:
    latest_intraday_stop_monitor_df = latest_intraday_stop_monitor_df.sort_values(
        by=["是否触发卖出", "历史最大浮盈"],
        ascending=[False, False],
        na_position="last"
    )

latest_low_efficiency_sell_plan_df = pd.DataFrame(latest_low_efficiency_sell_plan_records)
if not latest_low_efficiency_sell_plan_df.empty:
    latest_low_efficiency_sell_plan_df = latest_low_efficiency_sell_plan_df.sort_values(
        by=["持有交易日", "当前浮盈"],
        ascending=[False, True],
        na_position="last"
    )

# =========================
# 20. 输出
# =========================
output_dir = os.path.join(BASE_DIR, "输出", "金叉执行结果")
os.makedirs(output_dir, exist_ok=True)
output_file = os.path.join(output_dir, f"金叉买入_前高回撤10%卖出策略_中证800_Tushare版_{end_date}.xlsx")
return_curve_image_file = os.path.join(output_dir, f"收益走势图_策略_vs_{benchmark_name}_Tushare版_{end_date}.png")


def save_return_curve_image(df, image_file):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        import matplotlib.ticker as mtick
        import matplotlib.dates as mdates
    except ImportError:
        print("未安装 matplotlib，跳过 PNG 收益走势图。可执行：python3 -m pip install matplotlib")
        return False

    plot_df = df.copy()
    plot_df["日期"] = pd.to_datetime(plot_df["日期"])

    plt.rcParams["font.sans-serif"] = [
        "Arial Unicode MS",
        "PingFang SC",
        "Heiti SC",
        "SimHei",
        "DejaVu Sans",
    ]
    plt.rcParams["axes.unicode_minus"] = False

    fig, ax = plt.subplots(figsize=(13.5, 7.2), dpi=180)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("#fbfbfb")

    ax.plot(
        plot_df["日期"],
        plot_df["策略累计收益"],
        label="策略累计收益",
        color="#2563eb",
        linewidth=2.3,
    )
    ax.plot(
        plot_df["日期"],
        plot_df[f"{benchmark_name}累计收益"],
        label=f"{benchmark_name}累计收益",
        color="#64748b",
        linewidth=2.0,
        linestyle="--",
    )
    ax.axhline(0, color="#111827", linewidth=0.8, alpha=0.55)

    latest_date = plot_df["日期"].iloc[-1].strftime("%Y-%m-%d")
    strategy_latest = plot_df["策略累计收益"].iloc[-1]
    benchmark_latest = plot_df[f"{benchmark_name}累计收益"].iloc[-1]
    excess_latest = strategy_latest - benchmark_latest
    title = (
        f"金叉策略累计收益 vs {benchmark_name}  |  截至 {latest_date}\n"
        f"策略 {strategy_latest:.2%}    基准 {benchmark_latest:.2%}    超额 {excess_latest:.2%}"
    )
    ax.set_title(title, fontsize=15, fontweight="bold", color="#111827", pad=18)
    ax.set_xlabel("日期", fontsize=11, color="#374151", labelpad=10)
    ax.set_ylabel("累计收益率", fontsize=11, color="#374151", labelpad=10)
    ax.yaxis.set_major_formatter(mtick.PercentFormatter(1.0))
    ax.xaxis.set_major_locator(mdates.MonthLocator(interval=3))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m"))
    ax.grid(True, axis="y", color="#d1d5db", linewidth=0.8, alpha=0.7)
    ax.grid(True, axis="x", color="#e5e7eb", linewidth=0.5, alpha=0.35)
    ax.legend(
        loc="upper left",
        frameon=True,
        facecolor="white",
        edgecolor="#e5e7eb",
        framealpha=0.95,
        fontsize=10,
    )
    for spine in ["top", "right"]:
        ax.spines[spine].set_visible(False)
    ax.spines["left"].set_color("#9ca3af")
    ax.spines["bottom"].set_color("#9ca3af")
    ax.tick_params(axis="x", labelrotation=35, labelsize=9, colors="#4b5563")
    ax.tick_params(axis="y", labelsize=9, colors="#4b5563")

    fig.tight_layout()
    fig.savefig(image_file, bbox_inches="tight")
    plt.close(fig)
    return True


def add_return_curve_image(writer, sheet_name, df, image_file):
    sheet = writer.sheets[sheet_name]
    header_row = 35
    first_data_row = header_row + 1

    for cell in sheet[header_row]:
        cell.font = Font(bold=True)

    if os.path.exists(image_file):
        image = OpenpyxlImage(image_file)
        image.width = 1050
        image.height = 560
        sheet.add_image(image, "A1")

    for col_letter in ["B", "C"]:
        for cell in sheet[col_letter][header_row:]:
            cell.number_format = "0.00%"
    for col_letter in ["D", "E"]:
        for cell in sheet[col_letter][header_row:]:
            cell.number_format = "0.0000"
    sheet.freeze_panes = f"A{first_data_row}"
    sheet.column_dimensions["A"].width = 14
    for col_letter in ["B", "C", "D", "E"]:
        sheet.column_dimensions[col_letter].width = 18


save_return_curve_image(return_curve_df, return_curve_image_file)

with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
    nav_analysis.to_frame("净值").to_excel(writer, sheet_name="净值")
    return_curve_df.to_excel(writer, sheet_name="收益走势图", index=False, startrow=34)
    add_return_curve_image(writer, "收益走势图", return_curve_df, return_curve_image_file)
    stats.to_excel(writer, sheet_name="策略指标", index=False)
    metrics_sheet = writer.sheets["策略指标"]
    metrics_sheet.cell(row=len(stats) + 3, column=1, value="平仓原因统计")
    close_reason_stats.to_excel(
        writer,
        sheet_name="策略指标",
        startrow=len(stats) + 3,
        index=False,
    )
    position_df.to_excel(writer, sheet_name="每日持仓", index=False)
    current_holding_df.to_excel(writer, sheet_name="当前持仓", startrow=2, index=False)
    current_holding_sheet = writer.sheets["当前持仓"]
    current_holding_sheet.cell(row=1, column=1, value="当前已使用仓位")
    current_holding_sheet.cell(row=1, column=2, value=current_used_position)
    golden_trigger_df.to_excel(writer, sheet_name="最新金叉触发", index=False)
    golden_exec_df.to_excel(writer, sheet_name="今日执行买入", index=False)
    hard_filter_excluded_df.to_excel(writer, sheet_name="硬性条件剔除", index=False)
    stop_df.to_excel(writer, sheet_name="最新交易日卖出", index=False)
    latest_intraday_stop_monitor_df.to_excel(writer, sheet_name="最新盘中卖出监控", index=False)
    latest_low_efficiency_sell_plan_df.to_excel(writer, sheet_name="明日低效持仓卖出", index=False)

print("\n【最新信号】")
print(f"最新金叉触发数量: {len(golden_trigger_list)}")
print(f"昨日金叉今日执行数量: {len(golden_exec_list)}")
print(f"硬性条件剔除提示数量: {len(hard_filter_excluded_df)}")
print(f"最新交易日卖出数量: {len(stop_list)}")
print(f"最新盘中卖出监控数量: {len(latest_intraday_stop_monitor_df)}")
print(f"明日低效持仓开盘卖出数量: {len(latest_low_efficiency_sell_plan_df)}")
print(f"当前持仓数量: {len(current_holding_df)}")

print("输出完成：", output_file)
