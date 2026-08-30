"""全 A 股质量成长 100 研究版策略入口。

这个文件固定了 2026-08-24 的 800/100/50 对比中，全 A 股 100 只版本的口径：

- 全部 A 股历史截面为选股范围；
- 每月调仓，固定 100 只，按行业参考名额选股；
- 20% 成分缓冲；
- 行业内与全市场混合评分；
- 按实际公告日可得数据自算 TTM ROE；
- 开启业绩预告/快报，置信权重分别为 0.40/0.70；
- 前复权行情，默认单边交易成本 0.5%；
- 沪深 300 全收益指数为主基准。

为避免复制整个计算引擎后出现两份逻辑分叉，本文件只保留研究参数和
执行入口，底层数据、评分、选股、权重与回测计算复用同目录下已验证的
`行业比较混合评分_全A股质量成长800指数.py`。
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import sys
from dataclasses import asdict, replace
from datetime import date
from pathlib import Path


SCRIPT_DIR = Path(__file__).resolve().parent
ENGINE_PATH = SCRIPT_DIR / "行业比较混合评分_全A股质量成长800指数.py"
OUTPUT_DIR = SCRIPT_DIR.parent / "输出" / "全A股质量成长100"

# 研究主参数：如需做系统性参数实验，优先在这里修改。
STRATEGY_NAME = "全A股质量成长100"
TARGET_CONSTITUENTS = 100
CONSTITUENT_BUFFER_RATIO = 0.20
DEFAULT_COST_RATE = 0.005
DEFAULT_START_DATE = "2022-01-01"
DEFAULT_END_DATE = "2026-08-20"
EARNINGS_NOTICE_CONFIDENCE = 0.40
EARNINGS_EXPRESS_CONFIDENCE = 0.70


def load_engine():
    script_dir = str(SCRIPT_DIR)
    if script_dir not in sys.path:
        sys.path.insert(0, script_dir)
    spec = importlib.util.spec_from_file_location("quality_growth_100_engine", ENGINE_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法加载底层策略引擎：{ENGINE_PATH}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def make_config(engine, earnings_guidance_enabled: bool = True):
    """显式覆盖关键参数，不依赖底层引擎可能变化的默认值。"""
    config = replace(
        engine.hybrid_config(TARGET_CONSTITUENTS),
        max_constituents=TARGET_CONSTITUENTS,
        rebalance_frequency="M",
        stock_factor_scope="blended",
        constituent_selection_mode="industry_quota",
        constituent_buffer_ratio=CONSTITUENT_BUFFER_RATIO,
        industry_quota_min_score=45.0,
        industry_target_mode="industry_score",
        quality_data_mode="calculated_ttm",
        valuation_data_mode="calculated_pit",
        earnings_guidance_enabled=earnings_guidance_enabled,
        earnings_notice_confidence=EARNINGS_NOTICE_CONFIDENCE,
        earnings_express_confidence=EARNINGS_EXPRESS_CONFIDENCE,
    )
    engine.validate_config(config)
    return config


def default_output_path(as_of: str) -> Path:
    stamp = as_of.replace("-", "")
    return OUTPUT_DIR / f"{STRATEGY_NAME}_行业比较混合评分_{stamp}.xlsx"


def find_previous_components(as_of: str, current_output: Path) -> Path | None:
    cutoff = as_of.replace("-", "")
    candidates: list[tuple[str, Path]] = []
    for path in current_output.parent.glob(f"{STRATEGY_NAME}_行业比较混合评分_*.xlsx"):
        stamp = path.stem.rsplit("_", 1)[-1]
        if len(stamp) == 8 and stamp.isdigit() and stamp < cutoff:
            candidates.append((stamp, path))
    return max(candidates, default=("", None), key=lambda item: item[0])[1]


def run_build(engine, args, config, output_path: Path) -> Path:
    previous_path = args.previous_components or find_previous_components(args.as_of, output_path)
    previous_constituents = (
        engine.load_previous_constituents(previous_path)
        if previous_path is not None
        else None
    )
    selected, industry_summary, all_scores, diagnostics = engine.build_index(
        args.as_of,
        config,
        args.benchmark_weights,
        previous_constituents,
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    engine.export_result(
        selected,
        industry_summary,
        all_scores,
        diagnostics,
        config,
        output_path,
    )
    print(json.dumps(diagnostics, ensure_ascii=False, indent=2))
    print(f"上一期成分：{previous_path or '无（本期不使用历史成分缓冲）'}")
    print(f"单期构建结果：{output_path}")
    return output_path


def relabel_backtest_workbook(output_path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(output_path)
    if "绩效汇总" in workbook.sheetnames:
        workbook["绩效汇总"]["A1"] = f"{STRATEGY_NAME}回测绩效汇总"
    workbook.save(output_path)


def run_backtest(engine, args, config, output_path: Path) -> Path:
    daily, periods, yearly, metrics = engine.run_backtest(
        args.start_date,
        args.end_date,
        args.cost_rate,
        config,
    )
    metrics.setdefault("模型配置", {}).update({
        "策略名称": STRATEGY_NAME,
        "选股范围": "全部A股历史截面",
        "目标成分数": TARGET_CONSTITUENTS,
        "单边交易成本": args.cost_rate,
        "回测开始日": args.start_date,
        "回测结束日": args.end_date,
    })
    output_path.parent.mkdir(parents=True, exist_ok=True)
    engine.write_backtest_outputs(output_path, daily, periods, yearly, metrics)
    relabel_backtest_workbook(output_path)
    print(json.dumps(metrics, ensure_ascii=False, indent=2))
    print(f"回测结果：{output_path}")
    return output_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--task",
        choices=["config", "build", "backtest", "both"],
        default="backtest",
        help="config 仅打印参数；build 构建单期成分；backtest 运行回测",
    )
    parser.add_argument("--as-of", default=date.today().isoformat())
    parser.add_argument("--start-date", default=DEFAULT_START_DATE)
    parser.add_argument("--end-date", default=DEFAULT_END_DATE)
    parser.add_argument("--cost-rate", type=float, default=DEFAULT_COST_RATE)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--benchmark-weights", type=Path)
    parser.add_argument("--previous-components", type=Path)
    parser.add_argument(
        "--disable-earnings-guidance",
        action="store_true",
        help="研究开关：关闭业绩预告/快报（默认为开启，与原100只结果一致）",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    engine = load_engine()
    args.as_of = engine.normalize_date(args.as_of)
    config = make_config(
        engine,
        earnings_guidance_enabled=not args.disable_earnings_guidance,
    )
    output_path = args.output or default_output_path(args.as_of)

    if args.task == "config":
        print(json.dumps(asdict(config), ensure_ascii=False, indent=2))
        return
    if args.task in {"build", "both"}:
        run_build(engine, args, config, output_path)
    if args.task in {"backtest", "both"}:
        run_backtest(engine, args, config, output_path)


if __name__ == "__main__":
    main()
