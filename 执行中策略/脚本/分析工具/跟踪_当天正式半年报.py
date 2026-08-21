import argparse
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd


SCRIPT_DIR = Path(__file__).resolve().parent
SCRIPT_ROOT = SCRIPT_DIR.parent
OUTPUT_DIR = SCRIPT_ROOT.parent / "输出" / "半年报跟踪"
A_SHARE_SECTOR_ID = "a001010100000000"
REPORT_BATCH_SIZE = 800
MARKET_BATCH_SIZE = 200


# 正式财报指标。金额字段必须搭配 rptType=1（合并报表），
# 否则 Wind 可能 ErrorCode=0 但返回全空。
REPORT_FIELDS = {
    "ann_date": "stm_issuingdate",
    "revenue": "tot_oper_rev",
    "revenue_yoy": "yoy_or",
    "net_profit": "np_belongto_parcomsh",
    "net_profit_yoy": "yoynetprofit",
    "deducted_net_profit": "deductedprofit",
    "deducted_net_profit_yoy": "yoynetprofit_deducted",
}


QUARTER_CALC_FIELDS = {
    "revenue": "tot_oper_rev",
    "net_profit": "np_belongto_parcomsh",
    "deducted_net_profit": "deductedprofit",
}


MARKET_FIELDS = {
    "industry": ("wicsname2024", "tradeDate={trade_date};industryType=1;"),
    "pe_ttm": ("pe_ttm", "tradeDate={trade_date}"),
}


PROFIT_NOTICE_FIELDS = {
    "notice_rpt_date": "profitnotice_lastrptdate",
    "notice_date": "profitnotice_date",
    "notice_type": "profitnotice_style",
    "notice_net_profit_min": "profitnotice_netprofitmin",
    "notice_net_profit_max": "profitnotice_netprofitmax",
    "notice_yoy_min": "profitnotice_changemin",
    "notice_yoy_max": "profitnotice_changemax",
}


# Wind A 股券商一致预期是年度口径，不是半年度口径。
# 脚本取披露前一交易日、近 180 天机构最新预测均值，只用于计算半年完成率。
CONSENSUS_FIELDS = {
    "consensus_net_profit_fy": "west_netprofit",
    "consensus_np_institutions": "west_instnum_np",
}


DISPLAY_COLUMNS = [
    "ann_date",
    "wind_code",
    "sec_name",
    "industry",
    "rpt_date",
    "revenue",
    "revenue_yoy",
    "net_profit",
    "net_profit_yoy",
    "deducted_net_profit",
    "deducted_net_profit_yoy",
    "q2_revenue",
    "q2_revenue_yoy",
    "q2_net_profit",
    "q2_net_profit_yoy",
    "q2_deducted_net_profit",
    "q2_deducted_net_profit_yoy",
    "notice_date",
    "notice_type",
    "notice_net_profit_min",
    "notice_net_profit_mid",
    "notice_net_profit_max",
    "notice_vs_mid_pct",
    "notice_comparison",
    "notice_yoy_min",
    "notice_yoy_mid",
    "notice_yoy_max",
    "notice_yoy_diff_ppt",
    "consensus_asof",
    "consensus_net_profit_fy",
    "consensus_np_institutions",
    "h1_np_completion_pct",
    "pe_ttm",
]


OUTPUT_COLUMN_NAMES = {
    "ann_date": "披露日",
    "wind_code": "代码",
    "sec_name": "名称",
    "industry": "行业",
    "rpt_date": "报告期",
    "revenue": "H1营业收入(亿元)",
    "revenue_yoy": "H1营收同比(%)",
    "net_profit": "H1归母净利润(亿元)",
    "net_profit_yoy": "H1归母净利润同比(%)",
    "deducted_net_profit": "H1扣非净利润(亿元)",
    "deducted_net_profit_yoy": "H1扣非净利润同比(%)",
    "q2_revenue": "Q2营业收入(亿元)",
    "q2_revenue_yoy": "Q2营收同比(%)",
    "q2_net_profit": "Q2归母净利润(亿元)",
    "q2_net_profit_yoy": "Q2归母净利润同比(%)",
    "q2_deducted_net_profit": "Q2扣非净利润(亿元)",
    "q2_deducted_net_profit_yoy": "Q2扣非净利润同比(%)",
    "notice_date": "业绩预告日",
    "notice_type": "业绩预告类型",
    "notice_net_profit_min": "预告净利润下限(亿元)",
    "notice_net_profit_mid": "预告净利润中值(亿元)",
    "notice_net_profit_max": "预告净利润上限(亿元)",
    "notice_vs_mid_pct": "较预告中值(%)",
    "notice_comparison": "业绩预告比较",
    "notice_yoy_min": "预告同比下限(%)",
    "notice_yoy_mid": "预告同比中值(%)",
    "notice_yoy_max": "预告同比上限(%)",
    "notice_yoy_diff_ppt": "实际同比-预告中值(百分点)",
    "consensus_asof": "一致预期截点",
    "consensus_net_profit_fy": "券商一致预期归母净利润-全年(亿元)",
    "consensus_np_institutions": "净利润预测机构数",
    "h1_np_completion_pct": "半年归母净利润完成率(%)",
    "pe_ttm": "PE_TTM",
}


MONEY_COLUMNS = [
    "revenue",
    "net_profit",
    "deducted_net_profit",
    "q2_revenue",
    "q2_net_profit",
    "q2_deducted_net_profit",
    "notice_net_profit_min",
    "notice_net_profit_mid",
    "notice_net_profit_max",
    "consensus_net_profit_fy",
]
PERCENT_COLUMNS = [
    "revenue_yoy",
    "net_profit_yoy",
    "deducted_net_profit_yoy",
    "q2_revenue_yoy",
    "q2_net_profit_yoy",
    "q2_deducted_net_profit_yoy",
    "notice_vs_mid_pct",
    "notice_yoy_min",
    "notice_yoy_mid",
    "notice_yoy_max",
    "notice_yoy_diff_ppt",
    "h1_np_completion_pct",
]


TERMINAL_COLUMNS = [
    "ann_date",
    "wind_code",
    "sec_name",
    "industry",
    "revenue",
    "revenue_yoy",
    "net_profit",
    "net_profit_yoy",
    "deducted_net_profit",
    "deducted_net_profit_yoy",
    "q2_revenue",
    "q2_revenue_yoy",
    "q2_net_profit",
    "q2_net_profit_yoy",
    "q2_deducted_net_profit",
    "q2_deducted_net_profit_yoy",
    "notice_net_profit_min",
    "notice_net_profit_max",
    "notice_comparison",
    "consensus_net_profit_fy",
    "consensus_np_institutions",
    "h1_np_completion_pct",
]


def parse_yyyy_mm_dd(value):
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()
    return datetime.strptime(str(value), "%Y-%m-%d").date()


def yyyymmdd(value):
    return value.strftime("%Y%m%d")


def default_rpt_date(query_date):
    """七月起跟踪当年半年报，一至六月跟踪上年半年报。"""
    year = query_date.year if query_date.month >= 7 else query_date.year - 1
    return date(year, 6, 30)


def default_lookback_days(query_date):
    """覆盖上一个工作日至查询日，避免周末披露在周一被漏掉。"""
    previous_workday = query_date - timedelta(days=1)
    while previous_workday.weekday() >= 5:
        previous_workday -= timedelta(days=1)
    return (query_date - previous_workday).days + 1


def ensure_wind_ok(data, label):
    if data.ErrorCode != 0:
        raise RuntimeError(f"Wind {label} failed: {data.ErrorCode} {getattr(data, 'Data', None)}")


def wind_data_to_rows(data, requested_fields, codes):
    by_field = {str(field).lower(): values for field, values in zip(data.Fields, data.Data)}
    rows = []
    for row_index, code in enumerate(codes):
        row = {"wind_code": code}
        for standard, wind_field in requested_fields.items():
            values = by_field.get(wind_field.lower())
            row[standard] = values[row_index] if values is not None and row_index < len(values) else None
        rows.append(row)
    return rows


def fetch_wss_rows(w, codes, requested_fields, options, label, batch_size=MARKET_BATCH_SIZE):
    rows = []
    fields_text = ",".join(requested_fields.values())
    for start in range(0, len(codes), batch_size):
        batch = codes[start:start + batch_size]
        data = w.wss(batch, fields_text, options)
        ensure_wind_ok(data, label)
        rows.extend(wind_data_to_rows(data, requested_fields, batch))
    return pd.DataFrame(rows)


def previous_trading_day(value):
    result = value - timedelta(days=1)
    while result.weekday() >= 5:
        result -= timedelta(days=1)
    return result


def normalized_midpoint(lower, upper):
    lower = pd.to_numeric(lower, errors="coerce")
    upper = pd.to_numeric(upper, errors="coerce")
    return pd.concat([lower, upper], axis=1).mean(axis=1)


def compare_with_notice(row):
    actual = row.get("net_profit")
    lower = row.get("notice_net_profit_min")
    upper = row.get("notice_net_profit_max")
    if pd.isna(lower) and pd.isna(upper):
        return "无业绩预告"
    if pd.notna(lower) and pd.notna(upper) and lower > upper:
        lower, upper = upper, lower
    if pd.notna(actual) and pd.notna(upper) and actual > upper:
        return "超预告上限"
    if pd.notna(actual) and pd.notna(lower) and actual < lower:
        return "低于预告下限"
    if pd.notna(actual):
        return "符合预告区间"
    return "无法比较"


def fetch_expectation_enrichment(w, df, rpt_date):
    codes = df["wind_code"].drop_duplicates().tolist()
    if not codes:
        return df

    notice = fetch_wss_rows(
        w,
        codes,
        PROFIT_NOTICE_FIELDS,
        f"rptDate={yyyymmdd(rpt_date)}",
        "wss profit notice comparison",
    )
    notice["notice_rpt_date"] = pd.to_datetime(notice["notice_rpt_date"], errors="coerce")
    valid_notice = notice["notice_rpt_date"].dt.date == rpt_date
    notice_columns = [column for column in PROFIT_NOTICE_FIELDS if column != "notice_rpt_date"]
    notice.loc[~valid_notice, notice_columns] = pd.NA
    notice["notice_date"] = pd.to_datetime(notice["notice_date"], errors="coerce").dt.strftime("%Y-%m-%d")
    notice = notice.drop(columns=["notice_rpt_date"])
    df = df.merge(notice, on="wind_code", how="left")

    consensus_frames = []
    announcement_dates = pd.to_datetime(df["ann_date"], errors="coerce").dt.date
    df["consensus_asof"] = announcement_dates.map(previous_trading_day)
    for asof_date, group in df.groupby("consensus_asof", dropna=True):
        group_codes = group["wind_code"].drop_duplicates().tolist()
        options = (
            f"tradeDate={yyyymmdd(asof_date)};"
            f"year={rpt_date.year};westPeriod=180;unit=1"
        )
        frame = fetch_wss_rows(
            w,
            group_codes,
            CONSENSUS_FIELDS,
            options,
            "wss pre-disclosure consensus",
        )
        frame["consensus_asof"] = asof_date.strftime("%Y-%m-%d")
        consensus_frames.append(frame)
    df["consensus_asof"] = df["consensus_asof"].map(
        lambda value: value.strftime("%Y-%m-%d") if pd.notna(value) else None
    )
    if consensus_frames:
        consensus = pd.concat(consensus_frames, ignore_index=True)
        df = df.merge(consensus, on=["wind_code", "consensus_asof"], how="left")

    numeric_columns = [
        "net_profit",
        "net_profit_yoy",
        "revenue",
        "notice_net_profit_min",
        "notice_net_profit_max",
        "notice_yoy_min",
        "notice_yoy_max",
        *CONSENSUS_FIELDS,
    ]
    for column in numeric_columns:
        if column in df.columns:
            df[column] = pd.to_numeric(df[column], errors="coerce")

    df["notice_net_profit_mid"] = normalized_midpoint(
        df["notice_net_profit_min"], df["notice_net_profit_max"]
    )
    df["notice_yoy_mid"] = normalized_midpoint(df["notice_yoy_min"], df["notice_yoy_max"])
    valid_mid = df["notice_net_profit_mid"].notna() & (df["notice_net_profit_mid"] != 0)
    df["notice_vs_mid_pct"] = pd.NA
    df.loc[valid_mid, "notice_vs_mid_pct"] = (
        (df.loc[valid_mid, "net_profit"] - df.loc[valid_mid, "notice_net_profit_mid"])
        / df.loc[valid_mid, "notice_net_profit_mid"].abs()
        * 100
    )
    df["notice_yoy_diff_ppt"] = df["net_profit_yoy"] - df["notice_yoy_mid"]
    df["notice_comparison"] = df.apply(compare_with_notice, axis=1)

    df["h1_np_completion_pct"] = pd.NA
    valid_np_consensus = df["consensus_net_profit_fy"].notna() & (df["consensus_net_profit_fy"] > 0)
    df.loc[valid_np_consensus, "h1_np_completion_pct"] = (
        df.loc[valid_np_consensus, "net_profit"]
        / df.loc[valid_np_consensus, "consensus_net_profit_fy"]
        * 100
    )
    return df


def fetch_a_share_codes(w):
    data = w.wset("sectorconstituent", f"sectorid={A_SHARE_SECTOR_ID};field=wind_code,sec_name")
    ensure_wind_ok(data, "wset sectorconstituent")
    by_field = {str(field).lower(): values for field, values in zip(data.Fields, data.Data)}
    codes = by_field.get("wind_code") or []
    names = by_field.get("sec_name") or codes
    rows = []
    seen = set()
    for code, name in zip(codes, names):
        if not code or code in seen:
            continue
        seen.add(code)
        rows.append((str(code), name or str(code)))
    if not rows:
        raise RuntimeError(f"Wind 全 A 成分返回字段异常: {data.Fields}")
    return rows


def fetch_formal_reports(w, query_date, lookback_days, rpt_date):
    constituents = fetch_a_share_codes(w)
    codes = [item[0] for item in constituents]
    code_to_name = dict(constituents)
    fields_text = ",".join(REPORT_FIELDS.values())
    options = (
        f"tradeDate={yyyymmdd(query_date)};"
        f"rptDate={yyyymmdd(rpt_date)};"
        "rptType=1"
    )
    rows = []
    # 大批次可将全 A 扫描压缩到约 7 次请求，降低 Wind 连续请求中断的概率。
    for start in range(0, len(codes), REPORT_BATCH_SIZE):
        batch = codes[start:start + REPORT_BATCH_SIZE]
        print(f"拉取正式半年报 WSS: {start + 1}-{start + len(batch)} / {len(codes)}")
        data = w.wss(batch, fields_text, options)
        ensure_wind_ok(data, "wss formal half-year report")
        batch_rows = wind_data_to_rows(data, REPORT_FIELDS, batch)
        for row in batch_rows:
            row["sec_name"] = code_to_name.get(row["wind_code"], row["wind_code"])
        rows.extend(batch_rows)

    df = pd.DataFrame(rows)
    if df.empty:
        return df, options

    df["ann_date"] = pd.to_datetime(df["ann_date"], errors="coerce")
    start_date = query_date - timedelta(days=lookback_days - 1)
    ann_dates = df["ann_date"].dt.date
    df = df[(ann_dates >= start_date) & (ann_dates <= query_date)].copy()
    df["ann_date"] = df["ann_date"].dt.strftime("%Y-%m-%d")
    df["rpt_date"] = rpt_date.strftime("%Y-%m-%d")
    return df, options


def calculate_q2_enrichment(w, df, query_date, rpt_date):
    """用半年累计值减一季度累计值，计算当年及上年 Q2 单季度值和同比。"""
    codes = df["wind_code"].drop_duplicates().tolist()
    if not codes:
        return df

    periods = {
        "current_q1": date(rpt_date.year, 3, 31),
        "prior_h1": date(rpt_date.year - 1, 6, 30),
        "prior_q1": date(rpt_date.year - 1, 3, 31),
    }
    for prefix, period in periods.items():
        values = fetch_wss_rows(
            w,
            codes,
            QUARTER_CALC_FIELDS,
            (
                f"tradeDate={yyyymmdd(query_date)};"
                f"rptDate={yyyymmdd(period)};"
                "rptType=1"
            ),
            f"wss {prefix} cumulative report",
            batch_size=REPORT_BATCH_SIZE,
        )
        values = values.rename(
            columns={column: f"{prefix}_{column}" for column in QUARTER_CALC_FIELDS}
        )
        df = df.merge(values, on="wind_code", how="left")

    helper_columns = []
    for metric in QUARTER_CALC_FIELDS:
        current_h1 = pd.to_numeric(df[metric], errors="coerce")
        current_q1_column = f"current_q1_{metric}"
        prior_h1_column = f"prior_h1_{metric}"
        prior_q1_column = f"prior_q1_{metric}"
        helper_columns.extend([current_q1_column, prior_h1_column, prior_q1_column])

        current_q1 = pd.to_numeric(df[current_q1_column], errors="coerce")
        prior_h1 = pd.to_numeric(df[prior_h1_column], errors="coerce")
        prior_q1 = pd.to_numeric(df[prior_q1_column], errors="coerce")
        q2_column = f"q2_{metric}"
        q2_yoy_column = f"{q2_column}_yoy"
        df[q2_column] = current_h1 - current_q1
        prior_q2 = prior_h1 - prior_q1
        valid_comparison = df[q2_column].notna() & prior_q2.notna() & (prior_q2 != 0)
        df[q2_yoy_column] = pd.NA
        df.loc[valid_comparison, q2_yoy_column] = (
            (df.loc[valid_comparison, q2_column] - prior_q2.loc[valid_comparison])
            / prior_q2.loc[valid_comparison].abs()
            * 100
        )

    return df.drop(columns=helper_columns)


def fetch_market_enrichment(w, codes, trade_date):
    if not codes:
        return pd.DataFrame(columns=["wind_code", *MARKET_FIELDS])
    result = pd.DataFrame({"wind_code": codes})
    for column, (field, option_template) in MARKET_FIELDS.items():
        values = []
        options = option_template.format(trade_date=yyyymmdd(trade_date))
        for start in range(0, len(codes), MARKET_BATCH_SIZE):
            batch = codes[start:start + MARKET_BATCH_SIZE]
            data = w.wss(batch, field, options)
            ensure_wind_ok(data, f"wss {field}")
            field_values = data.Data[0] if data.Data else []
            values.extend(field_values[:len(batch)])
            if len(field_values) < len(batch):
                values.extend([None] * (len(batch) - len(field_values)))
        result[column] = values
    return result


def build_table(w, query_date, lookback_days, rpt_date):
    df, source_options = fetch_formal_reports(w, query_date, lookback_days, rpt_date)
    if df.empty:
        return pd.DataFrame(columns=DISPLAY_COLUMNS), source_options

    df = calculate_q2_enrichment(w, df, query_date, rpt_date)
    df = fetch_expectation_enrichment(w, df, rpt_date)

    number_columns = [
        *MONEY_COLUMNS,
        *PERCENT_COLUMNS,
        "consensus_np_institutions",
        "pe_ttm",
    ]
    for column in number_columns:
        if column in df.columns:
            df[column] = pd.to_numeric(df[column], errors="coerce")

    codes = df["wind_code"].drop_duplicates().tolist()
    enrichment = fetch_market_enrichment(w, codes, query_date)
    df = df.merge(enrichment, on="wind_code", how="left")

    for column in DISPLAY_COLUMNS:
        if column not in df.columns:
            df[column] = pd.NA
    df = df[DISPLAY_COLUMNS].sort_values(
        ["ann_date", "net_profit_yoy", "wind_code"],
        ascending=[False, False, True],
        na_position="last",
    )
    return df, source_options


def prepare_output_table(df):
    shown = df.copy()
    for column in MONEY_COLUMNS:
        if column in shown.columns:
            shown[column] = (pd.to_numeric(shown[column], errors="coerce") / 100_000_000).round(2)
    for column in PERCENT_COLUMNS + ["pe_ttm"]:
        if column in shown.columns:
            shown[column] = pd.to_numeric(shown[column], errors="coerce").round(2)
    for column in ["consensus_np_institutions"]:
        if column in shown.columns:
            shown[column] = pd.to_numeric(shown[column], errors="coerce").round().astype("Int64")
    return shown.rename(columns=OUTPUT_COLUMN_NAMES)


def format_for_display(df):
    shown = prepare_output_table(df)
    for column in [OUTPUT_COLUMN_NAMES[item] for item in MONEY_COLUMNS]:
        shown[column] = shown[column].map(lambda value: "" if pd.isna(value) else f"{value:,.2f}")
    for column in [OUTPUT_COLUMN_NAMES[item] for item in PERCENT_COLUMNS] + ["PE_TTM"]:
        shown[column] = shown[column].map(lambda value: "" if pd.isna(value) else f"{value:.2f}")
    for column in ["净利润预测机构数"]:
        shown[column] = shown[column].map(lambda value: "" if pd.isna(value) else str(int(value)))
    terminal_names = [OUTPUT_COLUMN_NAMES[column] for column in TERMINAL_COLUMNS]
    return shown[terminal_names]


def save_outputs(df, query_date, rpt_date):
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    stem = f"A股正式半年报跟踪_{rpt_date.year}_{query_date:%Y%m%d}"
    csv_path = OUTPUT_DIR / f"{stem}.csv"
    xlsx_path = OUTPUT_DIR / f"{stem}.xlsx"
    output_df = prepare_output_table(df)
    output_df.to_csv(csv_path, index=False, encoding="utf-8-sig")

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        output_df.to_excel(writer, index=False, sheet_name="正式半年报")
        notes = pd.DataFrame(
            [
                ["业绩预告比较", "正式半年报归母净利润与同一 06-30 报告期的业绩预告区间比较。"],
                ["超预告上限", "实际归母净利润高于业绩预告上限。"],
                ["一致预期口径", "取正式披露前一交易日、近 180 天机构最新年度预测的均值。"],
                ["半年完成率", "半年实际值 / 全年一致预期。该指标用于观察进度，不等同于半年报超预期。"],
                ["Q2 单季度", "当年半年累计值减当年一季度累计值；Q2 同比使用上年同期同口径单季度值计算。"],
                ["机构数", "对对应年度归母净利润给出预测的机构家数；家数越少，参考时越需谨慎。"],
                ["空值", "Wind 无同期业绩预告、无有效券商预测，或该行业不适用对应指标。"],
            ],
            columns=["项目", "说明"],
        )
        notes.to_excel(writer, index=False, sheet_name="口径说明")
        ws = writer.book["正式半年报"]
        ws.freeze_panes = "E2"
        ws.auto_filter.ref = ws.dimensions
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)

        for column_index, column_name in enumerate(output_df.columns, start=1):
            values = output_df[column_name].astype("string").fillna("").head(200)
            width = min(max([len(str(column_name)) + 4, *[len(value) + 2 for value in values]]), 28)
            ws.column_dimensions[get_column_letter(column_index)].width = width

        alternate_fill = PatternFill("solid", fgColor="EAF2F8")
        for row_index in range(2, ws.max_row + 1):
            if row_index % 2 == 0:
                for cell in ws[row_index]:
                    cell.fill = alternate_fill

        positive_fill = PatternFill("solid", fgColor="E2F0D9")
        negative_fill = PatternFill("solid", fgColor="FCE4D6")
        for column_name in [
            "H1营收同比(%)",
            "H1归母净利润同比(%)",
            "Q2营收同比(%)",
            "Q2归母净利润同比(%)",
        ]:
            column_index = output_df.columns.get_loc(column_name) + 1
            column_letter = get_column_letter(column_index)
            cell_range = f"{column_letter}2:{column_letter}{max(ws.max_row, 2)}"
            ws.conditional_formatting.add(
                cell_range, CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=positive_fill)
            )
            ws.conditional_formatting.add(
                cell_range, CellIsRule(operator="lessThan", formula=["0"], fill=negative_fill)
            )
        notes_ws = writer.book["口径说明"]
        notes_ws.freeze_panes = "A2"
        notes_ws.column_dimensions["A"].width = 22
        notes_ws.column_dimensions["B"].width = 90
        for cell in notes_ws[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
    return csv_path, xlsx_path


def print_summary(df):
    if df.empty:
        return
    yoy = pd.to_numeric(df["net_profit_yoy"], errors="coerce")
    profit = pd.to_numeric(df["net_profit"], errors="coerce")
    print(
        f"摘要: 净利润同比增长 {int((yoy > 0).sum())} 家"
        f" | 同比下降 {int((yoy < 0).sum())} 家"
        f" | 亏损 {int((profit < 0).sum())} 家"
    )
    notice = df["notice_comparison"]
    consensus_count = int(pd.to_numeric(df["consensus_net_profit_fy"], errors="coerce").notna().sum())
    print(
        f"比较: 有业绩预告 {int((notice != '无业绩预告').sum())} 家"
        f" | 超预告上限 {int((notice == '超预告上限').sum())} 家"
        f" | 有全年一致预期 {consensus_count} 家"
    )


def main():
    parser = argparse.ArgumentParser(description="跟踪最近新披露的 A 股正式半年报及核心指标。")
    parser.add_argument("--date", default=date.today().strftime("%Y-%m-%d"), help="查询日期，默认今天")
    parser.add_argument(
        "--lookback-days",
        type=int,
        default=None,
        help="披露日回看天数；默认覆盖上一个工作日至查询日",
    )
    parser.add_argument("--rpt-date", default=None, help="报告期，默认自动取最近半年末，如 2026-06-30")
    parser.add_argument("--no-save", action="store_true", help="只在终端显示，不保存 CSV/XLSX")
    args = parser.parse_args()

    query_date = parse_yyyy_mm_dd(args.date)
    rpt_date = parse_yyyy_mm_dd(args.rpt_date) if args.rpt_date else default_rpt_date(query_date)
    lookback_days = default_lookback_days(query_date) if args.lookback_days is None else args.lookback_days
    if lookback_days < 1:
        parser.error("--lookback-days 必须大于等于 1")
    if (rpt_date.month, rpt_date.day) != (6, 30):
        parser.error("这是半年报跟踪脚本，--rpt-date 必须是某年 06-30")

    from WindPy import w

    w.start()
    try:
        df, source_options = build_table(w, query_date, lookback_days, rpt_date)
    finally:
        w.close()

    start_date = query_date - timedelta(days=lookback_days - 1)
    print("Wind 数据源: wss 正式财务报告")
    print(f"Wind 参数: {source_options}")
    print(
        f"查询披露日: {start_date:%Y-%m-%d} 至 {query_date:%Y-%m-%d}"
        f" | 报告期: {rpt_date:%Y-%m-%d} | 公司数: {len(df)}"
    )
    print_summary(df)
    if df.empty:
        print("没有找到符合条件的 A 股正式半年报。")
    else:
        print(format_for_display(df).to_string(index=False, max_colwidth=24))

    if not args.no_save:
        csv_path, xlsx_path = save_outputs(df, query_date, rpt_date)
        print(f"CSV: {csv_path}")
        print(f"Excel: {xlsx_path}")


if __name__ == "__main__":
    main()
