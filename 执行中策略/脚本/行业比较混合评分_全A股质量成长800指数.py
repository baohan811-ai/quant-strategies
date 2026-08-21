"""
行业比较＋混合评分的全 A 股质量成长 800 指数（单文件版）。

目标：
1. 在全部 A 股中，从基本面、估值、业绩稳定性、成长性、周期阶段、
   竞争格局、动量和低规模八个维度评分，选总分最高的 800 只股票；
2. 组合行业权重以沪深 300 的万得一级行业权重为锚；
3. 行业内以自由流通市值为主分配权重，并对高分股票适度倾斜；
4. 行业吸引力及行业动量决定相对沪深 300 的权重偏离，偏离幅度受限；
5. 指标在万得一级行业内缩尾和标准化；换样使用 20% 排名缓冲，
   不按行业预先分配成分股数量。

数据原则：财报只使用 available_date <= 截止日的数据，避免未来函数。
本脚本优先读取项目现有 SQLite；沪深 300 当前成分及官方权重默认从 Wind
获取，也可用 --benchmark-weights 传入本地 CSV/XLSX（industry,target_weight）。
指数构建、历史回测、绩效计算和文件输出均包含在本文件中。
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sqlite3
from dataclasses import asdict, dataclass, replace
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd


SCRIPT_DIR = Path(__file__).resolve().parent
BASE_DIR = SCRIPT_DIR.parent
FUNDAMENTAL_DB_PATH = BASE_DIR / "缓存" / "全部A股_基础基本面.sqlite3"
MARKET_DB_PATH = BASE_DIR / "缓存" / "本地行情数据库.sqlite3"
OUTPUT_DIR = BASE_DIR / "输出" / "全A股质量成长800"
LEGACY_OUTPUT_DIR = BASE_DIR / "输出"
PROFIT_NOTICE_OUTPUT_DIR = BASE_DIR / "输出" / "业绩预告跟踪"
PERFORMANCE_EXPRESS_OUTPUT_DIR = BASE_DIR / "输出" / "业绩快报跟踪"
BUILD_OUTPUT_PATTERN = re.compile(
    r"^全A股质量成长800_行业比较混合评分_(\d{8})\.xlsx$"
)

ALL_A_UNIVERSE_NAME = "全部A股"
INDUSTRY_SYSTEM = "wind_level1"
BENCHMARK_CODE = "000300.SH"


@dataclass(frozen=True)
class StrategyConfig:
    max_constituents: int = 800
    rebalance_frequency: str = "M"
    industry_classification: str = "wind_level1"
    min_data_coverage: float = 0.50
    min_market_cap: float = 0.0
    exclude_st: bool = True

    fundamental_weight: float = 0.25
    valuation_weight: float = 0.20
    stability_weight: float = 0.15
    growth_weight: float = 0.20
    cycle_weight: float = 0.10
    competition_weight: float = 0.10

    history_quarters: int = 12
    exceptional_score_quantile: float = 0.90
    industry_score_tilt: float = 0.15
    max_industry_relative_deviation: float = 0.25
    max_off_benchmark_industry_weight: float = 0.01
    # 评分倾斜强度：0.50 时约在正负 1.39 个标准差触及 2.0/0.5 倍边界。
    stock_score_tilt: float = 0.50
    max_stock_weight: float = 0.05
    valuation_min_coverage: float = 0.90
    stale_valuation_warning_days: int = 45
    winsorize_lower_quantile: float = 0.05
    winsorize_upper_quantile: float = 0.95
    constituent_buffer_ratio: float = 0.20
    scoring_mode: str = "industry_relative"
    stock_factor_scope: str = "blended"
    constituent_selection_mode: str = "global_top"
    industry_quota_min_score: float = 45.0
    industry_target_mode: str = "selected_high_score_mass"
    industry_between_tilt_strength: float = 1.0
    hybrid_momentum_weight: float = 0.10
    hybrid_low_size_weight: float = 0.05
    industry_momentum_weight: float = 0.10
    industry_history_months: int = 36
    industry_history_min_observations: int = 12
    earnings_guidance_enabled: bool = True
    earnings_notice_confidence: float = 0.40
    earnings_express_confidence: float = 0.70


SCORE_WEIGHTS = {
    "基本面得分": "fundamental_weight",
    "估值得分": "valuation_weight",
    "稳定性得分": "stability_weight",
    "成长性得分": "growth_weight",
    "周期阶段得分": "cycle_weight",
    "竞争格局得分": "competition_weight",
}


def validate_config(config: StrategyConfig) -> None:
    total = sum(getattr(config, attr) for attr in SCORE_WEIGHTS.values())
    if not math.isclose(total, 1.0, abs_tol=1e-9):
        raise ValueError(f"六维评分权重合计必须为 1，当前为 {total:.6f}")
    if config.max_constituents <= 0:
        raise ValueError("max_constituents 必须大于 0")
    if config.rebalance_frequency != "M":
        raise ValueError("当前确认的调仓频率为月度（M）")
    if config.industry_classification != INDUSTRY_SYSTEM:
        raise ValueError("当前确认的行业分类为万得一级（wind_level1）")
    if not 0 <= config.winsorize_lower_quantile < config.winsorize_upper_quantile <= 1:
        raise ValueError("缩尾分位数必须满足 0 <= lower < upper <= 1")
    if not 0 <= config.constituent_buffer_ratio < 0.5:
        raise ValueError("成分缓冲比例必须在 [0, 0.5) 内")
    if config.scoring_mode not in {"industry_relative", "hybrid"}:
        raise ValueError("scoring_mode 必须为 industry_relative 或 hybrid")
    if config.stock_factor_scope not in {"industry_only", "blended"}:
        raise ValueError("stock_factor_scope 必须为 industry_only 或 blended")
    if config.constituent_selection_mode not in {"industry_quota", "global_top"}:
        raise ValueError("constituent_selection_mode 必须为 industry_quota 或 global_top")
    if not 0 <= config.industry_quota_min_score <= 100:
        raise ValueError("行业名额最低总分必须在 [0, 100] 内")
    if config.industry_target_mode not in {"selected_high_score_mass", "industry_score"}:
        raise ValueError("industry_target_mode 必须为 selected_high_score_mass 或 industry_score")
    if config.hybrid_momentum_weight < 0 or config.hybrid_low_size_weight < 0:
        raise ValueError("动量和低规模因子权重不能为负")
    if config.hybrid_momentum_weight + config.hybrid_low_size_weight >= 1:
        raise ValueError("动量和低规模因子权重合计必须小于1")
    if not 0 <= config.industry_momentum_weight < 1:
        raise ValueError("行业动量权重必须在 [0, 1) 内")
    if config.stock_score_tilt < 0:
        raise ValueError("个股评分倾斜强度不能为负")
    if config.industry_history_months < config.industry_history_min_observations:
        raise ValueError("行业历史窗口不能短于最少有效观察期")
    if config.industry_history_min_observations < 2:
        raise ValueError("行业历史最少有效观察期必须大于等于2")
    if not 0 <= config.earnings_notice_confidence <= 1:
        raise ValueError("业绩预告置信权重必须在 [0, 1] 内")
    if not 0 <= config.earnings_express_confidence <= 1:
        raise ValueError("业绩快报置信权重必须在 [0, 1] 内")
    if config.earnings_express_confidence < config.earnings_notice_confidence:
        raise ValueError("业绩快报置信权重不应低于业绩预告")


def normalize_date(value: str | date | pd.Timestamp) -> str:
    return pd.Timestamp(value).strftime("%Y-%m-%d")


def finite_numeric(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce").replace([np.inf, -np.inf], np.nan)


def normalize_percent(series: pd.Series) -> pd.Series:
    """Wind 百分比字段可能为 15 或 0.15，统一为小数。"""
    values = finite_numeric(series)
    return values.where(values.abs() <= 2, values / 100.0)


def grouped_percentile(
    frame: pd.DataFrame,
    column: str,
    higher_is_better: bool = True,
    neutral: float = 0.5,
    lower_quantile: float = 0.05,
    upper_quantile: float = 0.95,
) -> pd.Series:
    """在万得一级行业内先缩尾，再转换为稳健的百分位分数。"""
    values = finite_numeric(frame[column])
    valid_group = frame["行业"].fillna("未分类").astype(str)
    grouped = values.groupby(valid_group)
    lower = grouped.transform(lambda value: value.quantile(lower_quantile))
    upper = grouped.transform(lambda value: value.quantile(upper_quantile))
    winsorized = values.clip(lower=lower, upper=upper)
    ranked = winsorized.groupby(valid_group).rank(pct=True, method="average")
    score = ranked if higher_is_better else 1.0 - ranked
    return score.fillna(neutral).clip(0.0, 1.0)


def positive_valuation_percentile(
    frame: pd.DataFrame,
    column: str,
    config: StrategyConfig,
) -> pd.Series:
    temp = frame.copy()
    temp[column] = finite_numeric(temp[column]).where(lambda value: value > 0)
    return grouped_percentile(
        temp,
        column,
        higher_is_better=False,
        neutral=0.35,
        lower_quantile=config.winsorize_lower_quantile,
        upper_quantile=config.winsorize_upper_quantile,
    )


def global_percentile(
    frame: pd.DataFrame,
    column: str,
    higher_is_better: bool = True,
    neutral: float = 0.5,
    lower_quantile: float = 0.05,
    upper_quantile: float = 0.95,
    positive_only: bool = False,
) -> pd.Series:
    """全市场缩尾百分位，用于保留行业之间的绝对差异。"""
    values = finite_numeric(frame[column])
    if positive_only:
        values = values.where(values > 0)
    lower = values.quantile(lower_quantile)
    upper = values.quantile(upper_quantile)
    winsorized = values.clip(lower=lower, upper=upper)
    ranked = winsorized.rank(pct=True, method="average")
    score = ranked if higher_is_better else 1.0 - ranked
    return score.fillna(neutral).clip(0.0, 1.0)


def blend_scores(
    industry_relative: pd.Series,
    global_absolute: pd.Series,
    industry_share: float,
) -> pd.Series:
    return industry_share * industry_relative + (1 - industry_share) * global_absolute


def load_universe(as_of: str) -> pd.DataFrame:
    with sqlite3.connect(MARKET_DB_PATH) as conn:
        snapshot_date = conn.execute(
            """
            SELECT MAX(snapshot_date)
            FROM universe_constituents_snapshot
            WHERE universe_name = ? AND snapshot_date <= ?
            """,
            [ALL_A_UNIVERSE_NAME, as_of],
        ).fetchone()[0]
        if snapshot_date:
            universe = pd.read_sql_query(
                """
                SELECT wind_code, sec_name
                FROM universe_constituents_snapshot
                WHERE universe_name = ? AND snapshot_date = ?
                ORDER BY wind_code
                """,
                conn,
                params=[ALL_A_UNIVERSE_NAME, snapshot_date],
            )
            universe_source = f"历史股票池截面:{snapshot_date}"
        else:
            universe = pd.read_sql_query(
                """
                SELECT wind_code, sec_name
                FROM stock_universe
                WHERE universe_name = ?
                ORDER BY wind_code
                """,
                conn,
                params=[ALL_A_UNIVERSE_NAME],
            )
            universe_source = "当前股票池"
    if universe.empty:
        raise RuntimeError(f"截止 {as_of} 没有可用的全A股股票池截面")
    universe["截止日"] = as_of
    universe["universe_source"] = universe_source
    return universe.drop_duplicates("wind_code", keep="last")


def load_industries(codes: list[str], as_of: str) -> pd.DataFrame:
    """历史截面优先；没有历史截面时使用当前行业映射并明确标记来源。"""
    params = [INDUSTRY_SYSTEM, ALL_A_UNIVERSE_NAME, as_of]
    with sqlite3.connect(MARKET_DB_PATH) as conn:
        snapshot_date = conn.execute(
            """
            SELECT MAX(snapshot_date)
            FROM stock_industry_snapshot
            WHERE classification_system = ? AND universe_name = ?
              AND snapshot_date <= ?
            """,
            params,
        ).fetchone()[0]
        if snapshot_date:
            industry = pd.read_sql_query(
                """
                SELECT wind_code, sec_name, industry_level1 AS industry
                FROM stock_industry_snapshot
                WHERE classification_system = ? AND universe_name = ?
                  AND snapshot_date = ?
                """,
                conn,
                params=[INDUSTRY_SYSTEM, ALL_A_UNIVERSE_NAME, snapshot_date],
            )
            source = f"历史行业截面:{snapshot_date}"
        else:
            industry = pd.read_sql_query(
                """
                SELECT wind_code, sec_name, industry_level1 AS industry
                FROM stock_industry
                WHERE classification_system = ?
                """,
                conn,
                params=[INDUSTRY_SYSTEM],
            )
            source = "当前行业映射"

    industry = industry[industry["wind_code"].isin(codes)].copy()
    industry = industry.drop_duplicates("wind_code", keep="last")
    industry["industry_source"] = source
    return industry


def load_latest_fundamentals(codes: list[str], as_of: str) -> pd.DataFrame:
    columns = [
        "trade_date", "wind_code", "pe_ttm", "pb_lf", "ps_ttm",
        "dividend_yield", "roe_ttm", "debt_to_assets",
        "revenue_yoy_qfa", "netprofit_yoy_qfa",
        "gross_profit_margin_qfa", "net_profit_margin_qfa",
        "qfa_report_period", "qfa_available_date",
    ]
    with sqlite3.connect(FUNDAMENTAL_DB_PATH) as conn:
        data = pd.read_sql_query(
            f"SELECT {','.join(columns)} FROM fundamentals WHERE trade_date <= ? ORDER BY wind_code, trade_date",
            conn,
            params=[as_of],
        )
    data = data[data["wind_code"].isin(codes)]
    if data.empty:
        raise RuntimeError(f"截止 {as_of} 没有可用基础面截面")
    latest = data.groupby("wind_code", as_index=False).last()
    latest = latest.rename(columns={"trade_date": "fundamental_source_date"})
    return latest


def choose_valuation_date(as_of: str, universe_size: int, config: StrategyConfig) -> str:
    minimum = max(1, int(universe_size * config.valuation_min_coverage))
    with sqlite3.connect(FUNDAMENTAL_DB_PATH) as conn:
        row = conn.execute(
            """
            SELECT trade_date
            FROM daily_valuation
            WHERE trade_date <= ? AND COALESCE(free_float_mkt_cap, mkt_cap_ard) > 0
            GROUP BY trade_date
            HAVING COUNT(*) >= ?
            ORDER BY trade_date DESC
            LIMIT 1
            """,
            [as_of, minimum],
        ).fetchone()
    if not row:
        raise RuntimeError(
            f"日频市值库在 {as_of} 前没有覆盖至少 {minimum} 只股票的完整截面"
        )
    return str(row[0])


def load_valuation(codes: list[str], as_of: str, config: StrategyConfig) -> tuple[pd.DataFrame, str]:
    valuation_date = choose_valuation_date(as_of, len(codes), config)
    with sqlite3.connect(FUNDAMENTAL_DB_PATH) as conn:
        valuation = pd.read_sql_query(
            """
            SELECT wind_code, mkt_cap_ard, free_float_mkt_cap
            FROM daily_valuation
            WHERE trade_date = ?
            """,
            conn,
            params=[valuation_date],
        )
    valuation = valuation[valuation["wind_code"].isin(codes)].copy()
    valuation["weighting_market_cap"] = valuation["free_float_mkt_cap"].fillna(
        valuation["mkt_cap_ard"]
    )
    valuation["valuation_source_date"] = valuation_date
    return valuation, valuation_date


def load_momentum_features(codes: list[str], as_of: str) -> pd.DataFrame:
    """计算12-1个月与6-1个月价格动量，跳过最近一个月。"""
    as_of_date = pd.Timestamp(as_of)
    targets = {
        "close_1m": as_of_date - pd.DateOffset(months=1),
        "close_6m": as_of_date - pd.DateOffset(months=6),
        "close_12m": as_of_date - pd.DateOffset(months=12),
    }
    code_set = set(codes)
    result = pd.DataFrame({"wind_code": codes})
    with sqlite3.connect(MARKET_DB_PATH) as conn:
        for column, target in targets.items():
            start = target - pd.Timedelta(days=20)
            prices = pd.read_sql_query(
                """
                SELECT trade_date, wind_code, close
                FROM daily_prices
                WHERE adjusted = 'F' AND trade_date >= ? AND trade_date <= ?
                ORDER BY wind_code, trade_date
                """,
                conn,
                params=[start.strftime("%Y-%m-%d"), target.strftime("%Y-%m-%d")],
            )
            prices = prices[prices["wind_code"].isin(code_set)]
            latest = prices.groupby("wind_code", as_index=False).last()
            latest = latest.rename(
                columns={"close": column, "trade_date": f"{column}_date"}
            )
            result = result.merge(
                latest[["wind_code", column, f"{column}_date"]],
                on="wind_code",
                how="left",
            )
    result["momentum_12_1"] = (
        finite_numeric(result["close_1m"]) / finite_numeric(result["close_12m"]) - 1
    )
    result["momentum_6_1"] = (
        finite_numeric(result["close_1m"]) / finite_numeric(result["close_6m"]) - 1
    )
    result["momentum_data_coverage"] = result[["momentum_12_1", "momentum_6_1"]].notna().mean(axis=1)
    return result


def load_report_history(codes: list[str], as_of: str, quarters: int) -> pd.DataFrame:
    with sqlite3.connect(FUNDAMENTAL_DB_PATH) as conn:
        reports = pd.read_sql_query(
            """
            SELECT rpt_date, wind_code, available_date,
                   revenue_yoy_qfa, netprofit_yoy_qfa,
                   gross_profit_margin_qfa, net_profit_margin_qfa
            FROM financial_reports
            WHERE available_date IS NOT NULL AND available_date <= ?
            ORDER BY wind_code, rpt_date
            """,
            conn,
            params=[as_of],
        )
    reports = reports[reports["wind_code"].isin(codes)].copy()
    if reports.empty:
        return reports
    return reports.groupby("wind_code", group_keys=False).tail(quarters)


def report_history_features(reports: pd.DataFrame) -> pd.DataFrame:
    output_columns = [
        "wind_code", "report_count", "growth_positive_ratio", "growth_volatility",
        "margin_volatility", "revenue_acceleration", "profit_acceleration",
    ]
    if reports.empty:
        return pd.DataFrame(columns=output_columns)

    for column in [
        "revenue_yoy_qfa", "netprofit_yoy_qfa",
        "gross_profit_margin_qfa", "net_profit_margin_qfa",
    ]:
        reports[column] = normalize_percent(reports[column])

    rows = []
    for code, group in reports.groupby("wind_code", sort=False):
        group = group.sort_values("rpt_date")
        revenue = group["revenue_yoy_qfa"]
        profit = group["netprofit_yoy_qfa"]
        gross_margin = group["gross_profit_margin_qfa"]
        net_margin = group["net_profit_margin_qfa"]
        growth_observations = pd.concat([revenue, profit], axis=0).dropna()
        growth_stds = [value.std() for value in (revenue, profit) if value.notna().sum() >= 3]
        margin_stds = [value.std() for value in (gross_margin, net_margin) if value.notna().sum() >= 3]
        revenue_acceleration = revenue.iloc[-1] - revenue.iloc[-5] if revenue.notna().sum() >= 5 else np.nan
        profit_acceleration = profit.iloc[-1] - profit.iloc[-5] if profit.notna().sum() >= 5 else np.nan
        rows.append({
            "wind_code": code,
            "report_count": int(len(group)),
            "growth_positive_ratio": float((growth_observations > 0).mean()) if len(growth_observations) else np.nan,
            "growth_volatility": float(np.mean(growth_stds)) if growth_stds else np.nan,
            "margin_volatility": float(np.mean(margin_stds)) if margin_stds else np.nan,
            "revenue_acceleration": revenue_acceleration,
            "profit_acceleration": profit_acceleration,
        })
    return pd.DataFrame(rows, columns=output_columns)


_EARNINGS_DISCLOSURE_FILES_IMPORTED = False


def ensure_earnings_disclosure_table(conn: sqlite3.Connection) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS earnings_disclosures (
            wind_code TEXT NOT NULL,
            rpt_date TEXT NOT NULL,
            ann_date TEXT NOT NULL,
            disclosure_type TEXT NOT NULL,
            profit_yoy_min REAL,
            profit_yoy_max REAL,
            profit_yoy_mid REAL,
            disclosure_style TEXT,
            source_file TEXT,
            updated_at TEXT NOT NULL,
            PRIMARY KEY (wind_code, rpt_date, ann_date, disclosure_type)
        )
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_earnings_disclosures_asof
        ON earnings_disclosures (ann_date, wind_code, rpt_date)
        """
    )


def import_earnings_disclosure_files() -> None:
    """把每日预告/快报跟踪文件汇入点时缓存；重复运行采用主键更新。"""
    global _EARNINGS_DISCLOSURE_FILES_IMPORTED
    if _EARNINGS_DISCLOSURE_FILES_IMPORTED:
        return

    sources = [
        (
            PROFIT_NOTICE_OUTPUT_DIR,
            "notice",
            ["同比下限(%)", "预告同比下限(%)"],
            ["同比上限(%)", "预告同比上限(%)"],
            ["同比中值(%)", "预告同比中值(%)"],
        ),
        (
            PERFORMANCE_EXPRESS_OUTPUT_DIR,
            "express",
            ["净利润同比下限(%)"],
            ["净利润同比上限(%)"],
            ["净利润同比(%)", "归母净利润同比(%)", "同比中值(%)"],
        ),
    ]

    def first_column(frame: pd.DataFrame, candidates: list[str]) -> str | None:
        return next((column for column in candidates if column in frame.columns), None)

    rows = []
    for directory, disclosure_type, min_names, max_names, mid_names in sources:
        if not directory.is_dir():
            continue
        for path in sorted(directory.glob("*.csv")):
            try:
                data = pd.read_csv(path, encoding="utf-8-sig")
            except (OSError, UnicodeError, pd.errors.ParserError):
                continue
            code_column = first_column(data, ["代码", "wind_code", "证券代码"])
            report_column = first_column(data, ["报告期", "rpt_date"])
            announce_column = first_column(data, ["披露日", "公告日期", "ann_date"])
            if not code_column or not report_column or not announce_column:
                continue
            min_column = first_column(data, min_names)
            max_column = first_column(data, max_names)
            mid_column = first_column(data, mid_names)
            style_column = first_column(data, ["预告类型", "快报类型", "disclosure_style"])
            for item in data.itertuples(index=False, name=None):
                values = dict(zip(data.columns, item))
                code = values.get(code_column)
                rpt_date = pd.to_datetime(values.get(report_column), errors="coerce")
                ann_date = pd.to_datetime(values.get(announce_column), errors="coerce")
                if pd.isna(code) or pd.isna(rpt_date) or pd.isna(ann_date):
                    continue
                yoy_min = pd.to_numeric(values.get(min_column), errors="coerce") if min_column else np.nan
                yoy_max = pd.to_numeric(values.get(max_column), errors="coerce") if max_column else np.nan
                yoy_mid = pd.to_numeric(values.get(mid_column), errors="coerce") if mid_column else np.nan
                if pd.isna(yoy_mid):
                    available = [value for value in [yoy_min, yoy_max] if pd.notna(value)]
                    yoy_mid = float(np.mean(available)) if available else np.nan
                if pd.isna(yoy_mid):
                    continue
                rows.append((
                    str(code),
                    rpt_date.strftime("%Y-%m-%d"),
                    ann_date.strftime("%Y-%m-%d"),
                    disclosure_type,
                    float(yoy_min) / 100 if pd.notna(yoy_min) else None,
                    float(yoy_max) / 100 if pd.notna(yoy_max) else None,
                    float(yoy_mid) / 100,
                    str(values.get(style_column)) if style_column and pd.notna(values.get(style_column)) else None,
                    str(path),
                    pd.Timestamp.now().isoformat(timespec="seconds"),
                ))
    with sqlite3.connect(FUNDAMENTAL_DB_PATH) as conn:
        ensure_earnings_disclosure_table(conn)
        if rows:
            conn.executemany(
                """
                INSERT INTO earnings_disclosures (
                    wind_code, rpt_date, ann_date, disclosure_type,
                    profit_yoy_min, profit_yoy_max, profit_yoy_mid,
                    disclosure_style, source_file, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(wind_code, rpt_date, ann_date, disclosure_type)
                DO UPDATE SET
                    profit_yoy_min = excluded.profit_yoy_min,
                    profit_yoy_max = excluded.profit_yoy_max,
                    profit_yoy_mid = excluded.profit_yoy_mid,
                    disclosure_style = excluded.disclosure_style,
                    source_file = excluded.source_file,
                    updated_at = excluded.updated_at
                """,
                rows,
            )
        conn.commit()
    _EARNINGS_DISCLOSURE_FILES_IMPORTED = True


def load_earnings_guidance(
    frame: pd.DataFrame,
    as_of: str,
    config: StrategyConfig,
) -> pd.DataFrame:
    """读取公告日不晚于截止日、且报告期晚于最新正式财报的预告或快报。"""
    output = pd.DataFrame({"wind_code": frame["wind_code"].astype(str)})
    if not config.earnings_guidance_enabled:
        return output
    import_earnings_disclosure_files()
    with sqlite3.connect(FUNDAMENTAL_DB_PATH) as conn:
        ensure_earnings_disclosure_table(conn)
        disclosures = pd.read_sql_query(
            """
            SELECT wind_code, rpt_date, ann_date, disclosure_type,
                   profit_yoy_min, profit_yoy_max, profit_yoy_mid,
                   disclosure_style, source_file
            FROM earnings_disclosures
            WHERE ann_date <= ?
            """,
            conn,
            params=[as_of],
        )
    if disclosures.empty:
        return output
    disclosures = disclosures[
        disclosures["wind_code"].isin(set(output["wind_code"]))
    ].copy()
    disclosures["rpt_date_dt"] = pd.to_datetime(disclosures["rpt_date"], errors="coerce")
    disclosures["ann_date_dt"] = pd.to_datetime(disclosures["ann_date"], errors="coerce")
    formal_period = pd.to_datetime(frame.get("qfa_report_period"), errors="coerce")
    formal_lookup = pd.Series(formal_period.to_numpy(), index=frame["wind_code"].astype(str))
    disclosures["formal_period_dt"] = disclosures["wind_code"].map(formal_lookup)
    disclosures = disclosures[
        disclosures["rpt_date_dt"].notna()
        & disclosures["ann_date_dt"].notna()
        & ((disclosures["rpt_date_dt"] - disclosures["ann_date_dt"]).dt.days <= 366)
        & (
            disclosures["formal_period_dt"].isna()
            | (disclosures["rpt_date_dt"] > disclosures["formal_period_dt"])
        )
    ].copy()
    if disclosures.empty:
        return output
    disclosures["priority"] = disclosures["disclosure_type"].map(
        {"notice": 1, "express": 2}
    ).fillna(0)
    disclosures = disclosures.sort_values(
        ["wind_code", "rpt_date_dt", "priority", "ann_date_dt"],
        ascending=[True, False, False, False],
    ).drop_duplicates("wind_code", keep="first")
    disclosures["earnings_guidance_confidence"] = disclosures["disclosure_type"].map({
        "notice": config.earnings_notice_confidence,
        "express": config.earnings_express_confidence,
    }).fillna(0.0)
    disclosures = disclosures.rename(columns={
        "rpt_date": "earnings_guidance_report_period",
        "ann_date": "earnings_guidance_announcement_date",
        "disclosure_type": "earnings_guidance_type",
        "profit_yoy_mid": "earnings_guidance_profit_yoy",
        "disclosure_style": "earnings_guidance_style",
        "source_file": "earnings_guidance_source",
    })
    columns = [
        "wind_code", "earnings_guidance_report_period",
        "earnings_guidance_announcement_date", "earnings_guidance_type",
        "earnings_guidance_profit_yoy", "earnings_guidance_confidence",
        "earnings_guidance_style", "earnings_guidance_source",
    ]
    return output.merge(disclosures[columns], on="wind_code", how="left")


def classify_cycle_stage(row: pd.Series) -> str:
    revenue = row.get("revenue_yoy_qfa")
    profit = row.get("netprofit_yoy_qfa")
    acceleration_values = [
        value for value in [row.get("revenue_acceleration"), row.get("profit_acceleration")]
        if pd.notna(value)
    ]
    acceleration = float(np.mean(acceleration_values)) if acceleration_values else np.nan
    if pd.isna(revenue) and pd.isna(profit):
        return "数据不足"
    revenue = -1.0 if pd.isna(revenue) else revenue
    profit = -1.0 if pd.isna(profit) else profit
    acceleration = 0.0 if pd.isna(acceleration) else acceleration
    if revenue > 0 and profit > 0 and acceleration >= 0:
        return "扩张"
    if acceleration > 0 and (revenue > 0 or profit > 0):
        return "复苏"
    if revenue > 0 or profit > 0:
        return "放缓"
    return "收缩"


def build_industry_history_scores(
    membership: pd.DataFrame,
    as_of: str,
    config: StrategyConfig,
) -> pd.DataFrame:
    """按行业自身历史位置计算行业六维得分，避免商业模式差异主导行业比较。"""
    metric_columns = [
        "pe_ttm", "pb_lf", "ps_ttm", "dividend_yield", "roe_ttm",
        "debt_to_assets", "revenue_yoy_qfa", "netprofit_yoy_qfa",
        "gross_profit_margin_qfa", "net_profit_margin_qfa",
    ]
    history_start = (
        pd.Timestamp(as_of)
        - pd.DateOffset(months=config.industry_history_months + 12)
    ).strftime("%Y-%m-%d")
    with sqlite3.connect(FUNDAMENTAL_DB_PATH) as conn:
        history = pd.read_sql_query(
            f"""
            SELECT trade_date, wind_code, {','.join(metric_columns)}
            FROM fundamentals
            WHERE trade_date >= ? AND trade_date <= ?
            ORDER BY trade_date, wind_code
            """,
            conn,
            params=[history_start, as_of],
        )
    membership = membership[["wind_code", "行业"]].drop_duplicates("wind_code")
    history = history.merge(membership, on="wind_code", how="inner")
    if history.empty:
        return pd.DataFrame(columns=["行业"])

    for column in [
        "dividend_yield", "roe_ttm", "debt_to_assets", "revenue_yoy_qfa",
        "netprofit_yoy_qfa", "gross_profit_margin_qfa", "net_profit_margin_qfa",
    ]:
        history[column] = normalize_percent(history[column])
    for column in ["pe_ttm", "pb_lf", "ps_ttm"]:
        history[column] = finite_numeric(history[column]).where(lambda value: value > 0)

    grouped = history.groupby(["trade_date", "行业"], sort=True)
    industry_history = grouped[metric_columns].median().reset_index()
    for source, target in [
        ("revenue_yoy_qfa", "revenue_positive_breadth"),
        ("netprofit_yoy_qfa", "profit_positive_breadth"),
    ]:
        valid_positive = (history[source] > 0).where(history[source].notna())
        breadth = valid_positive.groupby([history["trade_date"], history["行业"]]).mean()
        industry_history = industry_history.merge(
            breadth.rename(target).reset_index(),
            on=["trade_date", "行业"],
            how="left",
        )

    def historical_percentile(series: pd.Series, higher_is_better: bool = True) -> float:
        series = finite_numeric(series).tail(config.industry_history_months)
        if series.empty or pd.isna(series.iloc[-1]):
            return np.nan
        valid = series.dropna()
        if len(valid) < config.industry_history_min_observations:
            return np.nan
        score = float(valid.rank(pct=True, method="average").iloc[-1])
        return score if higher_is_better else 1.0 - score

    def weighted_valid(values: dict[str, float], weights: dict[str, float]) -> float:
        valid_keys = [key for key in weights if pd.notna(values.get(key))]
        total_weight = sum(weights[key] for key in valid_keys)
        if total_weight <= 0:
            return 0.5
        return float(sum(values[key] * weights[key] for key in valid_keys) / total_weight)

    rows = []
    for industry, group in industry_history.groupby("行业", sort=True):
        group = group.sort_values("trade_date").tail(
            config.industry_history_months + 12
        ).copy()
        revenue = group["revenue_yoy_qfa"]
        profit = group["netprofit_yoy_qfa"]
        revenue_acceleration = revenue - revenue.shift(12)
        profit_acceleration = profit - profit.shift(12)
        growth_volatility = pd.concat(
            [
                revenue.rolling(12, min_periods=6).std(),
                profit.rolling(12, min_periods=6).std(),
            ],
            axis=1,
        ).mean(axis=1, skipna=True)
        margin_volatility = pd.concat(
            [
                group["gross_profit_margin_qfa"].rolling(12, min_periods=6).std(),
                group["net_profit_margin_qfa"].rolling(12, min_periods=6).std(),
            ],
            axis=1,
        ).mean(axis=1, skipna=True)

        scores = {
            "roe": historical_percentile(group["roe_ttm"]),
            "gross_margin": historical_percentile(group["gross_profit_margin_qfa"]),
            "net_margin": historical_percentile(group["net_profit_margin_qfa"]),
            "low_debt": historical_percentile(group["debt_to_assets"], higher_is_better=False),
            "pe": historical_percentile(group["pe_ttm"], higher_is_better=False),
            "pb": historical_percentile(group["pb_lf"], higher_is_better=False),
            "ps": historical_percentile(group["ps_ttm"], higher_is_better=False),
            "dividend": historical_percentile(group["dividend_yield"]),
            "revenue_growth": historical_percentile(revenue),
            "profit_growth": historical_percentile(profit),
            "revenue_acceleration": historical_percentile(revenue_acceleration),
            "profit_acceleration": historical_percentile(profit_acceleration),
            "revenue_breadth": historical_percentile(group["revenue_positive_breadth"]),
            "profit_breadth": historical_percentile(group["profit_positive_breadth"]),
            "stable_growth": historical_percentile(growth_volatility, higher_is_better=False),
            "stable_margin": historical_percentile(margin_volatility, higher_is_better=False),
        }
        is_financial = industry == "金融"
        fundamental_weights = (
            {"roe": 1.0}
            if is_financial
            else {"roe": 0.40, "gross_margin": 0.20, "net_margin": 0.25, "low_debt": 0.15}
        )
        valuation_weights = (
            {"pe": 0.25, "pb": 0.55, "dividend": 0.20}
            if is_financial
            else {"pe": 0.40, "pb": 0.25, "ps": 0.20, "dividend": 0.15}
        )
        stability_weights = {
            "revenue_breadth": 0.20,
            "profit_breadth": 0.25,
            "stable_growth": 0.30,
        }
        if not is_financial:
            stability_weights["stable_margin"] = 0.25
        growth_weights = {
            "revenue_growth": 0.30,
            "profit_growth": 0.40,
            "revenue_acceleration": 0.10,
            "profit_acceleration": 0.20,
        }
        cycle_weights = {
            "revenue_growth": 0.35,
            "profit_growth": 0.35,
            "revenue_acceleration": 0.15,
            "profit_acceleration": 0.15,
        }
        competition_weights = (
            {"roe": 1.0}
            if is_financial
            else {"roe": 0.4375, "gross_margin": 0.3125, "net_margin": 0.25}
        )
        rows.append({
            "行业": industry,
            "行业基本面得分": weighted_valid(scores, fundamental_weights),
            "行业估值得分": weighted_valid(scores, valuation_weights),
            "行业稳定性得分": weighted_valid(scores, stability_weights),
            "行业成长性得分": weighted_valid(scores, growth_weights),
            "行业周期得分": weighted_valid(scores, cycle_weights),
            "行业竞争格局得分": weighted_valid(scores, competition_weights),
            "行业历史观察月数": int(min(len(group), config.industry_history_months)),
        })
    return pd.DataFrame(rows)


def build_scores(
    frame: pd.DataFrame,
    config: StrategyConfig,
    industry_history_scores: pd.DataFrame | None = None,
) -> pd.DataFrame:
    frame = frame.copy()

    def rank_score(
        column: str,
        higher_is_better: bool = True,
        neutral: float = 0.5,
    ) -> pd.Series:
        return grouped_percentile(
            frame,
            column,
            higher_is_better=higher_is_better,
            neutral=neutral,
            lower_quantile=config.winsorize_lower_quantile,
            upper_quantile=config.winsorize_upper_quantile,
        )

    def market_score(
        column: str,
        higher_is_better: bool = True,
        neutral: float = 0.5,
        positive_only: bool = False,
    ) -> pd.Series:
        return global_percentile(
            frame,
            column,
            higher_is_better=higher_is_better,
            neutral=neutral,
            lower_quantile=config.winsorize_lower_quantile,
            upper_quantile=config.winsorize_upper_quantile,
            positive_only=positive_only,
        )

    def industry_score_from_global(stock_score: pd.Series) -> pd.Series:
        industry_median = stock_score.groupby(frame["行业"]).median()
        industry_rank = industry_median.rank(pct=True, method="average")
        return frame["行业"].map(industry_rank).fillna(0.5)

    percent_columns = [
        "dividend_yield", "roe_ttm", "debt_to_assets", "revenue_yoy_qfa",
        "netprofit_yoy_qfa", "gross_profit_margin_qfa", "net_profit_margin_qfa",
    ]
    for column in percent_columns:
        frame[column] = normalize_percent(frame[column])

    industry_components = {
        "roe": rank_score("roe_ttm"),
        "gross_margin": rank_score("gross_profit_margin_qfa"),
        "net_margin": rank_score("net_profit_margin_qfa"),
        "low_debt": rank_score("debt_to_assets", higher_is_better=False),
        "pe": positive_valuation_percentile(frame, "pe_ttm", config),
        "pb": positive_valuation_percentile(frame, "pb_lf", config),
        "ps": positive_valuation_percentile(frame, "ps_ttm", config),
        "dividend": rank_score("dividend_yield"),
        "positive_history": rank_score("growth_positive_ratio"),
        "stable_growth": rank_score("growth_volatility", higher_is_better=False),
        "stable_margin": rank_score("margin_volatility", higher_is_better=False),
        "revenue_growth": rank_score("revenue_yoy_qfa"),
        "profit_growth": rank_score("netprofit_yoy_qfa"),
        "revenue_acceleration": rank_score("revenue_acceleration"),
        "profit_acceleration": rank_score("profit_acceleration"),
        "size": rank_score("weighting_market_cap"),
    }
    global_components = {
        "roe": market_score("roe_ttm"),
        "gross_margin": market_score("gross_profit_margin_qfa"),
        "net_margin": market_score("net_profit_margin_qfa"),
        "low_debt": market_score("debt_to_assets", higher_is_better=False),
        "pe": market_score("pe_ttm", higher_is_better=False, neutral=0.35, positive_only=True),
        "pb": market_score("pb_lf", higher_is_better=False, neutral=0.35, positive_only=True),
        "ps": market_score("ps_ttm", higher_is_better=False, neutral=0.35, positive_only=True),
        "dividend": market_score("dividend_yield"),
        "positive_history": market_score("growth_positive_ratio"),
        "stable_growth": market_score("growth_volatility", higher_is_better=False),
        "stable_margin": market_score("margin_volatility", higher_is_better=False),
        "revenue_growth": market_score("revenue_yoy_qfa"),
        "profit_growth": market_score("netprofit_yoy_qfa"),
        "revenue_acceleration": market_score("revenue_acceleration"),
        "profit_acceleration": market_score("profit_acceleration"),
        "size": market_score("weighting_market_cap"),
    }
    guidance_confidence = finite_numeric(
        frame.get("earnings_guidance_confidence", pd.Series(0.0, index=frame.index))
    ).fillna(0.0).clip(0.0, 1.0)
    if "earnings_guidance_profit_yoy" in frame.columns:
        guidance_industry_growth = rank_score("earnings_guidance_profit_yoy")
        guidance_global_growth = market_score("earnings_guidance_profit_yoy")
        guidance_industry_acceleration = rank_score("earnings_guidance_profit_acceleration")
        guidance_global_acceleration = market_score("earnings_guidance_profit_acceleration")
        industry_components["profit_growth"] = (
            (1 - guidance_confidence) * industry_components["profit_growth"]
            + guidance_confidence * guidance_industry_growth
        )
        global_components["profit_growth"] = (
            (1 - guidance_confidence) * global_components["profit_growth"]
            + guidance_confidence * guidance_global_growth
        )
        industry_components["profit_acceleration"] = (
            (1 - guidance_confidence) * industry_components["profit_acceleration"]
            + guidance_confidence * guidance_industry_acceleration
        )
        global_components["profit_acceleration"] = (
            (1 - guidance_confidence) * global_components["profit_acceleration"]
            + guidance_confidence * guidance_global_acceleration
        )
    global_momentum_score = (
        0.50 * market_score("momentum_12_1")
        + 0.50 * market_score("momentum_6_1")
    )
    industry_momentum_score = (
        0.50 * rank_score("momentum_12_1")
        + 0.50 * rank_score("momentum_6_1")
    )
    global_low_size_score = market_score("weighting_market_cap", higher_is_better=False)
    industry_low_size_score = rank_score("weighting_market_cap", higher_is_better=False)
    if config.stock_factor_scope == "industry_only":
        frame["动量得分"] = industry_momentum_score
        frame["低规模得分"] = industry_low_size_score
    else:
        frame["动量得分"] = global_momentum_score
        frame["低规模得分"] = global_low_size_score

    global_fundamental = (
        0.40 * global_components["roe"] + 0.20 * global_components["gross_margin"]
        + 0.25 * global_components["net_margin"] + 0.15 * global_components["low_debt"]
    )
    global_valuation = (
        0.40 * global_components["pe"] + 0.25 * global_components["pb"]
        + 0.20 * global_components["ps"] + 0.15 * global_components["dividend"]
    )
    global_stability = (
        0.45 * global_components["positive_history"]
        + 0.30 * global_components["stable_growth"]
        + 0.25 * global_components["stable_margin"]
    )
    global_growth = (
        0.30 * global_components["revenue_growth"]
        + 0.40 * global_components["profit_growth"]
        + 0.10 * global_components["revenue_acceleration"]
        + 0.20 * global_components["profit_acceleration"]
    )
    global_cycle = (
        0.35 * global_components["revenue_growth"]
        + 0.35 * global_components["profit_growth"]
        + 0.15 * global_components["revenue_acceleration"]
        + 0.15 * global_components["profit_acceleration"]
    )
    global_competition = (
        0.35 * global_components["roe"] + 0.25 * global_components["gross_margin"]
        + 0.20 * global_components["net_margin"] + 0.20 * global_components["size"]
    )

    legacy_industry_scores = {
        "行业基本面得分": industry_score_from_global(global_fundamental),
        "行业估值得分": industry_score_from_global(global_valuation),
        "行业稳定性得分": industry_score_from_global(global_stability),
        "行业成长性得分": industry_score_from_global(global_growth),
        "行业周期得分": industry_score_from_global(global_cycle),
        "行业竞争格局得分": industry_score_from_global(global_competition),
    }
    industry_score_columns = [
        "行业基本面得分", "行业估值得分", "行业稳定性得分",
        "行业成长性得分", "行业周期得分", "行业竞争格局得分",
    ]
    if industry_history_scores is not None and not industry_history_scores.empty:
        history_lookup = industry_history_scores.set_index("行业")
        for column in industry_score_columns:
            frame[column] = frame["行业"].map(history_lookup[column]).fillna(0.5)
        frame["行业历史观察月数"] = frame["行业"].map(
            history_lookup["行业历史观察月数"]
        ).fillna(0).astype(int)
    else:
        # 兼容直接调用 build_scores 的场景；正式构建会传入行业自身历史得分。
        for column, values in legacy_industry_scores.items():
            frame[column] = values
        frame["行业历史观察月数"] = 0
    frame["行业动量得分"] = industry_score_from_global(global_momentum_score)
    industry_base_score = sum(
        frame[industry_column] * getattr(config, weight_attr)
        for industry_column, weight_attr in {
            "行业基本面得分": "fundamental_weight",
            "行业估值得分": "valuation_weight",
            "行业稳定性得分": "stability_weight",
            "行业成长性得分": "growth_weight",
            "行业周期得分": "cycle_weight",
            "行业竞争格局得分": "competition_weight",
        }.items()
    )
    if config.scoring_mode == "hybrid":
        frame["行业综合得分"] = (
            (1 - config.industry_momentum_weight) * industry_base_score
            + config.industry_momentum_weight * frame["行业动量得分"]
        )
    else:
        frame["行业综合得分"] = industry_base_score

    if config.scoring_mode == "hybrid":
        industry_only = config.stock_factor_scope == "industry_only"
        fundamental_components = {
            key: blend_scores(
                industry_components[key], global_components[key], 1.0 if industry_only else 0.70
            )
            for key in ["roe", "gross_margin", "net_margin"]
        }
        fundamental_components["low_debt"] = industry_components["low_debt"]
        frame["基本面得分"] = (
            0.40 * fundamental_components["roe"]
            + 0.20 * fundamental_components["gross_margin"]
            + 0.25 * fundamental_components["net_margin"]
            + 0.15 * fundamental_components["low_debt"]
        )
        valuation_components = {
            key: blend_scores(
                industry_components[key], global_components[key], 1.0 if industry_only else 0.70
            )
            for key in ["pe", "pb", "ps", "dividend"]
        }
        frame["估值得分"] = (
            0.40 * valuation_components["pe"] + 0.25 * valuation_components["pb"]
            + 0.20 * valuation_components["ps"] + 0.15 * valuation_components["dividend"]
        )
        stability_components = {
            key: blend_scores(
                industry_components[key], global_components[key], 1.0 if industry_only else 0.80
            )
            for key in ["positive_history", "stable_growth", "stable_margin"]
        }
        frame["稳定性得分"] = (
            0.45 * stability_components["positive_history"]
            + 0.30 * stability_components["stable_growth"]
            + 0.25 * stability_components["stable_margin"]
        )
        growth_components = {
            key: blend_scores(
                industry_components[key], global_components[key], 1.0 if industry_only else 0.50
            )
            for key in ["revenue_growth", "profit_growth", "revenue_acceleration", "profit_acceleration"]
        }
        frame["成长性得分"] = (
            0.30 * growth_components["revenue_growth"]
            + 0.40 * growth_components["profit_growth"]
            + 0.10 * growth_components["revenue_acceleration"]
            + 0.20 * growth_components["profit_acceleration"]
        )
        individual_cycle = (
            0.35 * growth_components["revenue_growth"]
            + 0.35 * growth_components["profit_growth"]
            + 0.15 * growth_components["revenue_acceleration"]
            + 0.15 * growth_components["profit_acceleration"]
        )
        if industry_only:
            frame["周期阶段得分"] = individual_cycle
        else:
            frame["周期阶段得分"] = (
                0.70 * legacy_industry_scores["行业周期得分"]
                + 0.30 * individual_cycle
            )
        industry_competition = (
            0.35 * industry_components["roe"] + 0.25 * industry_components["gross_margin"]
            + 0.20 * industry_components["net_margin"] + 0.20 * industry_components["size"]
        )
        frame["竞争格局得分"] = (
            industry_competition
            if industry_only
            else 0.80 * industry_competition + 0.20 * global_competition
        )
    else:
        frame["基本面得分"] = (
            0.40 * industry_components["roe"] + 0.20 * industry_components["gross_margin"]
            + 0.25 * industry_components["net_margin"] + 0.15 * industry_components["low_debt"]
        )
        frame["估值得分"] = (
            0.40 * industry_components["pe"] + 0.25 * industry_components["pb"]
            + 0.20 * industry_components["ps"] + 0.15 * industry_components["dividend"]
        )
        frame["稳定性得分"] = (
            0.45 * industry_components["positive_history"]
            + 0.30 * industry_components["stable_growth"]
            + 0.25 * industry_components["stable_margin"]
        )
        frame["成长性得分"] = (
            0.30 * industry_components["revenue_growth"]
            + 0.40 * industry_components["profit_growth"]
            + 0.10 * industry_components["revenue_acceleration"]
            + 0.20 * industry_components["profit_acceleration"]
        )
        frame["周期阶段得分"] = (
            0.35 * industry_components["revenue_growth"]
            + 0.35 * industry_components["profit_growth"]
            + 0.15 * industry_components["revenue_acceleration"]
            + 0.15 * industry_components["profit_acceleration"]
        )
        frame["竞争格局得分"] = (
            0.35 * industry_components["roe"] + 0.25 * industry_components["gross_margin"]
            + 0.20 * industry_components["net_margin"] + 0.20 * industry_components["size"]
        )
    frame["周期阶段"] = frame.apply(classify_cycle_stage, axis=1)

    required = [
        "roe_ttm", "debt_to_assets", "pe_ttm", "pb_lf", "ps_ttm",
        "revenue_yoy_qfa", "netprofit_yoy_qfa", "gross_profit_margin_qfa",
        "net_profit_margin_qfa", "weighting_market_cap",
    ]
    frame["数据完整度"] = frame[required].notna().mean(axis=1)
    six_dimension_score = sum(
        frame[score_column] * getattr(config, weight_attr)
        for score_column, weight_attr in SCORE_WEIGHTS.items()
    )
    if config.scoring_mode == "hybrid":
        residual_weight = 1 - config.hybrid_momentum_weight - config.hybrid_low_size_weight
        raw_score = (
            residual_weight * six_dimension_score
            + config.hybrid_momentum_weight * frame["动量得分"]
            + config.hybrid_low_size_weight * frame["低规模得分"]
        )
    else:
        raw_score = six_dimension_score
    coverage_multiplier = 0.85 + 0.15 * frame["数据完整度"]
    frame["原始总分"] = raw_score
    frame["总分"] = (raw_score * coverage_multiplier * 100).clip(0, 100)
    return frame


def ensure_benchmark_weight_table(conn: sqlite3.Connection) -> None:
    conn.execute("""
        CREATE TABLE IF NOT EXISTS index_industry_weight_snapshot (
            snapshot_date TEXT NOT NULL,
            index_code TEXT NOT NULL,
            classification_system TEXT NOT NULL,
            industry_level1 TEXT NOT NULL,
            industry_weight REAL NOT NULL,
            constituent_count INTEGER NOT NULL,
            source_field TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            PRIMARY KEY (
                snapshot_date, index_code, classification_system, industry_level1
            )
        )
    """)


def load_local_benchmark_industry_weights(as_of: str) -> tuple[pd.Series | None, str | None]:
    with sqlite3.connect(MARKET_DB_PATH) as conn:
        ensure_benchmark_weight_table(conn)
        snapshot_date = conn.execute(
            """
            SELECT MAX(snapshot_date)
            FROM index_industry_weight_snapshot
            WHERE index_code = ? AND classification_system = ?
              AND snapshot_date <= ?
            """,
            [BENCHMARK_CODE, INDUSTRY_SYSTEM, as_of],
        ).fetchone()[0]
        if not snapshot_date:
            return None, None
        rows = pd.read_sql_query(
            """
            SELECT industry_level1, industry_weight
            FROM index_industry_weight_snapshot
            WHERE snapshot_date = ? AND index_code = ?
              AND classification_system = ?
            """,
            conn,
            params=[snapshot_date, BENCHMARK_CODE, INDUSTRY_SYSTEM],
        )
    weights = rows.set_index("industry_level1")["industry_weight"]
    return weights / weights.sum(), snapshot_date


def save_benchmark_industry_weights(
    as_of: str,
    weights: pd.Series,
    counts: pd.Series,
) -> None:
    updated_at = pd.Timestamp.now().isoformat(timespec="seconds")
    rows = [
        (
            as_of, BENCHMARK_CODE, INDUSTRY_SYSTEM, str(industry), float(weight),
            int(counts.get(industry, 0)), "Wind indexconstituent.i_weight", updated_at,
        )
        for industry, weight in weights.items()
    ]
    with sqlite3.connect(MARKET_DB_PATH) as conn:
        ensure_benchmark_weight_table(conn)
        conn.executemany(
            """
            INSERT INTO index_industry_weight_snapshot (
                snapshot_date, index_code, classification_system, industry_level1,
                industry_weight, constituent_count, source_field, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(snapshot_date, index_code, classification_system, industry_level1)
            DO UPDATE SET
                industry_weight = excluded.industry_weight,
                constituent_count = excluded.constituent_count,
                source_field = excluded.source_field,
                updated_at = excluded.updated_at
            """,
            rows,
        )
        conn.commit()


def fetch_benchmark_industry_weights(as_of: str, industry_map: dict[str, str]) -> pd.Series:
    try:
        from WindPy import w
    except ImportError as exc:
        raise RuntimeError("未安装 WindPy，请使用 --benchmark-weights 提供本地行业权重") from exc

    start_result = w.start()
    if getattr(start_result, "ErrorCode", 0) != 0:
        raise RuntimeError(f"Wind 启动失败: {getattr(start_result, 'ErrorCode', None)}")
    result = w.wset(
        "indexconstituent",
        f"date={as_of.replace('-', '')};windcode={BENCHMARK_CODE};field=wind_code,sec_name,i_weight",
    )
    if result.ErrorCode != 0 or not result.Data:
        raise RuntimeError(f"沪深300成分权重获取失败: ErrorCode={result.ErrorCode}")
    fields = {str(field).lower(): values for field, values in zip(result.Fields, result.Data)}
    rows = pd.DataFrame({
        "wind_code": fields["wind_code"],
        "benchmark_weight": pd.to_numeric(fields["i_weight"], errors="coerce") / 100.0,
    })
    rows["行业"] = rows["wind_code"].map(industry_map).fillna("未分类")
    weights = rows.groupby("行业")["benchmark_weight"].sum(min_count=1).dropna()
    if weights.sum() <= 0:
        raise RuntimeError("沪深300行业权重为空")
    weights = weights / weights.sum()
    save_benchmark_industry_weights(as_of, weights, rows.groupby("行业").size())
    return weights


def load_benchmark_weights(path: Path) -> pd.Series:
    if path.suffix.lower() in {".xlsx", ".xls"}:
        frame = pd.read_excel(path)
    else:
        frame = pd.read_csv(path)
    required = {"industry", "target_weight"}
    if not required.issubset(frame.columns):
        raise ValueError(f"本地行业权重文件必须包含列: {sorted(required)}")
    values = finite_numeric(frame["target_weight"])
    if values.max() > 1.5:
        values = values / 100.0
    weights = values.groupby(frame["industry"].astype(str)).sum().dropna()
    if weights.sum() <= 0:
        raise ValueError("本地行业权重合计必须大于 0")
    return weights / weights.sum()


def load_previous_constituents(path: Path) -> set[str]:
    if path.suffix.lower() in {".xlsx", ".xls"}:
        frame = pd.read_excel(path, sheet_name=0)
    else:
        frame = pd.read_csv(path)
    code_column = next(
        (column for column in ["wind_code", "证券代码"] if column in frame.columns),
        None,
    )
    if code_column is None:
        raise ValueError("上一期成分文件必须包含 wind_code 或 证券代码列")
    return set(frame[code_column].dropna().astype(str))


def find_latest_previous_components(
    as_of: str,
    search_directories: list[Path],
) -> Path | None:
    """查找日期严格早于本次截止日的最近一份单期指数构建文件。"""
    cutoff = pd.Timestamp(as_of)
    candidates = []
    visited_directories = set()
    for directory in search_directories:
        directory = directory.resolve()
        if directory in visited_directories or not directory.is_dir():
            continue
        visited_directories.add(directory)
        for path in directory.glob("全A股质量成长800_行业比较混合评分_*.xlsx"):
            if path.name.startswith("~$"):
                continue
            match = BUILD_OUTPUT_PATTERN.fullmatch(path.name)
            if match is None:
                continue
            file_date = pd.to_datetime(match.group(1), format="%Y%m%d", errors="coerce")
            if pd.notna(file_date) and file_date < cutoff:
                candidates.append((file_date, path))
    if not candidates:
        return None
    return max(candidates, key=lambda item: item[0])[1]


def select_constituents_with_buffer(
    eligible: pd.DataFrame,
    previous_constituents: set[str] | None,
    config: StrategyConfig,
) -> pd.DataFrame:
    """固定数量选股；核心区直接入选，缓冲区内优先保留原成分股。"""
    ranked = eligible.sort_values(
        ["总分", "weighting_market_cap"], ascending=[False, False]
    ).copy()
    ranked["全市场评分排名"] = np.arange(1, len(ranked) + 1)
    target_count = min(config.max_constituents, len(ranked))
    if target_count == 0:
        return ranked.head(0)
    if target_count < config.max_constituents or not previous_constituents:
        selected = ranked.head(target_count).copy()
        selected["入选方式"] = "按分入选"
    else:
        core_cutoff = int(math.floor(config.max_constituents * (1 - config.constituent_buffer_ratio)))
        buffer_cutoff = int(math.ceil(config.max_constituents * (1 + config.constituent_buffer_ratio)))
        core = ranked[ranked["全市场评分排名"] <= core_cutoff].copy()
        core["入选方式"] = "核心区直接入选"
        retained = ranked[
            ranked["全市场评分排名"].between(core_cutoff + 1, buffer_cutoff)
            & ranked["wind_code"].isin(previous_constituents)
        ].head(target_count - len(core)).copy()
        retained["入选方式"] = "缓冲区原成分保留"
        chosen_codes = set(core["wind_code"]) | set(retained["wind_code"])
        fill = ranked[~ranked["wind_code"].isin(chosen_codes)].head(
            target_count - len(core) - len(retained)
        ).copy()
        fill["入选方式"] = "按分补足"
        selected = pd.concat([core, retained, fill], ignore_index=False)
        selected = selected.sort_values(
            ["总分", "weighting_market_cap"], ascending=[False, False]
        )
    selected["评分排名"] = np.arange(1, len(selected) + 1)
    return selected


def allocate_industry_quotas(
    eligible: pd.DataFrame,
    industry_targets: pd.DataFrame,
    config: StrategyConfig,
) -> pd.DataFrame:
    """按目标行业权重分配整数成分名额，并将不足行业的名额重新分配。"""
    target_count = min(config.max_constituents, len(eligible))
    eligible_counts = eligible.groupby("行业").size()
    quotas = industry_targets[["行业", "目标行业权重"]].copy()
    quotas["行业合格股票数"] = quotas["行业"].map(eligible_counts).fillna(0).astype(int)
    quotas["行业理想名额"] = quotas["目标行业权重"] * target_count
    quotas["行业参考名额"] = 0

    # 最高平均数法：正常情况下逼近目标权重；某行业容量不足时，缺口按
    # 其他行业目标权重比例自然重新分配，而不是平均摊给所有行业。
    for _ in range(target_count):
        available = quotas["行业参考名额"] < quotas["行业合格股票数"]
        if not available.any():
            break
        choices = quotas.loc[available].copy()
        choices["名额优先值"] = (
            choices["目标行业权重"] / (choices["行业参考名额"] + 1)
        )
        chosen_index = choices.sort_values(
            ["名额优先值", "目标行业权重", "行业"],
            ascending=[False, False, True],
        ).index[0]
        quotas.loc[chosen_index, "行业参考名额"] += 1
    if int(quotas["行业参考名额"].sum()) != target_count:
        raise RuntimeError("行业参考名额无法分配到目标成分数量")
    return quotas


def select_constituents_by_industry_quota(
    eligible: pd.DataFrame,
    industry_targets: pd.DataFrame,
    previous_constituents: set[str] | None,
    config: StrategyConfig,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """按行业目标权重分名额，再在各行业内部按评分及缓冲规则选股。"""
    ranked = eligible.sort_values(
        ["总分", "weighting_market_cap"], ascending=[False, False]
    ).copy()
    ranked["全市场评分排名"] = np.arange(1, len(ranked) + 1)
    ranked["行业内评分排名"] = ranked.groupby("行业")["总分"].rank(
        ascending=False,
        method="first",
    ).astype(int)
    quotas = allocate_industry_quotas(ranked, industry_targets, config)
    quota_lookup = quotas.set_index("行业")["行业参考名额"]
    selected_frames = []

    for industry, group in ranked.groupby("行业", sort=True):
        quota = int(quota_lookup.get(industry, 0))
        if quota <= 0:
            continue
        group = group.sort_values(
            ["总分", "weighting_market_cap"], ascending=[False, False]
        ).copy()
        if not previous_constituents or quota >= len(group):
            chosen = group.head(quota).copy()
            chosen["入选方式"] = "按行业名额入选"
        else:
            core_cutoff = max(
                1,
                int(math.floor(quota * (1 - config.constituent_buffer_ratio))),
            )
            buffer_cutoff = min(
                len(group),
                int(math.ceil(quota * (1 + config.constituent_buffer_ratio))),
            )
            core = group.head(core_cutoff).copy()
            core["入选方式"] = "行业核心区直接入选"
            retained = group.iloc[core_cutoff:buffer_cutoff]
            retained = retained[retained["wind_code"].isin(previous_constituents)].head(
                quota - len(core)
            ).copy()
            retained["入选方式"] = "行业缓冲区原成分保留"
            chosen_codes = set(core["wind_code"]) | set(retained["wind_code"])
            fill = group[~group["wind_code"].isin(chosen_codes)].head(
                quota - len(core) - len(retained)
            ).copy()
            fill["入选方式"] = "按行业评分补足"
            chosen = pd.concat([core, retained, fill], ignore_index=False)
        chosen["行业参考名额"] = quota
        selected_frames.append(chosen)

    if not selected_frames:
        return ranked.head(0), quotas
    selected = pd.concat(selected_frames, ignore_index=False).sort_values(
        ["总分", "weighting_market_cap"], ascending=[False, False]
    )
    selected["评分排名"] = np.arange(1, len(selected) + 1)
    return selected, quotas


def bounded_industry_targets(
    selected: pd.DataFrame,
    benchmark: pd.Series,
    config: StrategyConfig,
    scored_universe: pd.DataFrame | None = None,
) -> pd.DataFrame:
    exceptional_cutoff = selected["总分"].quantile(config.exceptional_score_quantile)
    exceptional = selected[selected["总分"] >= exceptional_cutoff]
    alpha_mass = exceptional.groupby("行业")["总分"].sum()
    alpha_mass = alpha_mass / alpha_mass.sum() if alpha_mass.sum() > 0 else benchmark.copy()
    if scored_universe is not None and "行业综合得分" in scored_universe:
        industry_scores = scored_universe.groupby("行业")["行业综合得分"].first()
    else:
        industry_scores = pd.Series(dtype=float)

    industries = sorted(set(selected["行业"]) | set(benchmark.index))
    selected_counts = selected.groupby("行业").size()
    rows = []
    for industry in industries:
        benchmark_weight = float(benchmark.get(industry, 0.0))
        score_weight = float(alpha_mass.get(industry, 0.0))
        industry_score = float(industry_scores.get(industry, 0.5))
        if config.industry_target_mode == "industry_score":
            attractiveness_signal = float(np.clip(2 * industry_score - 1, -1, 1))
            if benchmark_weight > 0:
                proposed = benchmark_weight * (
                    1
                    + config.max_industry_relative_deviation
                    * config.industry_between_tilt_strength
                    * attractiveness_signal
                )
            else:
                proposed = config.max_off_benchmark_industry_weight * max(attractiveness_signal, 0)
        else:
            attractiveness_signal = np.nan
            proposed = (
                (1 - config.industry_score_tilt) * benchmark_weight
                + config.industry_score_tilt * score_weight
            )
        if benchmark_weight > 0:
            lower = benchmark_weight * (1 - config.max_industry_relative_deviation)
            upper = benchmark_weight * (1 + config.max_industry_relative_deviation)
        else:
            lower = 0.0
            upper = config.max_off_benchmark_industry_weight
        constituent_count = int(selected_counts.get(industry, 0))
        capacity = constituent_count * config.max_stock_weight
        upper = min(upper, capacity)
        lower = min(lower, upper)
        if constituent_count == 0:
            lower = upper = proposed = 0.0
        rows.append({
            "行业": industry,
            "沪深300行业权重": benchmark_weight,
            "高分股权重信号": score_weight,
            "行业吸引力得分": industry_score,
            "行业吸引力信号": attractiveness_signal,
            "建议行业权重": proposed,
            "行业权重下限": lower,
            "行业权重上限": upper,
            "5%个股上限承载能力": capacity,
        })
    result = pd.DataFrame(rows)
    lower = result["行业权重下限"].to_numpy(dtype=float)
    upper = result["行业权重上限"].to_numpy(dtype=float)
    if lower.sum() > 1 + 1e-9 or upper.sum() < 1 - 1e-9:
        raise RuntimeError("行业上下限不可行：无法在约束内使行业权重合计为 100%")
    target = np.clip(result["建议行业权重"].to_numpy(dtype=float), lower, upper)
    for _ in range(len(target) * 2 + 2):
        residual = 1.0 - target.sum()
        if abs(residual) <= 1e-12:
            break
        capacity = (upper - target) if residual > 0 else (target - lower)
        available = capacity.sum()
        if available <= 1e-15:
            raise RuntimeError("行业权重投影失败：约束没有剩余调整空间")
        target += np.sign(residual) * min(abs(residual), available) * capacity / available
        target = np.clip(target, lower, upper)
    result["目标行业权重"] = target
    result["相对沪深300偏离"] = result["目标行业权重"] - result["沪深300行业权重"]
    return result


def cap_and_redistribute(raw: pd.Series, total: float, cap: float) -> pd.Series:
    if raw.empty or total <= 0:
        return pd.Series(0.0, index=raw.index)
    raw = finite_numeric(raw).fillna(0).clip(lower=0)
    if raw.sum() <= 0:
        raw[:] = 1.0
    if total > cap * len(raw) + 1e-12:
        raise RuntimeError(
            f"{len(raw)} 只股票在 {cap:.2%} 个股上限下无法承载 {total:.2%} 权重"
        )
    weights = raw / raw.sum() * total
    fixed = pd.Series(False, index=weights.index)
    for _ in range(len(weights) + 1):
        over = (weights > cap + 1e-12) & ~fixed
        if not over.any():
            break
        weights.loc[over] = cap
        fixed |= over
        remaining = total - weights.loc[fixed].sum()
        free = ~fixed
        if not free.any() or remaining <= 0:
            break
        free_raw = raw.loc[free]
        if free_raw.sum() > 0:
            weights.loc[free] = remaining * free_raw / free_raw.sum()
        else:
            weights.loc[free] = remaining / int(free.sum())
    return weights


def assign_stock_weights(
    selected: pd.DataFrame,
    industry_targets: pd.DataFrame,
    config: StrategyConfig,
) -> pd.DataFrame:
    selected = selected.copy()
    targets = industry_targets.set_index("行业")["目标行业权重"]
    selected["指数权重"] = 0.0
    selected["行业内市值权重"] = 0.0
    selected["评分倾斜倍数"] = 1.0
    for industry, indexes in selected.groupby("行业").groups.items():
        subset = selected.loc[indexes]
        target = float(targets.get(industry, 0.0))
        market_cap = finite_numeric(subset["weighting_market_cap"]).clip(lower=0)
        base = market_cap / market_cap.sum() if market_cap.sum() > 0 else pd.Series(1 / len(subset), index=indexes)
        score_std = subset["总分"].std(ddof=0)
        zscore = (subset["总分"] - subset["总分"].median()) / (score_std if score_std > 1e-9 else 1.0)
        tilt = np.exp(config.stock_score_tilt * zscore.clip(-2, 2)).clip(0.50, 2.00)
        raw = base * tilt
        selected.loc[indexes, "行业内市值权重"] = base
        selected.loc[indexes, "评分倾斜倍数"] = tilt
        selected.loc[indexes, "指数权重"] = cap_and_redistribute(
            raw, target, config.max_stock_weight
        )
    selected["指数权重"] = selected["指数权重"] / selected["指数权重"].sum()
    return selected


def build_index(
    as_of: str,
    config: StrategyConfig,
    benchmark_path: Path | None = None,
    previous_constituents: set[str] | None = None,
):
    validate_config(config)
    universe = load_universe(as_of)
    codes = universe["wind_code"].tolist()
    industries = load_industries(codes, as_of)
    fundamentals = load_latest_fundamentals(codes, as_of)
    valuation, valuation_date = load_valuation(codes, as_of, config)
    if config.scoring_mode == "hybrid":
        momentum = load_momentum_features(codes, as_of)
    else:
        momentum = pd.DataFrame({
            "wind_code": codes,
            "momentum_12_1": np.nan,
            "momentum_6_1": np.nan,
            "momentum_data_coverage": 0.0,
        })
    history = report_history_features(load_report_history(codes, as_of, config.history_quarters))

    frame = (
        universe
        .merge(industries[["wind_code", "industry", "industry_source"]], on="wind_code", how="left")
        .merge(fundamentals, on="wind_code", how="left")
        .merge(valuation, on="wind_code", how="left")
        .merge(momentum, on="wind_code", how="left")
        .merge(history, on="wind_code", how="left")
    )
    frame["行业"] = frame["industry"].fillna("未分类")
    frame["证券名称"] = frame["sec_name"]
    guidance = load_earnings_guidance(frame, as_of, config)
    frame = frame.merge(guidance, on="wind_code", how="left")
    frame["earnings_guidance_profit_acceleration"] = (
        finite_numeric(frame.get(
            "earnings_guidance_profit_yoy",
            pd.Series(np.nan, index=frame.index),
        ))
        - normalize_percent(frame["netprofit_yoy_qfa"])
    )
    industry_history_scores = build_industry_history_scores(
        frame[["wind_code", "行业"]],
        as_of,
        config,
    )
    frame = build_scores(frame, config, industry_history_scores)

    eligible = frame[frame["数据完整度"] >= config.min_data_coverage].copy()
    eligible = eligible[finite_numeric(eligible["weighting_market_cap"]) > config.min_market_cap]
    if config.exclude_st:
        eligible = eligible[
            ~eligible["证券名称"].fillna("").str.upper().str.contains(r"(?:^|\*)ST|PT", regex=True)
        ]
    eligible = eligible.sort_values(["总分", "weighting_market_cap"], ascending=[False, False])

    industry_map = frame.set_index("wind_code")["行业"].to_dict()
    benchmark_source = "外部文件"
    if benchmark_path is not None:
        benchmark = load_benchmark_weights(benchmark_path)
    else:
        benchmark, benchmark_date = load_local_benchmark_industry_weights(as_of)
        if benchmark is None:
            benchmark = fetch_benchmark_industry_weights(as_of, industry_map)
            benchmark_source = f"Wind:{as_of}"
        else:
            benchmark_source = f"本地沪深300行业权重:{benchmark_date}"
    if config.constituent_selection_mode == "industry_quota":
        selection_pool = eligible[
            eligible["总分"] >= config.industry_quota_min_score
        ].copy()
        if len(selection_pool) < min(config.max_constituents, len(eligible)):
            raise RuntimeError(
                f"行业名额最低总分 {config.industry_quota_min_score:.2f} 后仅剩"
                f" {len(selection_pool)} 只股票，无法选满 {config.max_constituents} 只"
            )
        preliminary_industry_targets = bounded_industry_targets(
            selection_pool,
            benchmark,
            config,
            frame,
        )
        selected, quota_table = select_constituents_by_industry_quota(
            selection_pool,
            preliminary_industry_targets,
            previous_constituents,
            config,
        )
    else:
        selected = select_constituents_with_buffer(eligible, previous_constituents, config)
        quota_table = pd.DataFrame({
            "行业": selected["行业"].drop_duplicates(),
        })
        quota_table["行业理想名额"] = pd.NA
        quota_table["行业参考名额"] = quota_table["行业"].map(
            selected.groupby("行业").size()
        ).astype(int)
        quota_table["行业合格股票数"] = quota_table["行业"].map(
            eligible.groupby("行业").size()
        ).fillna(0).astype(int)
    if selected.empty:
        raise RuntimeError("没有股票通过入选条件")

    industry_targets = bounded_industry_targets(selected, benchmark, config, frame)
    selected = assign_stock_weights(selected, industry_targets, config)

    actual = selected.groupby("行业")["指数权重"].agg(["sum", "count"]).reset_index()
    actual = actual.rename(columns={"sum": "实际行业权重", "count": "入选数量"})
    industry_summary = industry_targets.merge(
        industry_history_scores,
        on="行业",
        how="left",
    ).merge(
        quota_table[["行业", "行业理想名额", "行业参考名额", "行业合格股票数"]],
        on="行业",
        how="left",
    ).merge(actual, on="行业", how="left").fillna(
        {"实际行业权重": 0.0, "入选数量": 0}
    )
    industry_summary["行业权重校验差"] = (
        industry_summary["实际行业权重"] - industry_summary["目标行业权重"]
    )

    diagnostics = {
        "截止日": as_of,
        "股票池数量": len(universe),
        "符合最低数据要求数量": len(eligible),
        "最终成分数量": len(selected),
        "行业数量": selected["行业"].nunique(),
        "市值截面日期": valuation_date,
        "股票池来源": str(universe["universe_source"].iloc[0]),
        "行业权重来源": benchmark_source,
        "市值截面滞后天数": int((pd.Timestamp(as_of) - pd.Timestamp(valuation_date)).days),
        "平均数据完整度": float(selected["数据完整度"].mean()),
        "平均动量数据完整度": float(selected["momentum_data_coverage"].mean()),
        "完整动量历史成分数量": int((selected["momentum_data_coverage"] >= 1.0).sum()),
        "权重合计": float(selected["指数权重"].sum()),
        "最大个股权重": float(selected["指数权重"].max()),
        "个股评分倾斜强度": config.stock_score_tilt,
        "个股评分模式": config.scoring_mode,
        "个股因子比较范围": config.stock_factor_scope,
        "成分选择模式": config.constituent_selection_mode,
        "行业名额最低总分": config.industry_quota_min_score,
        "行业权重模式": config.industry_target_mode,
        "行业比较方法": "相对行业自身历史",
        "行业历史窗口月数": config.industry_history_months,
        "行业历史最少观察月数": config.industry_history_min_observations,
        "使用上一期成分缓冲": bool(previous_constituents),
        "缓冲区保留数量": int(selected["入选方式"].str.contains("缓冲区原成分保留").sum()),
        "入选股票最差全市场排名": int(selected["全市场评分排名"].max()),
        "业绩预告快报功能": config.earnings_guidance_enabled,
        "全市场有效业绩预告快报数量": int(
            frame.get("earnings_guidance_confidence", pd.Series(dtype=float)).fillna(0).gt(0).sum()
        ),
        "成分股有效业绩预告快报数量": int(
            selected.get("earnings_guidance_confidence", pd.Series(dtype=float)).fillna(0).gt(0).sum()
        ),
        "全市场有效业绩预告数量": int(
            frame.get("earnings_guidance_type", pd.Series(dtype=str)).eq("notice").sum()
        ),
        "全市场有效业绩快报数量": int(
            frame.get("earnings_guidance_type", pd.Series(dtype=str)).eq("express").sum()
        ),
        "成分股有效业绩预告数量": int(
            selected.get("earnings_guidance_type", pd.Series(dtype=str)).eq("notice").sum()
        ),
        "成分股有效业绩快报数量": int(
            selected.get("earnings_guidance_type", pd.Series(dtype=str)).eq("express").sum()
        ),
    }
    return selected, industry_summary, frame, diagnostics


def export_result(
    selected: pd.DataFrame,
    industry_summary: pd.DataFrame,
    all_scores: pd.DataFrame,
    diagnostics: dict,
    config: StrategyConfig,
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    selected_columns = [
        "评分排名", "全市场评分排名", "行业内评分排名", "行业参考名额", "入选方式",
        "wind_code", "证券名称", "行业", "总分", "原始总分",
        *SCORE_WEIGHTS.keys(), "动量得分", "低规模得分", "行业综合得分",
        "行业基本面得分", "行业估值得分", "行业稳定性得分", "行业成长性得分",
        "行业周期得分", "行业竞争格局得分", "行业动量得分", "行业历史观察月数",
        "周期阶段", "数据完整度", "指数权重",
        "行业内市值权重", "评分倾斜倍数", "weighting_market_cap",
        "mkt_cap_ard", "free_float_mkt_cap", "pe_ttm", "pb_lf", "ps_ttm",
        "dividend_yield", "roe_ttm", "debt_to_assets", "revenue_yoy_qfa",
        "netprofit_yoy_qfa", "gross_profit_margin_qfa", "net_profit_margin_qfa",
        "momentum_12_1", "momentum_6_1", "momentum_data_coverage",
        "report_count", "qfa_report_period", "qfa_available_date",
        "earnings_guidance_report_period", "earnings_guidance_announcement_date",
        "earnings_guidance_type", "earnings_guidance_style",
        "earnings_guidance_profit_yoy", "earnings_guidance_profit_acceleration",
        "earnings_guidance_confidence", "earnings_guidance_source",
        "fundamental_source_date", "valuation_source_date", "industry_source",
    ]
    selected_columns = [column for column in selected_columns if column in selected.columns]
    ranking_columns = [
        "wind_code", "证券名称", "行业", "总分", *SCORE_WEIGHTS.keys(),
        "动量得分", "低规模得分", "行业综合得分", "行业基本面得分",
        "行业估值得分", "行业稳定性得分", "行业成长性得分", "行业周期得分",
        "行业竞争格局得分", "行业动量得分", "行业历史观察月数",
        "周期阶段", "数据完整度", "weighting_market_cap",
        "earnings_guidance_report_period", "earnings_guidance_announcement_date",
        "earnings_guidance_type", "earnings_guidance_profit_yoy",
        "earnings_guidance_profit_acceleration", "earnings_guidance_confidence",
    ]
    ranking_columns = [column for column in ranking_columns if column in all_scores.columns]
    parameters = pd.DataFrame(
        [(key, value) for key, value in asdict(config).items()],
        columns=["参数", "值"],
    )
    diagnostics_frame = pd.DataFrame(diagnostics.items(), columns=["检查项", "结果"])
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        selected[selected_columns].to_excel(writer, sheet_name="800成分与权重", index=False)
        industry_summary.to_excel(writer, sheet_name="行业权重", index=False)
        all_scores.sort_values("总分", ascending=False)[ranking_columns].to_excel(
            writer, sheet_name="全市场评分", index=False
        )
        diagnostics_frame.to_excel(writer, sheet_name="诊断", index=False)
        parameters.to_excel(writer, sheet_name="参数", index=False)


TRADING_COST_RATE = 0.0025


def load_signal_dates(start_date: str, end_date: str) -> list[str]:
    with sqlite3.connect(FUNDAMENTAL_DB_PATH) as conn:
        rows = conn.execute(
            """
            SELECT DISTINCT trade_date FROM fundamentals
            WHERE trade_date >= ? AND trade_date <= ?
            ORDER BY trade_date
            """,
            [start_date, end_date],
        ).fetchall()
    return [row[0] for row in rows]


def load_trading_dates(start_date: str, end_date: str) -> pd.DatetimeIndex:
    with sqlite3.connect(MARKET_DB_PATH) as conn:
        rows = conn.execute(
            """
            SELECT DISTINCT trade_date FROM daily_prices
            WHERE adjusted = 'F' AND trade_date >= ? AND trade_date <= ?
            ORDER BY trade_date
            """,
            [start_date, end_date],
        ).fetchall()
    return pd.DatetimeIndex(pd.to_datetime([row[0] for row in rows]))


def next_trading_date(trading_dates: pd.DatetimeIndex, signal_date: str):
    candidates = trading_dates[trading_dates > pd.Timestamp(signal_date)]
    return candidates[0] if len(candidates) else None


def load_period_prices(
    codes: list[str],
    start_date,
    end_date,
    trading_dates: pd.DatetimeIndex,
) -> pd.DataFrame:
    if not codes:
        return pd.DataFrame(index=trading_dates)
    start_text = pd.Timestamp(start_date).strftime("%Y-%m-%d")
    end_text = pd.Timestamp(end_date).strftime("%Y-%m-%d")
    frames = []
    with sqlite3.connect(MARKET_DB_PATH) as conn:
        for start in range(0, len(codes), 800):
            batch = codes[start:start + 800]
            marks = ",".join("?" for _ in batch)
            frames.append(pd.read_sql_query(
                f"""
                SELECT trade_date, wind_code, close
                FROM daily_prices
                WHERE adjusted = 'F' AND trade_date >= ? AND trade_date <= ?
                  AND wind_code IN ({marks})
                """,
                conn,
                params=[start_text, end_text, *batch],
            ))
    raw = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
    if raw.empty:
        return pd.DataFrame(index=trading_dates, columns=codes, dtype=float)
    raw["trade_date"] = pd.to_datetime(raw["trade_date"])
    matrix = raw.pivot(index="trade_date", columns="wind_code", values="close")
    return matrix.reindex(index=trading_dates, columns=codes).ffill()


def fetch_benchmark_close(code: str, name: str, start_date: str, end_date: str) -> pd.Series:
    try:
        from WindPy import w
    except ImportError as exc:
        raise RuntimeError("WindPy不可用，无法取得沪深300基准行情") from exc
    result = w.start()
    if getattr(result, "ErrorCode", 0) != 0:
        raise RuntimeError(f"Wind启动失败: {getattr(result, 'ErrorCode', None)}")
    try:
        data = w.wsd(code, "close", start_date, end_date, "")
        if data.ErrorCode != 0 or not data.Data:
            raise RuntimeError(f"基准行情获取失败: {code}, ErrorCode={data.ErrorCode}")
        return pd.Series(
            pd.to_numeric(data.Data[0], errors="coerce"),
            index=pd.DatetimeIndex(pd.to_datetime(data.Times)),
            name=name,
        ).dropna()
    finally:
        w.close()


def performance_metrics(returns: pd.Series, trading_days_per_year: int = 252) -> dict:
    returns = pd.Series(returns).dropna()
    if returns.empty:
        return {}
    nav = (1 + returns).cumprod()
    if isinstance(returns.index, pd.DatetimeIndex) and len(returns) > 1:
        years = max((returns.index[-1] - returns.index[0]).days / 365.25, 1 / 365.25)
    else:
        years = max(len(returns) / trading_days_per_year, 1 / trading_days_per_year)
    cagr = nav.iloc[-1] ** (1 / years) - 1
    volatility = returns.std(ddof=1) * math.sqrt(trading_days_per_year)
    sharpe = (
        returns.mean() / returns.std(ddof=1) * math.sqrt(trading_days_per_year)
        if returns.std(ddof=1) > 0 else np.nan
    )
    drawdown = nav / nav.cummax() - 1
    max_drawdown = drawdown.min()
    return {
        "累计收益": float(nav.iloc[-1] - 1),
        "年化收益": float(cagr),
        "年化波动": float(volatility),
        "夏普比率_无风险利率0": float(sharpe),
        "最大回撤": float(max_drawdown),
        "卡玛比率": float(cagr / abs(max_drawdown)) if max_drawdown < 0 else np.nan,
    }


def run_backtest(
    start_date: str,
    end_date: str,
    cost_rate: float,
    config: StrategyConfig,
):
    signal_dates = load_signal_dates(start_date, end_date)
    trading_dates = load_trading_dates(start_date, end_date)
    schedules = []
    previous_constituents = None
    for signal_date in signal_dates:
        entry_date = next_trading_date(trading_dates, signal_date)
        if entry_date is None:
            continue
        benchmark, benchmark_date = load_local_benchmark_industry_weights(signal_date)
        if benchmark is None:
            continue
        selected, _, _, diagnostics = build_index(
            signal_date,
            config,
            previous_constituents=previous_constituents,
        )
        schedules.append({
            "signal_date": pd.Timestamp(signal_date),
            "entry_date": entry_date,
            # 回测计算只保留代码和权重，不再累积历次成分明细。
            "selected": selected[["wind_code", "指数权重"]].copy(),
            "benchmark_date": benchmark_date,
            "universe_count": diagnostics["股票池数量"],
            "eligible_count": diagnostics["符合最低数据要求数量"],
            "valuation_date": diagnostics["市值截面日期"],
            "guidance_count": diagnostics["成分股有效业绩预告快报数量"],
            "notice_count": diagnostics["成分股有效业绩预告数量"],
            "express_count": diagnostics["成分股有效业绩快报数量"],
        })
        previous_constituents = set(selected["wind_code"])
        print(f"完成截面 {signal_date}: 入选{len(selected)}只", flush=True)

    if not schedules:
        raise RuntimeError("没有可回测的月度调仓截面")

    latest_date = trading_dates.max()
    gross_nav = 1.0
    net_nav = 1.0
    daily_rows = []
    period_rows = []
    previous_end_weights = pd.Series(dtype=float)
    for idx, schedule in enumerate(schedules):
        entry = schedule["entry_date"]
        exit_date = schedules[idx + 1]["entry_date"] if idx + 1 < len(schedules) else latest_date
        if exit_date <= entry:
            continue
        selected = schedule["selected"].set_index("wind_code")
        codes = selected.index.tolist()
        period_dates = trading_dates[(trading_dates >= entry) & (trading_dates <= exit_date)]
        prices = load_period_prices(codes, entry, exit_date, period_dates)
        entry_prices = prices.loc[entry].replace(0, np.nan)
        tradable = entry_prices.notna()
        weights = selected.loc[tradable, "指数权重"].astype(float)
        weights /= weights.sum()
        prices = prices.loc[:, weights.index]
        relative = prices.divide(entry_prices.loc[weights.index], axis=1)
        portfolio_factor = relative.mul(weights, axis=1).sum(axis=1)

        all_codes = previous_end_weights.index.union(weights.index)
        before = previous_end_weights.reindex(all_codes, fill_value=0.0)
        after = weights.reindex(all_codes, fill_value=0.0)
        turnover = float((after - before).abs().sum() / 2.0) if idx > 0 else 1.0
        cost = turnover * cost_rate
        gross_path = gross_nav * portfolio_factor
        net_path = net_nav * (1 - cost) * portfolio_factor
        for day in period_dates:
            daily_rows.append({
                "日期": day,
                "策略毛净值": float(gross_path.loc[day]),
                "策略费后净值": float(net_path.loc[day]),
            })
        gross_nav = float(gross_path.iloc[-1])
        net_nav = float(net_path.iloc[-1])
        period_return = float(portfolio_factor.iloc[-1] - 1)
        period_rows.append({
            "信号日": schedule["signal_date"],
            "交易日": entry,
            "期末日": exit_date,
            "基准权重日": schedule["benchmark_date"],
            "市值截面日": schedule["valuation_date"],
            "股票池数量": schedule["universe_count"],
            "合格数量": schedule["eligible_count"],
            "目标成分数量": len(selected),
            "交易成分数量": int(tradable.sum()),
            "入场缺价数量": int((~tradable).sum()),
            "有效业绩预告快报数量": schedule["guidance_count"],
            "有效业绩预告数量": schedule["notice_count"],
            "有效业绩快报数量": schedule["express_count"],
            "单边换手率": turnover,
            "交易成本": cost,
            "区间毛收益": period_return,
            "区间费后收益": float((1 - cost) * (1 + period_return) - 1),
        })
        end_values = weights * relative.iloc[-1]
        previous_end_weights = end_values / end_values.sum()

    daily = pd.DataFrame(daily_rows).drop_duplicates("日期", keep="last").sort_values("日期")
    start_text = daily["日期"].min().strftime("%Y-%m-%d")
    end_text = daily["日期"].max().strftime("%Y-%m-%d")
    benchmark_price = fetch_benchmark_close("000300.SH", "沪深300价格指数", start_text, end_text)
    benchmark_total = fetch_benchmark_close("H00300.CSI", "沪深300全收益指数", start_text, end_text)
    daily_index = pd.DatetimeIndex(daily["日期"])
    benchmark_price = benchmark_price.reindex(daily_index).ffill().bfill()
    benchmark_total = benchmark_total.reindex(daily_index).ffill().bfill()
    daily["沪深300价格净值"] = benchmark_price.to_numpy() / benchmark_price.iloc[0]
    daily["沪深300全收益净值"] = benchmark_total.to_numpy() / benchmark_total.iloc[0]
    daily["策略毛收益"] = daily["策略毛净值"].pct_change().fillna(daily["策略毛净值"].iloc[0] - 1)
    daily["策略费后收益"] = daily["策略费后净值"].pct_change().fillna(daily["策略费后净值"].iloc[0] - 1)
    daily["沪深300价格收益"] = daily["沪深300价格净值"].pct_change().fillna(0.0)
    daily["沪深300全收益"] = daily["沪深300全收益净值"].pct_change().fillna(0.0)
    daily["费后超额净值"] = daily["策略费后净值"] / daily["沪深300全收益净值"]
    daily["策略费后回撤"] = daily["策略费后净值"] / daily["策略费后净值"].cummax() - 1

    periods = pd.DataFrame(period_rows)
    daily_lookup = daily.set_index("日期")
    periods["沪深300价格区间收益"] = [
        float(daily_lookup.loc[row.期末日, "沪深300价格净值"] / daily_lookup.loc[row.交易日, "沪深300价格净值"] - 1)
        for row in periods.itertuples(index=False)
    ]
    periods["沪深300全收益区间收益"] = [
        float(daily_lookup.loc[row.期末日, "沪深300全收益净值"] / daily_lookup.loc[row.交易日, "沪深300全收益净值"] - 1)
        for row in periods.itertuples(index=False)
    ]
    periods["费后超额收益"] = periods["区间费后收益"] - periods["沪深300全收益区间收益"]

    metrics = {
        "策略毛收益": performance_metrics(daily_lookup["策略毛收益"]),
        "策略费后收益": performance_metrics(daily_lookup["策略费后收益"]),
        "沪深300价格指数": performance_metrics(daily_lookup["沪深300价格收益"]),
        "沪深300全收益指数": performance_metrics(daily_lookup["沪深300全收益"]),
    }
    fee_metrics = metrics["策略费后收益"]
    fee_metrics["月度胜率"] = float((periods["区间费后收益"] > 0).mean())
    fee_metrics["月度跑赢率"] = float((periods["费后超额收益"] > 0).mean())
    fee_metrics["平均单边月换手"] = float(periods["单边换手率"].mean())
    fee_metrics["年化单边换手"] = float(periods["单边换手率"].mean() * 12)
    fee_metrics["累计交易成本估计"] = float(periods["交易成本"].sum())
    active_returns = daily_lookup["策略费后收益"] - daily_lookup["沪深300全收益"]
    tracking_error = active_returns.std(ddof=1) * math.sqrt(252)
    metrics["相对沪深300全收益"] = {
        "累计超额收益": float(daily["策略费后净值"].iloc[-1] - daily["沪深300全收益净值"].iloc[-1]),
        "几何年化超额收益": float(
            daily["费后超额净值"].iloc[-1]
            ** (365.25 / (daily["日期"].iloc[-1] - daily["日期"].iloc[0]).days) - 1
        ),
        "跟踪误差": float(tracking_error),
        "信息比率": float(active_returns.mean() / active_returns.std(ddof=1) * math.sqrt(252)),
        "日收益相关系数": float(daily["策略费后收益"].corr(daily["沪深300全收益"])),
    }
    metrics["模型配置"] = {
        "个股评分模式": config.scoring_mode,
        "个股因子比较范围": config.stock_factor_scope,
        "成分选择模式": config.constituent_selection_mode,
        "行业名额最低总分": config.industry_quota_min_score,
        "行业权重模式": config.industry_target_mode,
        "行业权重最大相对偏离": config.max_industry_relative_deviation,
        "行业比较方法": "相对行业自身历史",
        "行业历史窗口月数": config.industry_history_months,
        "行业历史最少观察月数": config.industry_history_min_observations,
        "成分缓冲比例": config.constituent_buffer_ratio,
        "个股评分倾斜强度": config.stock_score_tilt,
        "动量因子权重": config.hybrid_momentum_weight,
        "低规模因子权重": config.hybrid_low_size_weight,
        "行业动量权重": config.industry_momentum_weight,
        "业绩预告快报功能": config.earnings_guidance_enabled,
        "业绩预告置信权重": config.earnings_notice_confidence,
        "业绩快报置信权重": config.earnings_express_confidence,
    }

    yearly_nav = daily_lookup[["策略毛净值", "策略费后净值", "沪深300价格净值", "沪深300全收益净值"]].resample("YE").last()
    yearly_base = yearly_nav.shift(1)
    yearly_base.iloc[0] = 1.0
    yearly = (yearly_nav / yearly_base - 1).rename(columns={
        "策略毛净值": "策略毛收益",
        "策略费后净值": "策略费后收益",
        "沪深300价格净值": "沪深300价格收益",
        "沪深300全收益净值": "沪深300全收益",
    })
    yearly.index = yearly.index.year
    yearly["费后超额收益"] = yearly["策略费后收益"] - yearly["沪深300全收益"]
    yearly.index.name = "年度"
    return daily, periods, yearly.reset_index(), metrics


def write_backtest_outputs(
    output_path: Path,
    daily: pd.DataFrame,
    periods: pd.DataFrame,
    yearly: pd.DataFrame,
    metrics: dict,
) -> Path:
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    output_path.parent.mkdir(parents=True, exist_ok=True)
    metric_rows = []
    for category, values in metrics.items():
        if category == "模型配置":
            continue
        for metric, value in values.items():
            metric_rows.append({"类别": category, "指标": metric, "数值": value})
    metric_frame = pd.DataFrame(metric_rows)
    parameter_frame = pd.DataFrame(
        list(metrics.get("模型配置", {}).items()), columns=["参数", "值"]
    )

    writer_kwargs = {
        "engine": "openpyxl",
        "mode": "a" if output_path.exists() else "w",
    }
    if output_path.exists():
        writer_kwargs["if_sheet_exists"] = "replace"
    with pd.ExcelWriter(output_path, **writer_kwargs) as writer:
        workbook = writer.book
        if "历次成分权重" in workbook.sheetnames:
            del workbook["历次成分权重"]
        metric_frame.to_excel(writer, sheet_name="绩效汇总", index=False, startrow=2)
        yearly.to_excel(writer, sheet_name="年度收益", index=False)
        periods.to_excel(writer, sheet_name="月度调仓", index=False)
        daily.to_excel(writer, sheet_name="每日净值", index=False)
        parameter_frame.to_excel(writer, sheet_name="模型参数", index=False)

        dark_blue = "17365D"
        medium_blue = "4472C4"
        light_blue = "D9EAF7"
        light_gray = "E7E6E6"
        white = "FFFFFF"
        green = "E2F0D9"
        red = "FCE4D6"
        thin_gray = Side(style="thin", color="D9E1F2")

        summary = workbook["绩效汇总"]
        summary.merge_cells("A1:C1")
        summary["A1"] = "全A股质量成长800指数回测绩效汇总"
        summary["A1"].font = Font(name="Microsoft YaHei", size=16, bold=True, color=white)
        summary["A1"].fill = PatternFill("solid", fgColor=dark_blue)
        summary["A1"].alignment = Alignment(horizontal="left", vertical="center")
        summary.row_dimensions[1].height = 28

        backtest_sheet_names = ["绩效汇总", "年度收益", "月度调仓", "每日净值", "模型参数"]
        for sheet_name in backtest_sheet_names:
            sheet = workbook[sheet_name]
            sheet.sheet_view.showGridLines = False
            sheet.freeze_panes = "A2" if sheet.title != "绩效汇总" else "A4"
            header_row = 3 if sheet.title == "绩效汇总" else 1
            for cell in sheet[header_row]:
                if cell.value is not None:
                    cell.font = Font(name="Microsoft YaHei", bold=True, color=white)
                    cell.fill = PatternFill("solid", fgColor=medium_blue)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = Border(bottom=thin_gray)
            last_column = get_column_letter(sheet.max_column)
            sheet.auto_filter.ref = (
                f"A3:{last_column}{sheet.max_row}"
                if sheet.title == "绩效汇总"
                else f"A1:{last_column}{sheet.max_row}"
            )
            for row in sheet.iter_rows(min_row=header_row + 1):
                for cell in row:
                    cell.font = Font(name="Microsoft YaHei", size=10)
                    cell.alignment = Alignment(
                        horizontal="left" if isinstance(cell.value, str) else "right",
                        vertical="center",
                    )

        for row in range(4, summary.max_row + 1):
            metric = str(summary.cell(row, 2).value or "")
            value_cell = summary.cell(row, 3)
            if any(key in metric for key in ["夏普", "卡玛", "信息比率", "相关系数"]):
                value_cell.number_format = "0.00"
            else:
                value_cell.number_format = "0.00%"
            if "策略费后收益" == summary.cell(row, 1).value:
                summary.cell(row, 1).fill = PatternFill("solid", fgColor=light_blue)
            if "相对沪深300全收益" == summary.cell(row, 1).value:
                summary.cell(row, 1).fill = PatternFill("solid", fgColor=green)

        def format_table_sheet(sheet_name: str) -> None:
            sheet = workbook[sheet_name]
            headers = {cell.column: str(cell.value or "") for cell in sheet[1]}
            for column, header in headers.items():
                letter = get_column_letter(column)
                values = list(sheet.iter_cols(min_col=column, max_col=column, min_row=2, values_only=True))[0]
                if "日期" in header or header.endswith("日"):
                    for cell in sheet[letter][1:]:
                        cell.number_format = "yyyy-mm-dd"
                elif any(key in header for key in ["收益", "权重", "换手", "回撤", "波动", "成本"]):
                    for cell in sheet[letter][1:]:
                        cell.number_format = "0.00%"
                elif any(key in header for key in ["得分", "总分", "排名"]):
                    for cell in sheet[letter][1:]:
                        cell.number_format = "0.00"
                max_length = max(
                    [len(header)]
                    + [len(str(value)) for value in values[:500] if value is not None]
                )
                sheet.column_dimensions[letter].width = min(max(max_length + 2, 10), 22)

        for sheet_name in ["年度收益", "月度调仓", "每日净值", "模型参数"]:
            format_table_sheet(sheet_name)

        summary.column_dimensions["A"].width = 24
        summary.column_dimensions["B"].width = 24
        summary.column_dimensions["C"].width = 15
        parameter_sheet = workbook["模型参数"]
        parameter_sheet.column_dimensions["A"].width = 28
        parameter_sheet.column_dimensions["B"].width = 18
        for cell in parameter_sheet["B"][1:]:
            if isinstance(cell.value, float):
                cell.number_format = "0.00%"

        daily_sheet = workbook["每日净值"]
        headers = {cell.value: cell.column for cell in daily_sheet[1]}
        chart = LineChart()
        chart.title = "策略费后净值与沪深300全收益净值"
        chart.style = 13
        chart.height = 9
        chart.width = 17
        chart.y_axis.title = "净值"
        chart.x_axis.title = "日期"
        chart.legend.position = "b"
        date_ref = Reference(daily_sheet, min_col=headers["日期"], min_row=2, max_row=daily_sheet.max_row)
        for column_name in ["策略费后净值", "沪深300全收益净值"]:
            data_ref = Reference(
                daily_sheet,
                min_col=headers[column_name],
                min_row=1,
                max_row=daily_sheet.max_row,
            )
            chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(date_ref)
        summary.add_chart(chart, "E3")

        for sheet_name in backtest_sheet_names:
            sheet = workbook[sheet_name]
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0

    for legacy_name in [
        "每日净值.csv",
        "月度调仓与收益.csv",
        "历次成分与权重.csv",
        "年度收益.csv",
        "绩效汇总.json",
    ]:
        legacy_path = output_path.parent / legacy_name
        if legacy_path.exists():
            legacy_path.unlink()
    return output_path


def hybrid_config(top_n: int) -> StrategyConfig:
    return StrategyConfig(
        max_constituents=top_n,
        scoring_mode="hybrid",
        industry_target_mode="industry_score",
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="行业比较＋混合评分全A股质量成长800指数：单文件构建与回测"
    )
    parser.add_argument(
        "--task",
        choices=["build", "backtest", "both"],
        default="both",
        help="默认both：依次构建单期指数并运行历史回测",
    )
    parser.add_argument("--as-of", default=date.today().isoformat())
    parser.add_argument("--start-date", default="2022-01-01")
    parser.add_argument("--end-date", default=date.today().isoformat())
    parser.add_argument("--top-n", type=int, default=800)
    parser.add_argument("--cost-rate", type=float, default=TRADING_COST_RATE)
    parser.add_argument(
        "--stock-factor-scope",
        choices=["blended", "industry_only"],
        default="blended",
        help="个股因子比较范围；默认保留行业内与全市场混合评分",
    )
    parser.add_argument(
        "--constituent-selection-mode",
        choices=["global_top", "industry_quota"],
        default="global_top",
        help="成分选择模式；industry_quota 为实验性的行业参考名额模式",
    )
    parser.add_argument(
        "--industry-quota-min-score",
        type=float,
        default=45.0,
        help="行业参考名额模式的个股最低总分",
    )
    parser.add_argument(
        "--disable-earnings-guidance",
        action="store_true",
        help="禁用业绩预告/快报对成长和周期得分的修正",
    )
    parser.add_argument(
        "--earnings-notice-confidence",
        type=float,
        default=0.40,
        help="业绩预告置信权重，默认0.40",
    )
    parser.add_argument(
        "--earnings-express-confidence",
        type=float,
        default=0.70,
        help="业绩快报置信权重，默认0.70",
    )
    parser.add_argument("--benchmark-weights", type=Path)
    parser.add_argument("--previous-components", type=Path)
    parser.add_argument("--build-output", type=Path)
    parser.add_argument(
        "--backtest-output-dir",
        type=Path,
        help="兼容旧参数：指定整合后主输出文件的目录",
    )
    return parser.parse_args()


def run_single_build(args: argparse.Namespace, config: StrategyConfig) -> Path:
    as_of = normalize_date(args.as_of)
    output_path = args.build_output or (
        OUTPUT_DIR / f"全A股质量成长800_行业比较混合评分_{as_of.replace('-', '')}.xlsx"
    )
    previous_path = args.previous_components
    previous_source = "命令行指定"
    if previous_path is None:
        previous_path = find_latest_previous_components(
            as_of,
            [output_path.parent, OUTPUT_DIR, LEGACY_OUTPUT_DIR],
        )
        previous_source = "自动查找"
    previous_constituents = (
        load_previous_constituents(previous_path) if previous_path is not None else None
    )
    selected, industry_summary, all_scores, diagnostics = build_index(
        as_of,
        config,
        args.benchmark_weights,
        previous_constituents,
    )
    export_result(selected, industry_summary, all_scores, diagnostics, config, output_path)
    print(json.dumps(diagnostics, ensure_ascii=False, indent=2))
    if previous_path is None:
        print("提示：未找到早于本次截止日的历史成分文件，本次未应用历史成分缓冲。")
    else:
        print(f"上一期成分文件（{previous_source}）：{previous_path}")
    print(f"指数构建结果：{output_path}")
    return output_path


def run_history(
    args: argparse.Namespace,
    config: StrategyConfig,
    output_path: Path,
) -> Path:
    results = run_backtest(args.start_date, args.end_date, args.cost_rate, config)
    workbook_path = write_backtest_outputs(output_path, *results)
    print(json.dumps(results[-1], ensure_ascii=False, indent=2))
    print(f"回测结果已合并至主输出文件：{workbook_path}")
    return workbook_path


def main() -> None:
    args = parse_args()
    config = replace(
        hybrid_config(args.top_n),
        stock_factor_scope=args.stock_factor_scope,
        constituent_selection_mode=args.constituent_selection_mode,
        industry_quota_min_score=args.industry_quota_min_score,
        earnings_guidance_enabled=not args.disable_earnings_guidance,
        earnings_notice_confidence=args.earnings_notice_confidence,
        earnings_express_confidence=args.earnings_express_confidence,
    )
    validate_config(config)
    as_of = normalize_date(args.as_of)
    output_name = f"全A股质量成长800_行业比较混合评分_{as_of.replace('-', '')}.xlsx"
    output_path = args.build_output or (
        args.backtest_output_dir / output_name
        if args.backtest_output_dir is not None
        else OUTPUT_DIR / output_name
    )
    args.build_output = output_path
    if args.task in {"build", "both"}:
        output_path = run_single_build(args, config)
    elif not output_path.exists():
        print("未找到当期主输出文件，先生成单期指数结果再合并回测。")
        output_path = run_single_build(args, config)
    if args.task in {"backtest", "both"}:
        run_history(args, config, output_path)


if __name__ == "__main__":
    main()
